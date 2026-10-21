"""Word documents: contracts, forms, status tables. Each test pins a rule a real document needs."""
import shutil
import textwrap
from pathlib import Path

import pytest
from docx.oxml.ns import qn

from servantium_documents import docs as D, kit
from servantium_documents.brand import hexof
from servantium_documents.spec import SpecError

PKG = Path(__file__).resolve().parents[1]
EX = PKG / 'examples'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def write(tmp_path, body, name='doc.yaml'):
    p = tmp_path / name
    p.write_text(textwrap.dedent(body))
    return p


def make(path):
    """Build the docx in memory (no LibreOffice), return the python-docx Document."""
    spec = D.load(path)
    out = Path(spec['_path']).parent / 'out'
    out.mkdir(exist_ok=True)
    return D.Doc(spec).build(out)


def text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return '\n'.join(parts)


CONTRACT = '''
    document: Contract
    client: Acme Labs
    page_numbers: x_of_y
    cover: {title: Services Agreement}
    body:
      - clauses:
          - title: Definitions
            clauses:
              - '"Services" means the work in a SOW. The "Services" are performed remotely.'
          - title: Fees
            id: fees
            clauses:
              - title: Invoices
                id: invoices
                text: Paid within 30 days, see {ref:sched}.
                clauses: [First, {text: Second, clauses: [Sub a, Sub b]}]
      - signature: {parties: [{company: Acme Labs}, {}]}
      - schedule: {label: Schedule A, title: Rates, id: sched}
      - clauses:
          - title: Rates
            clauses: ['Rates per {ref:invoices}.']
'''


# ── every example builds ─────────────────────────────────────────────────────
@pytest.mark.parametrize('name', ['doc-msa', 'doc-sow', 'doc-status', 'doc-invoice'])
def test_examples_build(tmp_path, name):
    shutil.copy(EX / f'{name}.yaml', tmp_path / f'{name}.yaml')
    doc = make(tmp_path / f'{name}.yaml')
    assert doc.core_properties.author == kit.BRAND_NAME
    assert not D.banned_hits(D.load(tmp_path / f'{name}.yaml'))


@pytest.mark.skipif(not (shutil.which('soffice') or Path('/Applications/LibreOffice.app').exists()),
                    reason='LibreOffice not installed')
def test_build_writes_docx_and_pdf(tmp_path):
    shutil.copy(EX / 'doc-invoice.yaml', tmp_path / 'doc-invoice.yaml')
    r = D.build(tmp_path / 'doc-invoice.yaml')
    assert r['docx'].exists() and r['pdf'].exists() and r['docx'].parent == tmp_path / 'out'
    assert r['pages'] == 1


# ── contracts ────────────────────────────────────────────────────────────────
def test_clauses_use_real_multilevel_numbering(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    num = doc.part.numbering_part.element
    clauses = [a for a in num.findall(qn('w:abstractNum'))
               if a.find(qn('w:name')) is not None and a.find(qn('w:name')).get(qn('w:val')) == 'Clauses']
    assert len(clauses) == 1
    texts = [l.find(qn('w:lvlText')).get(qn('w:val')) for l in clauses[0].findall(qn('w:lvl'))]
    assert texts[:3] == ['%1.', '%1.%2', '%1.%2.%3']
    # every abstractNum precedes every num (schema order)
    tags = [c.tag for c in num]
    assert max(i for i, t in enumerate(tags) if t == qn('w:abstractNum')) < tags.index(qn('w:num'))
    # article headings are the built-in Heading 2, so Google keeps them as headings
    heads = [p for p in doc.paragraphs if p.style.name == 'Heading 2']
    assert [h.text for h in heads] == ['Definitions', 'Fees', 'Rates']
    assert all(h._p.pPr.numPr is not None for h in heads)


def test_schedule_restarts_numbering(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    ids = [p._p.pPr.numPr.numId.val for p in doc.paragraphs
           if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(set(ids)) == 2                               # main body, then Schedule A
    restart = [n for n in doc.part.numbering_part.element.findall(qn('w:num'))
               if int(n.get(qn('w:numId'))) == ids[-1]][0]
    assert restart.find(qn('w:lvlOverride')).find(qn('w:startOverride')).get(qn('w:val')) == '1'
    sched = [p for p in doc.paragraphs if p.style.name == 'Heading 1'][-1]
    assert sched.text == 'Schedule A: Rates' and sched.paragraph_format.page_break_before


def test_cross_references_print_the_number(tmp_path):
    t = text(make(write(tmp_path, CONTRACT)))
    assert 'Paid within 30 days, see Schedule A.' in t
    assert 'Rates per Section 2.1 of the Agreement.' in t


def test_unknown_reference_is_an_error(tmp_path):
    p = write(tmp_path, CONTRACT.replace('{ref:sched}', '{ref:nowhere}'))
    with pytest.raises(SpecError, match='nowhere'):
        make(p)


def test_defined_term_is_bold_where_first_defined_only(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    p = [p for p in doc.paragraphs if 'means the work' in p.text][0]
    bold = [r.text for r in p.runs if r.bold]
    assert bold == ['“Services”']


def test_signature_block_has_both_parties_and_lines(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    sig = [t for t in doc.tables if 'Acme Labs' in t.rows[0].cells[0].text][0]
    assert kit.ACTIVE.legal_name in sig.rows[0].cells[3].text
    labels = [sig.rows[r].cells[0].text for r in range(1, len(sig.rows))]
    assert labels == ['By:', 'Name:', 'Title:', 'Date:']
    assert sig.rows[1].cells[1]._tc.tcPr.find(qn('w:tcBorders')).find(qn('w:bottom')) is not None


def test_footer_is_brand_text_and_page_x_of_y(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    xml = doc.sections[0].footer._element.xml
    assert ' NUMPAGES ' in xml and ' PAGE ' in xml
    assert doc.sections[0].footer.paragraphs[0].text.startswith(kit.FOOTER_TEXT.split('  ')[0])


def test_draft_label_and_watermark(tmp_path):
    p = write(tmp_path, CONTRACT.replace('page_numbers: x_of_y', 'draft: true\n    watermark: Draft'))
    doc = make(p)
    hdr = doc.sections[0].header
    assert hdr.paragraphs[0].text.startswith('DRAFT')
    assert 'PowerPlusWaterMarkObject' in hdr._element.xml
    assert 'PowerPlusWaterMarkObject' in doc.sections[0].first_page_header._element.xml


# ── cover ────────────────────────────────────────────────────────────────────
def test_prepared_for_logo_lockup(tmp_path):
    from PIL import Image
    Image.new('RGB', (400, 100), (200, 60, 40)).save(tmp_path / 'acme.png')
    p = write(tmp_path, '''
        document: Plan
        client: Acme Labs
        prepared_for: {name: Acme Labs, logo: acme.png}
        cover: {title: A plan, meta: [[Date, 'October 1, 2026']]}
        body: [{p: Hello.}]
    ''')
    doc = make(p)
    lock = doc.tables[0]
    assert 'PREPARED FOR' in lock.cell(0, 1).text
    assert len(lock.cell(0, 1)._tc.xpath('.//pic:pic')) == 1           # the client's logo
    meta = doc.tables[1]
    assert [meta.cell(i, 0).text for i in range(2)] == ['PREPARED FOR', 'PREPARED BY']
    assert meta.cell(1, 1).text == kit.BRAND_NAME


def test_missing_prepared_for_logo_is_an_error(tmp_path):
    p = write(tmp_path, '''
        document: Plan
        client: Acme Labs
        prepared_for: {name: Acme Labs, logo: nope.png}
        cover: {title: A plan}
        body: [{p: Hello.}]
    ''')
    with pytest.raises(SpecError, match='nope.png'):
        D.load(p)


# ── tables ───────────────────────────────────────────────────────────────────
STATUS = '''
    document: Status
    client: Acme Labs
    cover: {layout: plain, title: Status}
    body:
      - table:
          head: [Area, Status, Trend, Fee]
          rag: [2]
          money: [3]
          rows: [[Scope, Green, At risk, 1200], [Budget, Red, Complete, 3400.5], [People, Amber, On track, 0]]
          foot: [[Total, '', '', 4600.5]]
'''


def test_rag_cells_are_shaded_in_brand_status_colours(tmp_path):
    t = make(write(tmp_path, STATUS)).tables[0]
    fill = lambda r, c: t.cell(r, c)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))
    assert fill(1, 1) == hexof(kit.SUCCESS)                # status green, not the brand primary
    assert fill(2, 1) == hexof(kit.CORAL)
    assert fill(3, 1) == hexof(kit.AMBER)
    assert fill(1, 2) == hexof(kit.AMBER)                  # "At risk" in a declared rag column
    assert fill(2, 2) == hexof(kit.SUCCESS)                # "Complete"
    assert len(t.cell(2, 1)._tc.tcPr.findall(qn('w:shd'))) == 1   # chip replaces the zebra fill


def test_table_header_repeats_numbers_sit_right_and_print_as_money(tmp_path):
    t = make(write(tmp_path, STATUS)).tables[0]
    assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    assert t.cell(2, 3).text == '$3,400.50'
    assert t.cell(4, 3).text == '$4,600.50'                # foot row
    assert t.cell(1, 3).paragraphs[0].alignment == 2       # RIGHT
    assert t.cell(0, 3).paragraphs[0].alignment == 2       # header follows its column


def test_short_table_is_kept_on_one_page(tmp_path):
    t = make(write(tmp_path, STATUS)).tables[0]
    keeps = [c.paragraphs[0].paragraph_format.keep_with_next for r in t.rows[:-1] for c in r.cells]
    assert all(keeps)


# ── forms ────────────────────────────────────────────────────────────────────
def test_invoice_computes_amounts_and_totals(tmp_path):
    p = write(tmp_path, '''
        document: Invoice 1
        client: Acme Labs
        cover: {layout: none}
        body:
          - form:
              type: invoice
              number: INV-1
              fields: [[Date, 'October 1, 2026']]
              items: {rows: [[Workshop, 2, 4500], [Travel, 1, 240.6]]}
              totals: {tax: {label: Sales tax, rate: 0.05}, paid: 1000}
    ''')
    doc = make(p)
    t = text(doc)
    assert '$9,000.00' in t                                # 2 x 4,500
    assert '$9,240.60' in t                                # subtotal
    assert '$462.03' in t                                  # 5% tax
    assert 'BALANCE DUE' in t and '$8,702.63' in t
    assert any(p.style.name == 'Title' and 'Invoice' in p.text
               for tb in doc.tables for c in tb.rows[0].cells for p in c.paragraphs)


def test_form_type_is_checked(tmp_path):
    p = write(tmp_path, '''
        document: X
        client: Acme Labs
        body: [{form: {type: receipt}}]
    ''')
    with pytest.raises(SpecError, match='form type'):
        D.load(p)


def test_change_order_shows_revised_value(tmp_path):
    p = write(tmp_path, '''
        document: CO
        client: Acme Labs
        body:
          - form:
              type: change_order
              items: {rows: [[More work, 10, 200]]}
              totals: {original: 50000}
              approvals: [{company: Acme Labs}, {}]
    ''')
    t = text(make(p))
    assert 'Revised contract value'.upper() in t and '$52,000.00' in t


# ── Word opens it without repair ─────────────────────────────────────────────
def test_property_elements_are_in_schema_order(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    make(write(tmp_path, STATUS, 'status.yaml'))
    roots = [doc.element] + [s.footer._element for s in doc.sections] + [s.header._element for s in doc.sections]
    for root in roots:
        for local, rank in D._ORDER.items():
            for el in root.iter(f'{W}{local}'):
                seq = [k.tag[len(W):] for k in el if k.tag.startswith(W) and k.tag[len(W):] in rank]
                assert seq == sorted(seq, key=rank.get), f'{local}: {seq}'
                assert len(seq) == len(set(seq)), f'{local} repeats an element: {seq}'


def test_docs_engine_types_no_brand_name_or_colour():
    src = (PKG / 'src' / 'servantium_documents' / 'docs.py').read_text()
    code = src.split('"""', 2)[2]                           # past the module docstring
    assert "'Servantium" not in code and '"Servantium' not in code
    assert 'RGBColor(0x' not in code


# ── fixes from the first demo batch ──────────────────────────────────────────
def test_rag_green_is_the_status_colour_under_another_brand(tmp_path):
    """Lattica's primary is indigo: a "Green" chip must still print the status green."""
    import subprocess, sys, json
    p = write(tmp_path, STATUS)
    code = f"""
import json
from pathlib import Path
from docx.oxml.ns import qn
from servantium_documents import docs as D, kit
from servantium_documents.brand import hexof
if not Path(kit.TEMPLATE_DOCX).exists():
    print(json.dumps({{'skip': True}})); raise SystemExit
spec = D.load({str(p)!r}); out = Path(spec['_path']).parent / 'out'; out.mkdir(exist_ok=True)
t = D.Doc(spec).build(out).tables[0]
fill = lambda r, c: t.cell(r, c)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))
print(json.dumps({{'green': fill(1, 1), 'red': fill(2, 1), 'success': hexof(kit.SUCCESS),
                   'coral': hexof(kit.CORAL), 'primary': hexof(kit.GREEN), 'brand': kit.ACTIVE.id}}))
"""
    import os
    env = dict(os.environ, SV_BRAND='lattica')
    r = subprocess.run([sys.executable, '-c', code], capture_output=True, text=True, env=env, cwd=PKG)
    assert r.returncode == 0, r.stderr
    got = json.loads(r.stdout.strip().splitlines()[-1])
    if got.get('skip'):
        pytest.skip('lattica document template not built')
    assert got['brand'] == 'lattica'
    assert got['green'] == got['success'] != got['primary']
    assert got['red'] == got['coral']


def test_foot_only_table_has_no_header_and_empty_table_is_skipped(tmp_path):
    doc = make(write(tmp_path, '''
        document: T
        client: Acme Labs
        cover: {layout: plain, title: T}
        body:
          - table: {head: [A, B], rows: [], foot: [[Total, 10]]}
          - table: {rows: []}
          - p: After.
    '''))
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.rows) == 1 and t.cell(0, 0).text == 'Total'
    assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is None


def test_long_table_splits_with_repeated_header(tmp_path):
    rows = ', '.join(f'[Row {i}, {i}]' for i in range(20))
    doc = make(write(tmp_path, f'''
        document: T
        client: Acme Labs
        cover: {{layout: plain, title: T}}
        body:
          - table: {{head: [Item, N], rows: [{rows}], foot: [[Total, 190]]}}
    '''))
    t = doc.tables[0]
    keep = [t.rows[i].cells[0].paragraphs[0].paragraph_format.keep_with_next for i in range(len(t.rows))]
    assert keep[0] is True                                  # header stays with the first row
    assert not any(keep[2:15])                              # the middle may split across pages
    assert keep[-3] and keep[-2]                            # the last rows stay with the total
    assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None


def test_change_order_labels_are_configurable_and_fit(tmp_path):
    doc = make(write(tmp_path, '''
        document: CO
        client: Acme Labs
        cover: {layout: none}
        body:
          - form:
              type: change_order
              items: {rows: [[More work, 10, 200]]}
              totals: {original: 50000, original_label: Contract value to date,
                       revised_label: New total contract value}
    '''))
    tot = [t for t in doc.tables if 'NEW TOTAL CONTRACT VALUE' in t.cell(t.rows.__len__() - 1, 0).text][0]
    assert tot.cell(0, 0).text == 'Contract value to date'
    w = int(tot._tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol'))[0].get(qn('w:w'))) / 1440
    from servantium_documents import kit as K
    assert w >= K.text_w_in('NEW TOTAL CONTRACT VALUE', 8.5, bold=True)


def test_subtotal_is_skipped_when_it_equals_the_total(tmp_path):
    doc = make(write(tmp_path, '''
        document: I
        client: Acme Labs
        cover: {layout: none}
        body:
          - form: {type: invoice, items: {rows: [[Work, 2, 100]]}}
    '''))
    t = text(doc)
    assert 'Subtotal' not in t and 'TOTAL' in t and '$200.00' in t


def test_items_table_is_kept_with_its_totals(tmp_path):
    doc = make(write(tmp_path, '''
        document: I
        client: Acme Labs
        cover: {layout: none}
        body:
          - form: {type: invoice, items: {rows: [[Work, 2, 100], [More, 1, 50]]}, totals: {tax: {rate: 0.1}}}
    '''))
    items = [t for t in doc.tables if t.cell(0, 0).text == 'DESCRIPTION'][0]
    last = items.rows[-1].cells[0].paragraphs[0]
    assert last.paragraph_format.keep_with_next and items._tbl.getnext().pPr.keepNext is not None


def test_form_shows_prepared_for_without_a_cover(tmp_path):
    from PIL import Image
    Image.new('RGB', (300, 100), (200, 60, 40)).save(tmp_path / 'acme.png')
    doc = make(write(tmp_path, '''
        document: I
        client: Acme Labs
        prepared_for: {name: Acme Labs, logo: acme.png}
        cover: {layout: none}
        body:
          - form: {type: invoice, number: INV-1, items: {rows: [[Work, 1, 100]]}}
    '''))
    head = doc.tables[0].cell(0, 1)
    assert 'PREPARED FOR' in head.text and len(head._tc.xpath('.//pic:pic')) == 1


def test_signed_copy_prints_signer_and_date(tmp_path):
    doc = make(write(tmp_path, '''
        document: S
        client: Acme Labs
        cover: {layout: plain, title: S}
        body:
          - signature:
              parties:
                - {company: Acme Labs, title: VP, signed: {by: Jordan Lee, date: 'November 1, 2026'}}
                - {}
    '''))
    sig = doc.tables[0]
    assert sig.cell(1, 0).text == 'By:' and sig.cell(1, 1).text == 'Jordan Lee'
    assert sig.cell(2, 1).text == 'Jordan Lee' and sig.cell(3, 1).text == 'VP'
    assert sig.cell(4, 1).text == 'November 1, 2026'
    assert sig.cell(1, 4).text == ''                        # the other party has not signed


def test_pagebreak_never_doubles_or_leaves_a_blank_page(tmp_path):
    doc = make(write(tmp_path, '''
        document: B
        client: Acme Labs
        cover: {layout: none}
        body:
          - form: {type: invoice, items: {rows: [[Work, 1, 100]]}}
          - pagebreak: true
          - pagebreak: true
          - p: Second page.
          - pagebreak: true
          - schedule: {label: Schedule A}
          - p: Third page.
    '''))
    xml = doc.element.body.xml
    assert xml.count('<w:br w:type="page"/>') == 0           # no break characters at all
    assert xml.count('<w:pageBreakBefore/>') == 2            # one hairline break + the schedule


def test_literal_bracket_is_plain_text(tmp_path):
    doc = make(write(tmp_path, '''
        document: L
        client: Acme Labs
        cover: {layout: plain, title: L}
        body:
          - p: "[Draft] see [the site](https://example.com) and [note]."
    '''))
    p = [p for p in doc.paragraphs if 'Draft' in p.text][0]
    assert p.text.startswith('[Draft] see') and '[note].' in p.text


@pytest.mark.skipif(not (shutil.which('soffice') or Path('/Applications/LibreOffice.app').exists()),
                    reason='LibreOffice not installed')
def test_build_returns_findings_as_strings(tmp_path):
    p = write(tmp_path, '''
        document: F
        client: Acme Labs
        cover: {layout: plain, title: F}
        body: [{p: We leverage the plan.}]
    ''')
    r = D.build(p)
    assert r['findings'] and all(isinstance(f, str) for f in r['findings'])
    assert 'leverage' in r['findings'][0]


def test_h1_breaks_the_page_only_when_asked(tmp_path):
    doc = make(write(tmp_path, '''
        document: H
        client: Acme Labs
        cover: {layout: plain, title: H}
        body:
          - h1: One
          - h1: Two
          - h1: {text: Three, page: true}
    '''))
    h = {p.text: p.paragraph_format.page_break_before for p in doc.paragraphs if p.style.name == 'Heading 1'}
    assert h == {'One': False, 'Two': False, 'Three': True}
    doc = make(write(tmp_path, '''
        document: H
        client: Acme Labs
        h1_new_page: true
        cover: {layout: plain, title: H}
        body: [{h1: One}, {p: x}, {h1: Two}]
    ''', 'h2.yaml'))
    h = {p.text: p.paragraph_format.page_break_before for p in doc.paragraphs if p.style.name == 'Heading 1'}
    assert h == {'One': False, 'Two': True}


def test_fourth_clause_level_prints_letters(tmp_path):
    doc = make(write(tmp_path, CONTRACT))
    a = [a for a in doc.part.numbering_part.element.findall(qn('w:abstractNum'))
         if a.find(qn('w:name')) is not None and a.find(qn('w:name')).get(qn('w:val')) == 'Clauses'][0]
    lvl3 = a.findall(qn('w:lvl'))[3]
    assert lvl3.find(qn('w:numFmt')).get(qn('w:val')) == 'lowerLetter'
    assert lvl3.find(qn('w:isLgl')) is None                 # isLgl would force "(1)"


def test_long_footer_never_pushes_the_page_number_to_a_second_line(tmp_path):
    long = 'Lattica, Inc.  ·  Invoices to Merrow Therapeutics, project LAT-PRJ-0342  ·  ar@lattica.example  ·  more words here'
    doc = make(write(tmp_path, CONTRACT.replace('page_numbers: x_of_y', f'page_numbers: x_of_y\n    footer: "{long}"')))
    fp = doc.sections[0].footer.paragraphs[0]
    label = fp.runs[0].text.rstrip('\t')
    size = fp.runs[0].font.size.pt
    from servantium_documents import kit as K
    assert size >= 7.5
    assert K.text_w_in(label, size) + K.text_w_in('Page 999 of 999', 9) <= 6.5


def test_form_approvals_are_compact(tmp_path):
    doc = make(write(tmp_path, '''
        document: C
        client: Acme Labs
        cover: {layout: none}
        body:
          - form: {type: invoice, items: {rows: [[W, 1, 1]]}, approvals: [{company: Acme Labs}, {}]}
    '''))
    sig = doc.tables[-1]
    h = int(sig.rows[1]._tr.trPr.find(qn('w:trHeight')).get(qn('w:val'))) / 1440
    assert h < 0.5 and sig.cell(1, 0).text == 'By:'
