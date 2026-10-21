"""Engine tests. Each one pins a bug or a rule that a real deck taught us."""
import re
import textwrap
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from servantium_documents import check as C, kit, spec as S, tokens as T
from servantium_documents.render import render

PKG = Path(__file__).resolve().parents[1]
TORTURE = PKG / 'examples' / 'torture-test.yaml'


def write(tmp_path, body):
    p = tmp_path / 'deck.yaml'
    p.write_text(textwrap.dedent(body))
    return p


# ── tokens: nothing typed by hand ────────────────────────────────────────────
def test_tokens_match_verdant():
    css = (PKG.parent / 'verdant' / 'tokens.css').read_text()
    for name, value in re.findall(r'--(color-[a-z0-9-]+)\s*:\s*(#[0-9a-fA-F]{6})\b', css):
        assert T.COLORS[name] == value.upper(), f'{name} drifted; run `npm run tokens`'
    assert T.FONTS['font-display'] == 'Playfair Display'


def test_engine_has_no_color_literals():
    src = (PKG / 'src' / 'servantium_documents').rglob('*.py')
    for f in src:
        if f.name == 'tokens.py':
            continue
        text = f.read_text()
        assert not re.search(r'RGBColor\(0x', text.replace('def _rgb', '')), f'{f.name} types a colour by hand'


# ── the spec format ──────────────────────────────────────────────────────────
def test_five_cards_are_refused(tmp_path):
    cards = '\n'.join(f'                  - {{title: Card {i}}}' for i in range(5))
    p = write(tmp_path, f'''
        deck: T
        client: _Sandbox
        slides:
          - content:
              title: Five cards is one too many for a slide.
              body:
                - cards:
{cards}
    ''')
    with pytest.raises(S.SpecError, match='twos, threes and fours'):
        S.load(p)


def test_comma_split_is_explained(tmp_path):
    p = write(tmp_path, '''
        deck: T
        client: _Sandbox
        slides:
          - content:
              title: A comma inside a flow map splits the value.
              body:
                - steps:
                    - {lead: One, text: Titles only, agreed first.}
                    - {lead: Two, text: Second}
    ''')
    with pytest.raises(S.SpecError, match='comma inside'):
        S.load(p)


def test_dates_stay_strings(tmp_path):
    p = write(tmp_path, '''
        deck: T
        client: _Sandbox
        date: 2026-09-29
        slides:
          - cover: {title: Cover}
    ''')
    assert S.load(p)['date'] == '2026-09-29'


# ── the checks: corrections that became rules ────────────────────────────────
def _findings(tmp_path, content):
    p = write(tmp_path, f'''
        deck: T
        client: _Sandbox
        slides:
          - content:
{textwrap.indent(textwrap.dedent(content), "              ")}
    ''')
    return {f.code for f in C.check(S.load(p)) if f.level == 'error'}


def test_em_dash_is_an_error(tmp_path):
    assert 'DASH' in _findings(tmp_path, '''
        title: Testing is done — time to pick a date.
        body: [{takeaway: {lead: One line.}}]
        notes: n
    ''')


def test_british_spelling_is_an_error(tmp_path):
    assert 'SPELL' in _findings(tmp_path, '''
        title: We organise the programme around three decisions.
        body: [{takeaway: {lead: One line.}}]
        notes: n
    ''')


def test_too_much_text_is_an_error(tmp_path):
    long = ' '.join(['word'] * 22)
    assert 'LINE' in _findings(tmp_path, f'''
        title: A slide whose card line runs far too long.
        body:
          - cards:
              - {{title: One, lines: ["{long}"]}}
              - {{title: Two}}
        notes: n
    ''')


def test_sentence_in_card_kicker_is_an_error(tmp_path):
    assert 'FIT' in _findings(tmp_path, '''
        title: A kicker is a corner tag, not a sentence.
        body:
          - cards:
              - {title: One, kicker: Gate 1 decides what comes next}
              - {title: Two, kicker: '02'}
        notes: n
    ''')


def test_wrapping_cta_is_an_error(tmp_path):
    p = write(tmp_path, '''
        deck: T
        client: _Sandbox
        slides:
          - closing:
              title: Next steps, agreed.
              asks: [One, Two]
              cta: A one-page proposal within a week of your answers
    ''')
    assert 'FIT' in {f.code for f in C.check(S.load(p)) if f.level == 'error'}


def test_two_emphasis_blocks_are_an_error(tmp_path):
    assert 'EMPHASIS' in _findings(tmp_path, '''
        title: A slide can only have one so-what.
        body:
          - takeaway: {lead: First.}
          - callout: {label: Second, text: also a so-what}
        notes: n
    ''')


# ── rendering ────────────────────────────────────────────────────────────────
def test_place_keeps_inherited_width():
    """Setting only top/height on an inherited placeholder wrote a ZERO width, and every
       measured title vanished. The first torture-test render caught it."""
    prs = kit.new_deck()
    s, _ = kit.content(prs, heading='A title that must stay visible.', sub='With a lede.')
    title = kit.placeholder(s, 'title')
    assert title.width > kit.Inches(10), 'title placeholder lost its width'
    assert kit.placeholder(s, kit.PH_LEDE).width > kit.Inches(10)


def test_torture_test_renders_clean(tmp_path):
    spec = S.load(TORTURE)
    assert not [f for f in C.check(spec) if f.level == 'error']
    out = tmp_path / 'torture.pptx'
    findings = render(spec, out, tmp_path)
    assert not [f for f in findings if f.code == 'OVERFLOW'], findings
    prs = Presentation(str(out))
    assert len(prs.slides) == len(spec['slides'])
    titled = [s for s in prs.slides if s.shapes.title is not None and s.shapes.title.text.strip()]
    assert len(titled) >= len(prs.slides) - 1, 'every slide but the quote should carry a real title'
    assert prs.core_properties.identifier == f'spec:{spec["_sha"]}'


def test_overflow_is_reported(tmp_path):
    cards = '\n'.join(f'                  - {{title: Card {i}, lines: [A line, Another line, A third line, A fourth line]}}' for i in range(4))
    p = write(tmp_path, f'''
        deck: T
        client: _Sandbox
        slides:
          - content:
              title: Too much stacked on one slide will overflow.
              lede: A lede that takes up room.
              body:
                - cards:
{cards}
                - table:
                    head: [A, B]
                    rows: [[a, b], [a, b], [a, b], [a, b], [a, b], [a, b], [a, b], [a, b]]
              notes: n
    ''')
    findings = render(S.load(p), tmp_path / 'o.pptx', tmp_path)
    assert any(f.code == 'OVERFLOW' for f in findings)


def test_ai_isms_are_errors(tmp_path):
    assert 'AIISM' in _findings(tmp_path, '''
        title: We quietly unlock value across the engagement landscape.
        body: [{takeaway: {lead: One line.}}]
        notes: n
    ''')


def test_repeated_composition_warns(tmp_path):
    slide = '''
          - content:
              title: A slide that looks like the one before it.
              body: [{cards: [{title: One}, {title: Two}]}, {takeaway: {lead: Same.}}]
              notes: n'''
    p = write(tmp_path, 'deck: T\nclient: _Sandbox\nslides:' + slide * 4)
    codes = {f.code for f in C.check(S.load(p))}
    assert 'VARIETY' in codes


def test_plan_pictures_render(tmp_path):
    """gantt, flow and stack render without overflow."""
    p = write(tmp_path, '''
        deck: T
        client: _Sandbox
        slides:
          - content:
              title: A plan over time is a Gantt chart.
              body:
                - gantt:
                    periods: [Jan, Feb, Mar, Apr]
                    rows:
                      - {name: Lane one, note: A team, bars: [{from: 1, to: 2, label: First, tone: forest}]}
                      - {name: Lane two, bars: [{from: 2, to: 4, label: Second, tone: amber}]}
                    gates: [{after: 2, label: Gate 1}, {after: 4, label: Decision}]
              notes: n
          - content:
              title: Every step's output is the next step's input.
              body:
                - flow:
                    steps: [Quote, SOW, Plan]
                    active: 2
                    inputs: [{icon: request_quote, title: Quote, meta: Q-1, tag: from 01, carried: true}]
                    node: {title: Servantium, verb: Drafts the SOW, chip: AI drafts · you approve}
                    output: {label: SOW, title: Scope as sold, rows: [{key: A, text: One}, {key: Open, text: Two, open: true}], foot: Traced}
              notes: n
          - content:
              title: Tools stand on one record.
              body:
                - stack:
                    tiers: [{label: Work, items: [{name: CRM, icon: cloud}, {name: Chat, icon: forum}]}]
                    foundation: {title: The record, text: Everything lands here., chips: [Quote, Plan]}
              notes: n
    ''')
    out = tmp_path / 'out' / 'deck.pptx'
    findings = render(S.load(p), out)
    assert not [f for f in findings if f.level == 'error']
    assert len(Presentation(str(out)).slides) == 3


def test_shipped_and_split_render(tmp_path):
    from PIL import Image
    for name, color in (('a.jpg', (40, 90, 60)), ('b.jpg', (20, 40, 80))):
        Image.new('RGB', (800, 450), color).save(tmp_path / name)
    p = write(tmp_path, '''
        deck: T
        client: _Sandbox
        slides:
          - content:
              title: We shipped a lot in eight months.
              body:
                - shipped:
                    groups:
                      - {title: One, items: [{icon: check, label: A}, {icon: check, label: B}]}
                      - {title: Two, items: [{label: C}, {label: D}]}
              notes: n
          - split:
              tag: Early results
              left: {image: a.jpg, label: Left, stat: 75%, headline: faster}
              right: {image: b.jpg, label: Right, stat: Day one, headline: staffed}
              notes: n
    ''')
    out = tmp_path / 'out' / 'deck.pptx'
    findings = render(S.load(p), out)
    assert not [f for f in findings if f.level == 'error']
    assert len(Presentation(str(out)).slides) == 2


def test_catchphrase_across_titles_warns(tmp_path):
    slide = '''
          - content:
              title: Everything lands on one record {n}.
              body: [{{takeaway: {{lead: Line {n}.}}}}]
              notes: n'''
    p = write(tmp_path, 'deck: T\nclient: _Sandbox\nslides:' + ''.join(slide.format(n=i) for i in range(4)))
    assert any(f.code == 'REPEAT' for f in C.check(S.load(p)))


def test_table_row_cell_count_is_checked(tmp_path):
    assert 'TABLE' in _findings(tmp_path, '''
        title: A table whose second row split on a comma.
        body:
          - table:
              head: [Line, How, Amount]
              rows:
                - [Services, "320 hours", "$48,000"]
                - [Seats, first three months, "30 x $45 x 3", "$4,050"]
        notes: n
    ''')


def test_diagram_lays_out_and_writes_drawio(tmp_path):
    import xml.dom.minidom
    from servantium_documents import diagram as D
    p = tmp_path / 'diagram.yaml'
    p.write_text(textwrap.dedent('''
        diagram: Flow
        client: _Sandbox
        lanes: [{key: a, name: Team A}, {key: b, name: Team B, note: Approvers}]
        steps:
          - {id: 1, num: 1, lanes: [a, b], kind: call, title: Kickoff, text: Both teams, target: 1 hr}
          - {id: 2, num: 2, lane: a, title: Draft, target: 2 hrs}
          - {id: 3, num: 3, lane: a, branch: true, kind: optional, title: Peer review}
          - {id: 4, num: 4, lane: b, kind: approval, title: Approve}
          - {id: 5, lane: a, kind: end, title: Done}
        loops: [{from: 4, to: 2, label: Rework}]
        phases: [{label: Day 1, from: 1, to: 2}, {label: Day 2, from: 4, to: 5}]
    '''))
    spec = D.load(p)
    W, H, items = D.layout(spec)
    assert W > 800 and H > 400
    svg = D.to_svg(W, H, items)
    assert svg.startswith('<svg') and 'marker-end' in svg
    xml.dom.minidom.parseString(D.to_drawio(spec, W, H, items))    # well-formed, opens in diagrams.net


def test_diagram_rejects_unknown_lane(tmp_path):
    from servantium_documents import diagram as D
    p = tmp_path / 'd.yaml'
    p.write_text('diagram: X\nclient: _Sandbox\nlanes: [{key: a, name: A}]\nsteps: [{id: 1, lane: z, title: T}]\n')
    with pytest.raises(S.SpecError):
        D.load(p)


def test_people_cards_fit_bios_and_logos(tmp_path):
    """Team bios: round photos, logos sized by area, and the card grows with the bio."""
    from PIL import Image
    Image.new('RGB', (400, 400), 'gray').save(tmp_path / 'a.png')
    Image.new('RGB', (512, 118), 'orange').save(tmp_path / 'wide.png')
    Image.new('RGB', (200, 200), 'blue').save(tmp_path / 'square.png')
    p = write(tmp_path, '''
        deck: People
        client: _Sandbox
        slides:
          - content:
              title: Two people who have done this before.
              body:
                - people:
                    - {photo: a.png, name: Ada Park, role: Founder, lines: [Ran delivery teams, Owned the P&L], logos: [wide.png, square.png]}
                    - {photo: a.png, name: Ben Ito, role: Founder, lines: [Built the platforms]}
    ''')
    findings = render(S.load(p), tmp_path / 'o.pptx', tmp_path)
    assert not [f for f in findings if f.level == 'error']
    pics = [sh for sh in Presentation(tmp_path / 'o.pptx').slides[0].shapes if sh.shape_type == 13]
    wide = max(pics, key=lambda sh: sh.width / sh.height)
    assert wide.height < kit.Inches(0.42)          # a heavy wordmark is not given full height


def test_consulting_pictures_render(tmp_path):
    """Adoption curve, approach matrix, zoom, a team slide and section trackers all render
       inside the slide with no overflow."""
    from PIL import Image, ImageDraw
    cut = Image.new('RGBA', (400, 400), (0, 0, 0, 0))
    ImageDraw.Draw(cut).ellipse((100, 40, 300, 380), fill=(120, 90, 70, 255))   # a cut-out "person"
    cut.save(tmp_path / 'p.png')
    Image.new('RGB', (512, 118), 'orange').save(tmp_path / 'wide.png')
    p = write(tmp_path, '''
        deck: Pictures
        client: _Sandbox
        slides:
          - section: {title: Where teams are, lede: The first part.}
          - content:
              title: AI is crossing into the early majority.
              body:
                - adoption: {here: Most teams are here, note: Rolling it out across the business}
          - team:
              title: We meet every customer where they are.
              points:
                - {icon: apps, lead: Your tools, text: We build inside what you use.}
                - {icon: forum, lead: Your process, text: We learn how you work first.}
              credentials: Implementations, internal tools and SaaS delivery.
              people:
                - {photo: p.png, name: Ada Park, role: Founder, logos: [wide.png]}
                - {photo: p.png, name: Ben Ito, role: Founder}
          - section: {title: How we work}
          - content:
              title: Each use case follows the same five steps.
              body:
                - approach:
                    stages: [Plan, Discover, Build, Test, Launch]
                    rows:
                      - {label: Your team, cells: [Pick it, Join, Review, Use it, Roll out]}
                      - {label: You get, tone: record, cells: [A brief, Interviews, Version one, A test, A workflow]}
          - content:
              title: Discovery starts with how each person works today.
              body:
                - zoom:
                    steps: [Plan, Discover, Build]
                    active: 2
                    title: Discovery workshops
                    rows:
                      - {key: "01", title: "How do you do it today?", text: Walk us through it.}
                      - {key: "02", title: "What have you tried?"}
    ''')
    findings = render(S.load(p), tmp_path / 'o.pptx', tmp_path)
    assert not [f for f in findings if f.level == 'error'], findings
    prs = Presentation(tmp_path / 'o.pptx')
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert any('How we work' in t for t in texts)      # the divider lists every section


def test_box_rows_and_title_tropes_are_flagged(tmp_path):
    findings = C.check(S.load(write(tmp_path, '''
        deck: Boxes
        client: _Sandbox
        slides:
          - content:
              title: Agenda
              body:
                - table: {head: [Time, Topic], rows: [["0:00", Welcome]]}
          - content:
              title: Consistency first, then robustness.
              body:
                - milestones:
                    - {when: May, title: One, text: Done.}
                    - {when: June, title: Two, text: Done.}
          - content:
              title: Then we build it into the tools you like best.
              body:
                - cards:
                    - {title: One}
                    - {title: Two}
    ''')))
    msgs = [(f.code, f.slide, f.message) for f in findings]
    assert any(c == 'VARIETY' and 'row of boxes' in m for c, _, m in msgs)
    assert any(c == 'TROPE' and n == 2 for c, n, _ in msgs)
    assert any(c == 'TROPE' and n == 3 for c, n, _ in msgs)
    assert not any(c == 'TITLE' and n == 1 and 'topic' in m for c, n, m in msgs)


def test_saved_pptx_opens_in_strict_readers(tmp_path):
    """Keynote and Google Slides refused decks that PowerPoint and LibreOffice opened: a
       decimal coordinate ("0.0") from the adoption curve, and a notes master the
       presentation never listed. Every saved deck must pass the package audit."""
    from servantium_documents import package
    p = write(tmp_path, '''
        deck: Strict
        client: _Sandbox
        slides:
          - content:
              title: AI is crossing into the early majority.
              body:
                - adoption: {here: Most teams are here}
              notes: Speaker notes create a notes master, which must be listed.
    ''')
    out = tmp_path / 'o.pptx'
    findings = render(S.load(p), out, tmp_path)
    assert not [f for f in findings if f.code == 'PPTX'], findings
    assert package.audit(out) == []


def test_staircase_questions_and_stepless_flow_render(tmp_path):
    """A pitch that pushes a client toward a bigger vision needed two new pictures: access
       that widens one proven step at a time, and stops where the client does the talking.
       A flow that runs across every stage (checks on everything agents send) has no stage
       bar. All three render inside the slide and pass the strict-reader audit."""
    from servantium_documents import package
    p = write(tmp_path, '''
        deck: Pitch
        client: _Sandbox
        slides:
          - content:
              title: Agents take on more, one step at a time.
              body:
                - staircase:
                    here: 1
                    steps:
                      - {label: Read, title: Read the record, text: Nothing is written back, when: Phase 2}
                      - {label: Log, title: Log what they did, text: "A response sent, a call made", when: Phase 2}
                      - {label: Request, title: Request documents, text: Past performance and certifications, when: Phase 3}
                      - {label: Generate, title: Generate documents, text: Drafts from your templates, when: Phase 3}
                      - {label: Propose, title: Propose plans, text: From your plan template, when: Later}
                - takeaway: {lead: "Each step opens when the last runs clean."}
          - content:
              title: Agent drafts are checked before they become the record.
              body:
                - flow:
                    inputs:
                      - {icon: travel_explore, title: Leads, meta: From the sites they watch, tag: Find}
                      - {icon: checklist, title: Requirements, meta: Filled in from the award, tag: Award}
                    node: {title: Servantium, verb: Checks what agents send, cycle: [Agent sends, Owner approves]}
                    output:
                      label: The record
                      title: Approved and traced
                      rows:
                        - {key: Source, text: Agent and time kept}
                        - {key: Owner, text: A named approver signs off, open: true}
          - content:
              title: Questions for you
              body:
                - questions:
                    note: Your answers shape the first sprint.
                    items:
                      - {q: "What do your agents run on?", why: So we connect the way security allows.}
                      - {q: "Which data must never reach an agent?"}
                      - {q: "Who approves what an agent sends?", why: So every draft has a named owner.}
                      - {q: "Is there a date this has to work by?"}
    ''')
    out = tmp_path / 'o.pptx'
    findings = render(S.load(p), out, tmp_path)
    assert not [f for f in findings if f.level == 'error'], findings
    assert package.audit(out) == []
    texts = ' '.join(sh.text_frame.text for sl in Presentation(out).slides for sh in sl.shapes if sh.has_text_frame)
    assert 'START HERE' in texts and '?' in texts


def test_navigation_titles_may_repeat(tmp_path):
    """A "Questions for you" stop at the end of each section, and a divider named for a
       section, are navigation. They repeat by design and must not read as a catchphrase."""
    p = write(tmp_path, '''
        deck: Nav
        client: _Sandbox
        slides:
          - section: {title: Your agents}
          - content:
              title: Questions for you
              body:
                - questions: {items: [{q: "What do they run on?"}, {q: "Who approves?"}]}
              notes: Stop here.
          - content:
              title: Questions for you
              body:
                - questions: {items: [{q: "Which data never moves?"}, {q: "What runs today?"}]}
              notes: Stop here.
          - content:
              title: Questions for you
              body:
                - questions: {items: [{q: "Is there a hard date?"}, {q: "Which months are off-limits?"}]}
              notes: Stop here.
          - content:
              title: Your agents propose plans for the PMO to approve.
              body:
                - takeaway: {lead: Agents draft and people decide.}
              notes: The agents section.
    ''')
    found = [f for f in C.check(S.load(p)) if f.code in ('REPEAT', 'TITLE')]
    assert not found, found


def test_middleout_and_a_gantt_that_leaves_room_for_its_takeaway(tmp_path):
    """The owner sees the platform in the middle of the sales-to-services stack, built left
       and right at once; every item joins the center with an arrow and long names grow their
       card instead of spilling. A six-row weekly Gantt with a takeaway ran past the footer:
       a Gantt now shrinks its rows to leave room for the emphasis block under it."""
    from servantium_documents import package
    p = write(tmp_path, '''
        deck: Middle
        client: _Sandbox
        slides:
          - content:
              title: We build left to sales and right to services at once.
              body:
                - middleout:
                    center: {title: Live now, items: [Quotes and contracts, Project plans], note: Every new deal starts here.}
                    left:
                      label: Toward sales
                      items:
                        - {icon: handshake, name: Engagements, text: Your pipeline moves in from the old CRM, when: Phase 1}
                        - {icon: travel_explore, name: Leads from government sites and grants, text: Ranked on fit, when: Phase 2}
                    right:
                      label: Toward services
                      items:
                        - {icon: schedule, name: Time capture, text: Hours against plan tasks, when: Phase 1}
                        - {icon: support_agent, name: Support, when: Phase 3}
                - takeaway: {lead: "Both sides move every phase."}
          - content:
              title: Every week runs on a Monday plan and a Friday close.
              body:
                - gantt:
                    periods: [Monday, Tuesday, Wednesday, Thursday, Friday]
                    rows:
                      - {name: Weekly plan, note: PMO, bars: [{from: 1, to: 1, label: Plan}]}
                      - {name: Time capture, note: Everyone, bars: [{from: 1, to: 5, label: Daily}]}
                      - {name: Approvals, note: Contracts, bars: [{from: 1, to: 5, label: As they come}]}
                      - {name: Time and expenses, note: Leads, bars: [{from: 5, to: 5, label: Approve}]}
                      - {name: Status debrief, note: Each project, bars: [{from: 5, to: 5, label: Debrief}]}
                      - {name: Sprint review, note: Both teams, bars: [{from: 5, to: 5, label: Review}]}
                - takeaway: {lead: "What closes on Friday", text: "is what the PMO sees on Monday."}
    ''')
    out = tmp_path / 'o.pptx'
    findings = render(S.load(p), out, tmp_path)
    assert not [f for f in findings if f.level == 'error'], findings
    assert package.audit(out) == []


def test_middleout_center_box_holds_four_items_and_a_note(tmp_path):
    """The advisors deck put four items and a note in the center: the note was drawn over
       the list and the last items fell out of the box. The center now grows to fit."""
    p = write(tmp_path, '''
        deck: Middle
        client: _Sandbox
        slides:
          - content:
              title: Buyers span headless to full UI.
              body:
                - middleout:
                    center: {title: One data model, note: "Every buyer, same record", items: [Quotes, Rate cards, Plans, Projects]}
                    left:
                      label: Headless
                      items:
                        - {name: MCP server, text: Their AI reads the record}
                        - {name: Agents, text: Drafts checked first}
                    right:
                      label: Full UI
                      items:
                        - {name: Grid, text: One screen}
                        - {name: Guided screens, text: For app teams}
    ''')
    out = tmp_path / 'o.pptx'
    render(S.load(p), out, tmp_path)
    slide = Presentation(out).slides[0]
    texts = {sh.text_frame.text: sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text}
    note, last = texts['Every buyer, same record'], texts['Projects']
    assert note.top >= last.top + last.height, 'the note sits below the last item'
    center = [sh for sh in slide.shapes if not sh.has_text_frame or not sh.text_frame.text]
    bottoms = [sh.top + sh.height for sh in center if sh.left <= note.left and sh.left + sh.width >= note.left + note.width
               and sh.top <= last.top]
    assert bottoms and max(bottoms) >= note.top + note.height, 'the note stays inside the center box'


def test_keynote_cannot_shrink_titles(tmp_path):
    """Keynote shrank every title on import: the template's placeholders inherit "shrink text
       on overflow", and the engine sized title boxes on 1.08 x the font size while Playfair
       sets a line at 1.33 x. Placeholders the engine fills now turn shrinking off, name
       their face, and are as tall as their text really sets."""
    from lxml import etree
    from servantium_documents import package
    p = write(tmp_path, '''
        deck: Shrink
        client: _Sandbox
        slides:
          - cover: {title: One System, subtitle: In phases, meta: For discussion}
          - section: {title: Today, lede: Four systems.}
          - content:
              title: Each engagement is spread across four systems.
              lede: One line of context under the title.
              body:
                - takeaway: {lead: The work happens once.}
              notes: Notes.
          - closing: {title: Next steps, asks: [Name owners., Send documents.]}
    ''')
    out = tmp_path / 'o.pptx'
    findings = render(S.load(p), out, tmp_path)
    assert not [f for f in findings if f.level == 'error'], findings
    assert package.audit(out) == []               # the audit flags any shrinkable placeholder
    prs = Presentation(out)
    for slide in prs.slides:
        for ph in slide.placeholders:
            if not ph.has_text_frame or not ph.text_frame.text.strip():
                continue
            bp = ph.text_frame._txBody.find(qn('a:bodyPr'))
            assert bp.find(qn('a:noAutofit')) is not None, ph.name
            assert all(r.font.name for para in ph.text_frame.paragraphs for r in para.runs), ph.name
    title = prs.slides[2].shapes.title
    need = kit.text_box_h(title.text_frame.text, kit.TITLE_W_IN, kit.TITLE_SIZE, font=kit.DISPLAY, line_spacing=1.04)
    assert title.height >= need
    assert title.text_frame.paragraphs[0].runs[0].font.name == kit.DISPLAY


def test_stack_note_sits_below_a_wrapped_tier_label(tmp_path):
    """A two-line tier label ("Five tools for one engagement") had its note drawn on top of the
       second line. The note now starts under however many lines the label took."""
    p = write(tmp_path, '''
        deck: Stack
        client: _Sandbox
        slides:
          - content:
              title: We win where the tech stack is the problem.
              body:
                - stack:
                    tiers:
                      - {label: Five tools for one engagement, note: Like a customer, items: [{name: Zendesk}, {name: Monday}]}
                      - {label: Short, note: Plain, items: [{name: CRM}, {name: Excel}]}
                    foundation: {title: One record, chips: [One, Two]}
    ''')
    out = tmp_path / 'o.pptx'
    render(S.load(p), out, tmp_path)
    shapes = {sh.text_frame.text: sh for sh in Presentation(out).slides[0].shapes if sh.has_text_frame and sh.text_frame.text}
    label, note = shapes['FIVE TOOLS FOR ONE ENGAGEMENT'], shapes['Like a customer']
    assert note.top >= label.top + label.height, 'the note starts below the wrapped label'


def test_full_bleed_slide_fills_the_page_and_carries_alt_text(tmp_path):
    """A cold-open beat is a designed picture with no title: it fills the slide edge to edge,
       and its alt text says what it shows."""
    from PIL import Image
    img = tmp_path / 'beat.png'
    Image.new('RGB', (1600, 900), (2, 62, 37)).save(img)
    p = write(tmp_path, f'''
        deck: Full
        client: _Sandbox
        slides:
          - full: {{image: {img}, alt: "October 2025. Two consultants start Servantium."}}
    ''')
    out = tmp_path / 'o.pptx'
    findings = render(S.load(p), out, tmp_path)
    assert not [f for f in findings if f.level == 'error'], findings
    pic = [sh for sh in Presentation(out).slides[0].shapes if sh.shape_type == 13][0]
    assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, kit.W, kit.H)
    assert 'Two consultants' in pic._element.nvPicPr.cNvPr.get('descr')
    assert not [f for f in C.check(S.load(p)) if f.level == 'error']


def test_process_map_and_pipeline_board_render(tmp_path):
    """The owner wanted a whole business on one slide (numbered steps, what each covers, what
       it replaces, a loop back to the start) and a sales board where every source feeds one
       ranked pipeline. Both render inside the slide and pass the strict-reader audit."""
    from PIL import Image
    from servantium_documents import package
    Image.new('RGBA', (64, 64), (200, 30, 30, 255)).save(tmp_path / 'logo.png')
    p = write(tmp_path, '''
        deck: Map
        client: _Sandbox
        slides:
          - content:
              title: The whole business, start to finish, on one map.
              body:
                - processmap:
                    groups: [{label: Sales, from: 1, to: 2}, {label: Live today, from: 3, to: 3, tone: live}]
                    stages:
                      - {title: Find, text: Sources, build: [Lead sources, Ranking], replaces: [logo.png]}
                      - {title: Pursue, live: [Past work], build: [Go or no-go]}
                      - {title: Deliver, build: [Time capture], later: [DCAA]}
                    loop: Past work feeds the next bid
          - content:
              title: Signals from every source land on one ranked board.
              body:
                - board:
                    sources: [SAM.gov, eBuy, Referrals]
                    columns:
                      - {title: Signal, cards: [{title: County EMS, meta: SAM.gov, score: "82", tag: SDVOSB}]}
                      - {title: Qualified, cards: []}
                      - {title: Awarded, cards: [{title: Installation C, tag: Handoff}]}
    ''')
    out = tmp_path / 'o.pptx'
    findings = render(S.load(p), out, tmp_path)
    assert not [f for f in findings if f.level == 'error'], findings
    assert package.audit(out) == []
