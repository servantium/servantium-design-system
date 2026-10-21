"""Workbook engine tests: the spec is strict, the formulas are real and recompute, the
Gantt bars are conditional formats, and nothing in sheets.py types a colour or a face."""
import datetime as dt
import re
import shutil
import textwrap
import zipfile
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from servantium_documents import kit, sheets
from servantium_documents.spec import SpecError

PKG = Path(__file__).resolve().parents[1]
EXAMPLES = PKG / 'examples'
SRC = PKG / 'src' / 'servantium_documents' / 'sheets.py'
HAVE_SOFFICE = bool(shutil.which('soffice')) or Path('/Applications/LibreOffice.app').exists()
needs_soffice = pytest.mark.skipif(not HAVE_SOFFICE, reason='LibreOffice not installed')


def write(tmp_path, body, name='book.yaml'):
    p = tmp_path / name
    p.write_text(textwrap.dedent(body))
    return p


def copy_example(tmp_path, name):
    dst = tmp_path / name
    shutil.copy(EXAMPLES / name, dst)
    return dst


def spec_of(name):
    return yaml.load((EXAMPLES / name).read_text(), Loader=sheets._StrictLoader)


MINIMAL = """\
    workbook: Test
    client: Acme
    date: 2026-10-20
    sheets:
      - name: Lines
        kind: table
        columns:
          - {key: item, header: Item}
          - {key: qty, header: Qty, format: int, total: sum, input: true}
          - {key: rate, header: Rate, format: currency, input: true}
          - {key: amount, header: Amount, format: currency, formula: "=[@qty]*[@rate]", total: sum}
        rows:
          - {item: A, qty: 2, rate: 10}
          - {item: B, qty: 3, rate: 7.5}
    """


# ── the spec is strict ───────────────────────────────────────────────────────
def test_examples_validate():
    for p in sorted(EXAMPLES.glob('workbook-*.yaml')):
        sheets.load(p)


@pytest.mark.parametrize('bad, expect', [
    ('workbook: T\nclient: A\nsheets: []\n', 'sheets'),
    ('workbook: T\nclient: A\ncolour: red\nsheets: [{name: X, kind: summary}]\n', "'colour'"),
    ('workbook: T\nclient: A\nsheets: [{name: X, kind: chart}]\n', "'chart'"),
    ('workbook: T\nclient: A\nsheets: [{name: X, kind: summary, colour: red}]\n', "'colour'"),
    ('workbook: T\nclient: A\nsheets: [{name: "a/b", kind: summary}]\n', 'name'),
    ('workbook: T\nclient: A\ndate: 20 Oct\nsheets: [{name: X, kind: summary}]\n', 'date'),
])
def test_schema_rejects_bad_specs(tmp_path, bad, expect):
    with pytest.raises(SpecError) as e:
        sheets.load(write(tmp_path, bad))
    assert expect in str(e.value)


def test_schema_rejects_unknown_column_key(tmp_path):
    body = MINIMAL.replace('{key: item, header: Item}', '{key: item, header: Item, colour: red}')
    with pytest.raises(SpecError, match='colour'):
        sheets.load(write(tmp_path, body))


def test_extra_row_keys_are_ignored(tmp_path):
    """Rows may bind straight from richer records: keys that are not columns are dropped."""
    body = MINIMAL.replace('{item: A, qty: 2, rate: 10}', '{item: A, qty: 2, rate: 10, price: 4, owner: Lee}')
    r = sheets.build(write(tmp_path, body), pdf=False)
    vals = [c.value for row in load_workbook(r['xlsx'])['Lines'].iter_rows() for c in row]
    assert 'Lee' not in vals and 4 not in vals


def test_formula_must_name_real_columns(tmp_path):
    body = MINIMAL.replace('=[@qty]*[@rate]', '=[@qty]*[@price]')
    with pytest.raises(SpecError, match=r'\[@price\]'):
        sheets.load(write(tmp_path, body))


def test_comma_split_is_explained(tmp_path):
    body = MINIMAL.replace('{key: item, header: Item}', '{key: item, header: Item, Inc}')
    with pytest.raises(SpecError, match='double quotes'):
        sheets.load(write(tmp_path, body))


def test_dates_stay_strings(tmp_path):
    spec = sheets.load(write(tmp_path, MINIMAL))
    assert spec['date'] == '2026-10-20'


@pytest.mark.parametrize('tasks, expect', [
    ('[{id: a, name: A}]', 'start date or depends_on'),
    ('[{id: a, name: A, start: 2026-11-02}]', 'end or duration_days'),
    ('[{id: a, name: A, start: 2026-11-02, end: 2026-11-01}]', 'before it starts'),
    ('[{id: a, name: A, start: 2026-11-02, duration_days: 2, depends_on: [zz]}]', '"zz" is not a task'),
    ('[{id: a, name: A, depends_on: [b], duration_days: 2}, {id: b, name: B, depends_on: [a], duration_days: 2}]',
     'loop'),
    ('[{id: p, name: P, start: 2026-11-02}, {id: a, name: A, parent: p, start: 2026-11-02, duration_days: 1}]',
     'is a phase'),
    ('[{id: a, name: A, start: 2026-11-02, milestone: true, duration_days: 3}]', 'start only'),
])
def test_gantt_rules(tmp_path, tasks, expect):
    body = f'workbook: P\nclient: A\nsheets:\n  - {{name: Plan, kind: gantt, tasks: {tasks}}}\n'
    with pytest.raises(SpecError, match=re.escape(expect)):
        sheets.load(write(tmp_path, body))


def test_form_items_follow_the_form(tmp_path):
    body = ('workbook: E\nclient: A\nsheets:\n  - {name: E, kind: form, form: expense_report, '
            'items: [{date: 2026-10-01, category: Meals, description: Lunch, amount: 12, qty: 2}]}\n')
    with pytest.raises(SpecError, match='qty not used on a expense_report'):
        sheets.load(write(tmp_path, body))


# ── formulas are real formulas ───────────────────────────────────────────────
def test_formulas_are_written_as_a1_references(tmp_path):
    r = sheets.build(write(tmp_path, MINIMAL), pdf=False)
    ws = load_workbook(r['xlsx'])['Lines']
    cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    amounts = [v for k, v in cells.items() if k.startswith('E') and isinstance(v, str) and v.startswith('=')]
    assert '=C7*D7' in amounts and '=C8*D8' in amounts                      # header on row 6
    assert '=SUM(E7:E8)' in cells.values() and '=SUM(C7:C8)' in cells.values()
    assert not any('[' in str(v) for v in cells.values() if str(v).startswith('=')), 'a template was not resolved'


def test_estimate_formulas_and_cross_sheet_refs(tmp_path):
    r = sheets.build(copy_example(tmp_path, 'workbook-estimate.yaml'), pdf=False)
    wb = load_workbook(r['xlsx'])
    est = wb['Estimate']
    formulas = [c.value for row in est.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
    assert sum(1 for f in formulas if re.fullmatch(r'=D\d+\*E\d+', f)) == 14     # every role line is live
    assert any(f.startswith('=SUM(F') and f.count(':') == 4 for f in formulas)   # grand total over 4 blocks
    summ = [c.value for row in wb['Summary'].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(re.match(r"='?Estimate'?!\$F\$\d+$", v) for v in summ), 'summary must point at the estimate'


def test_no_excel_tables_or_cached_values(tmp_path):
    """Sheets drops ListObjects; the file must carry plain ranges and calc-on-open."""
    r = sheets.build(copy_example(tmp_path, 'workbook-estimate.yaml'), pdf=False)
    with zipfile.ZipFile(r['xlsx']) as z:
        names = z.namelist()
        assert not any(n.startswith('xl/tables/') for n in names)
        wbxml = z.read('xl/workbook.xml').decode()
        core = z.read('docProps/core.xml').decode()
    assert 'fullCalcOnLoad="1"' in wbxml
    assert re.search(rf'<dc:creator[^>]*>{kit.BRAND_NAME}</dc:creator>', core)
    wb = load_workbook(r['xlsx'])
    for ws in wb.worksheets:
        assert not ws.tables
        assert ws.sheet_view.showGridLines is False
        assert ws.oddFooter.left.text == kit.FOOTER_TEXT
    assert wb['Estimate'].freeze_panes and wb['Estimate'].print_title_rows


# ── they recompute (LibreOffice is the independent calculator) ───────────────
def _recalc(tmp_path, xlsx):
    return load_workbook(sheets.recalculate(xlsx, tmp_path / 'recalc'), data_only=True)


def _numbers(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]


@needs_soffice
def test_estimate_totals_recompute(tmp_path):
    spec = spec_of('workbook-estimate.yaml')
    est = next(s for s in spec['sheets'] if s['name'] == 'Estimate')
    rows = [r for g in est['groups'] for r in g['rows']]
    fees = sum(r['hours'] * r['rate'] for r in rows)
    hours = sum(r['hours'] for r in rows)
    r = sheets.build(copy_example(tmp_path, 'workbook-estimate.yaml'), pdf=False)
    wb = _recalc(tmp_path, r['xlsx'])
    assert fees in _numbers(wb['Estimate']) and hours in _numbers(wb['Estimate'])
    summary = _numbers(wb['Summary'])
    assert fees in summary and hours in summary
    assert any(abs(v - fees / hours) < 1e-9 for v in summary)               # blended rate card
    first_group = sum(r['hours'] * r['rate'] for r in est['groups'][0]['rows'])
    assert first_group in summary                                           # [#fees.discover]


@needs_soffice
def test_editing_an_input_changes_the_total(tmp_path):
    """The point of real formulas: change an hour and the total follows."""
    r = sheets.build(write(tmp_path, MINIMAL), pdf=False)
    wb = load_workbook(r['xlsx'])
    ws = wb['Lines']
    ws['C7'] = 10                               # qty for item A: 2 -> 10
    edited = tmp_path / 'edited.xlsx'
    wb.save(edited)
    vals = _numbers(_recalc(tmp_path, edited)['Lines'])
    assert 10 * 10 + 3 * 7.5 in vals and 13 in vals


@needs_soffice
@pytest.mark.parametrize('name, expected', [
    ('workbook-invoice.yaml', lambda s: _invoice_total(s)),
    ('workbook-po.yaml', lambda s: _invoice_total(s)),
    ('workbook-expenses.yaml', lambda s: round(sum(i['amount'] for i in s['items']) - s.get('advance', 0), 2)),
])
def test_form_totals_recompute(tmp_path, name, expected):
    sh = spec_of(name)['sheets'][0]
    r = sheets.build(copy_example(tmp_path, name), pdf=False)
    vals = _numbers(_recalc(tmp_path, r['xlsx']).worksheets[0])
    want = expected(sh)
    assert any(abs(v - want) < 0.005 for v in vals), f'{name}: total {want} not found in {sorted(vals)[-5:]}'


def _invoice_total(sh):
    sub = sum(i['qty'] * i['rate'] for i in sh['items'])
    disc = -sub * sh.get('discount_rate', 0) - sh.get('discount', 0)
    tax = (sub + disc) * sh.get('tax_rate', 0)
    return round(sub + disc + tax + sh.get('shipping', 0), 2)


# ── the Gantt ────────────────────────────────────────────────────────────────
def test_gantt_bars_are_conditional_formats(tmp_path):
    r = sheets.build(copy_example(tmp_path, 'workbook-plan.yaml'), pdf=False)
    ws = load_workbook(r['xlsx'])['Plan']
    rules = [(str(rng.sqref), rule) for rng in ws.conditional_formatting for rule in rng.rules]
    formulas = [rule.formula[0] for _, rule in rules if rule.formula]
    assert len(rules) >= 8
    # bars compare each row's start/end with the date header (absolute row), so they move with dates
    bars = [f for f in formulas if 'TODAY' not in f]
    assert all(re.search(r'\$E\d+', f) and re.search(r'[A-Z]+\$\d+', f) for f in bars)
    assert any(re.search(r'\$E\d+<=[A-Z]+\$\d+\+6', f) for f in formulas)
    assert any('TODAY()' in f for f in formulas)
    for _, rule in rules:
        assert rule.type == 'expression' and rule.dxf.fill is not None    # solid fills only: Sheets keeps them
    cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(v.startswith('=MIN(') for v in cells) and any(v.startswith('=MAX(') for v in cells)
    assert any(sheets.DIAMOND in v and v.startswith('=IF(') for v in cells)
    # every milestone carries a diamond formula in every period, and every timeline cell its rule
    tasks = spec_of('workbook-plan.yaml')['sheets'][0]['tasks']
    milestones = sum(1 for t in tasks if t.get('milestone'))
    periods = len([c for c in ws[8] if c.column >= 11 and c.value])
    assert sum(1 for v in cells if sheets.DIAMOND in v and v.startswith('=')) == milestones * periods
    assert any(v.startswith('=WORKDAY(MAX(') for v in cells)                # dependencies are live


@needs_soffice
def test_gantt_dates_recompute(tmp_path):
    spec = spec_of('workbook-plan.yaml')
    plan = sheets.Plan(spec['sheets'][0]['tasks'])
    r = sheets.build(copy_example(tmp_path, 'workbook-plan.yaml'), pdf=False)
    ws = _recalc(tmp_path, r['xlsx'])['Plan']
    got = {}
    for row in ws.iter_rows(min_row=8):
        wbs = row[1].value
        if isinstance(wbs, str) and wbs in plan.wbs.values():
            got[wbs] = (row[4].value, row[5].value)
    assert len(got) == len(plan.order)
    for tid in plan.order:
        s, e = got[plan.wbs[tid]]
        assert (s.date(), e.date()) == plan.dates[tid], tid


# ── a PDF comes out ──────────────────────────────────────────────────────────
@needs_soffice
def test_build_writes_xlsx_and_pdf(tmp_path):
    r = sheets.build(copy_example(tmp_path, 'workbook-invoice.yaml'))
    assert r['xlsx'].name == 'Acme Logistics Invoice INV-1042.xlsx'
    assert r['xlsx'].parent == tmp_path / 'out'
    import pymupdf
    assert len(pymupdf.open(str(r['pdf']))) == 1                           # a form prints on one page


# ── branding comes from kit, never from literals ─────────────────────────────
def test_sheets_has_no_colour_or_font_literals():
    src = SRC.read_text()
    code = re.sub(r'"""[\s\S]*?"""', '', src)                                # docstrings may describe
    assert not re.search(r"['\"]#?[0-9A-Fa-f]{6}(?:[0-9A-Fa-f]{2})?['\"]", code), 'a colour is typed by hand'
    assert not re.search(r'RGBColor\(|rgb=', code)
    for face in ('Playfair', 'Source Sans', 'Calibri', 'Arial', 'Helvetica'):
        assert face not in code, f'font {face} is typed by hand'


# ── references, formats and layout added after the first demo round ──────────
REFS = """\
    workbook: Refs
    client: Acme
    date: 2026-10-20
    sheets:
      - name: Summary
        kind: summary
        metrics:
          - {label: Peak month, value: "=MAX([Monthly!spend])", format: currency}
          - {label: Expense cap, value: 500, format: currency}
          - {label: Average, key: avg, value: "=AVERAGE([Monthly!spend])", format: currency}
        table:
          columns:
            - {key: item, header: Item}
            - {key: amt, header: Amount, format: currency}
            - {key: over, header: Over cap, format: currency, formula: "=MAX(0,[@amt]-[#expense_cap])"}
            - {key: days, header: Days, format: int, formula: "=ROUNDUP([@amt]/[#avg],0)+INT(MOD(7,3))"}
          rows: [{item: a, amt: 700}]
      - name: Monthly
        kind: table
        columns:
          - {key: month, header: Month}
          - {key: spend, header: Spend, format: currency, total: sum}
          - {key: paid, header: Paid, format: yesno, total: count}
          - {key: ok, header: OK, format: check}
          - {key: note, header: Note}
        rows:
          - {month: Jan, spend: 100, paid: true, ok: true, note: true}
          - {month: Feb, spend: 300, paid: false, ok: false, note: false}
    """


def test_same_sheet_metric_and_cross_sheet_range_refs(tmp_path):
    r = sheets.build(write(tmp_path, REFS), pdf=False)
    wb = load_workbook(r['xlsx'])
    f = [c.value for row in wb['Summary'].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(re.fullmatch(r"=MAX\('?Monthly'?!\$C\$\d+:\$C\$\d+\)", v) for v in f), f
    assert any(re.fullmatch(r'=MAX\(0,C\d+-\$C\$\d+\)', v) for v in f), f
    assert any(v.startswith('=ROUNDUP(') and 'INT(MOD(' in v for v in f)


@needs_soffice
def test_metric_and_range_refs_recompute(tmp_path):
    r = sheets.build(write(tmp_path, REFS), pdf=False)
    vals = _numbers(_recalc(tmp_path, r['xlsx'])['Summary'])
    assert 300 in vals and 200 in vals and 200.0 in vals              # peak, average, over cap
    assert 5 in vals                                                  # ROUNDUP(700/200)+INT(MOD(7,3))


def test_booleans_read_yes_no_and_check(tmp_path):
    r = sheets.build(write(tmp_path, REFS), pdf=False)
    ws = load_workbook(r['xlsx'])['Monthly']
    vals = [c.value for row in ws.iter_rows() for c in row]
    assert vals.count('Yes') == 2 and vals.count('No') == 2            # yesno column and a text column
    assert sheets.CHECK in vals and True not in vals and False not in vals
    assert any(isinstance(v, str) and v.startswith('=COUNTIF(') and '"Yes"' in v for v in vals)


@pytest.mark.parametrize('fn', ['XLOOKUP', 'IFS', 'TEXTJOIN', 'MAXIFS'])
def test_functions_that_break_somewhere_are_refused(tmp_path, fn):
    body = MINIMAL.replace('=[@qty]*[@rate]', f'={fn}([@qty],[@rate])')
    with pytest.raises(SpecError, match=fn):
        sheets.build(write(tmp_path, body), pdf=False)


def test_safe_functions_pass_through():
    for f in ('=INT(A1)', '=MOD(A1,2)', '=SUMIF(A:A,"x",B:B)', '=SUMIFS(B:B,A:A,"x")', '=COUNTIF(A:A,1)',
              '=ROUNDUP(A1,0)', '=ROUNDDOWN(A1,0)', '=CEILING(A1,5)', '=FLOOR(A1,5)', '=IFERROR(A1/B1,0)'):
        assert sheets.unsafe_functions(f) == [], f
    assert sheets.unsafe_functions('=IF(A1="IFS(",1,2)') == []       # text inside quotes is not a call


def test_group_subtotals_without_a_grand_total(tmp_path):
    body = MINIMAL.replace("        rows:\n          - {item: A, qty: 2, rate: 10}\n          - {item: B, qty: 3, rate: 7.5}\n",
                           "        total_row: false\n        groups:\n"
                           "          - {name: One, rows: [{item: A, qty: 2, rate: 10}]}\n"
                           "          - {name: Two, rows: [{item: B, qty: 3, rate: 7.5}]}\n")
    assert 'total_row' in body
    r = sheets.build(write(tmp_path, body), pdf=False)
    vals = [c.value for row in load_workbook(r['xlsx'])['Lines'].iter_rows() for c in row]
    assert vals.count('Subtotal') == 2 and 'Total' not in vals


def test_form_skips_subtotal_when_it_is_the_total(tmp_path):
    body = ('workbook: E\nclient: A\nsheets:\n  - {name: E, kind: form, form: expense_report, '
            'items: [{date: 2026-10-01, category: Meals, description: Lunch, amount: 12}]}\n')
    r = sheets.build(write(tmp_path, body), pdf=False)
    vals = [c.value for row in load_workbook(r['xlsx'])['E'].iter_rows() for c in row]
    assert 'Subtotal' not in vals and 'Amount due' in vals


def test_party_blocks_stay_inside_the_print_area():
    widths = [48, 8, 8, 13, 15]                                        # an invoice's columns
    rows = sheets._party_rows(3, widths)
    for row in rows:
        for _, (a, b) in row:
            assert 0 <= a <= b < len(widths) and sum(widths[a:b + 1]) >= 26
    assert sum(len(r) for r in rows) == 3 and len(rows) == 2           # two a row, third below


def test_long_text_wraps_and_rows_grow(tmp_path):
    body = ('workbook: P\nclient: A\nsheets:\n  - name: Plan\n    kind: gantt\n'
            '    subtitle: ' + 'A long subtitle that goes on and on about the plan. ' * 6 + '\n'
            '    tasks:\n'
            '      - {id: a, name: "' + 'Very long task name ' * 4 + '", owner: "Nadia Rahman, PhD, Principal Engineer", '
            'start: 2026-11-02, duration_days: 3}\n')
    r = sheets.build(write(tmp_path, body), pdf=False)
    ws = load_workbook(r['xlsx'])['Plan']
    task_row = next(c.row for row in ws.iter_rows() for c in row if c.value == 'Nadia Rahman, PhD, Principal Engineer')
    assert ws.row_dimensions[task_row].height > 20
    sub_row = next(c.row for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('A long'))
    assert ws.row_dimensions[sub_row].height > 17
    assert ws.cell(row=sub_row, column=2).alignment.wrap_text
