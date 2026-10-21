"""Composed slide blocks.

kit holds primitives (rect, text, measured type, icons). This holds the layouts
that kept getting rewritten by hand in every build script: card rows, hero shots,
timelines, hub-and-spoke, tables, stat rows, strips.

Every block returns the y coordinate below itself, so slides compose top-down without
anyone hard-coding vertical positions. That is deliberate: nearly every layout bug in
the first decks came from a fixed offset that was right for one line of text and wrong
for two.
"""
from pathlib import Path
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

from .kit import (rect, text, title, picture, icon, block_h, text_w_in, est_lines,
                      GREEN, GREEN_600, GREEN_300, FOREST, FOREST_900, MINT, MINT_TINT,
                      INK, INK_800, INK_MUTED, CLOUD, MIST, WHITE, CORAL, SLATE, AMBER, TEAL,
                      DISPLAY, UI, W, H, MARGIN, BODY_W, FOOTER_Y, ASSETS, LOGO_ON_DARK)

from .kit import ON_FOREST as MIST_INK   # text on forest, a document token
from .kit import ON_FOREST_MUTED
from .kit import BRAND_SHORT, ACTIVE as _ACTIVE, SUCCESS

# Spec values that mean "the house": draw the active brand's own card or logo.
HOUSE_KEYS = {'house', 'servantium', _ACTIVE.id, BRAND_SHORT.lower()}
SAFE_BOTTOM = FOOTER_Y - Inches(0.16)     # nothing should cross this


def _cols(n, gap):
    return (BODY_W - gap * (n - 1)) // n


# ── strips and bands ─────────────────────────────────────────────────────────
def strip(s, y, runs, *, tone=GREEN, h=Inches(0.62)):
    """One-line dark strip. If the line needs to wrap, cut words rather than raise h."""
    rect(s, MARGIN, y, BODY_W, h, fill=FOREST)
    rect(s, MARGIN, y, Inches(0.05), h, fill=tone)
    text(s, MARGIN + Inches(0.34), y + (h - Inches(0.24)) // 2, BODY_W - Inches(0.7),
         Inches(0.26), [runs])
    return y + h


def band(s, y, heading, runs, *, tone=GREEN):
    """Dark band with a heading. Height measured from the body text."""
    inner = BODY_W - Inches(0.7)
    joined = ''.join(r[0] for r in runs)
    bh = block_h(joined, inner / 914400, 12, spacing=1.3)
    h = Inches(0.5) + bh + Inches(0.2)
    rect(s, MARGIN, y, BODY_W, h, fill=FOREST)
    rect(s, MARGIN, y, Inches(0.05), h, fill=tone)
    text(s, MARGIN + Inches(0.34), y + Inches(0.16), Inches(7.4), Inches(0.26),
         [(heading, dict(size=13, bold=True, color=WHITE))])
    text(s, MARGIN + Inches(0.34), y + Inches(0.5), inner, bh, [runs], line_spacing=1.3)
    return y + h


def callout(s, y, label, runs, *, tone=AMBER, icon_name='help'):
    h = Inches(0.72)
    rect(s, MARGIN, y, BODY_W, h, fill=FOREST)
    rect(s, MARGIN, y, Inches(0.05), h, fill=tone)
    icon(s, icon_name, MARGIN + Inches(0.32), y + Inches(0.18), Inches(0.36), color=tone)
    text(s, MARGIN + Inches(0.86), y + Inches(0.24), BODY_W - Inches(1.2), Inches(0.26),
         [[(label + ' ', dict(size=13, bold=True, color=tone))] + list(runs)])
    return y + h


# ── card rows ────────────────────────────────────────────────────────────────
def icon_cards(s, y, items, *, h=Inches(3.02), gap=Inches(0.32), fill=MIST, line=CLOUD):
    """items: (icon, kicker, title, [line, ...]). The workhorse.

    `h` is a MINIMUM. The card grows if the title wraps or the bullets are long,
    because a fixed height silently clipped the last bullet off the bottom.
    """
    n = len(items)
    cw = _cols(n, gap)
    inner = cw - Inches(0.6)
    bullet_w = (cw - Inches(0.78)) / 914400

    needed = h
    for _, _, name, lines in items:
        th = block_h(name, inner / 914400, 21, spacing=1.1, font=DISPLAY)
        tall = Inches(1.02) + th + Inches(0.22)
        for ln in lines:
            tall += block_h(ln, bullet_w, 11.5) + Inches(0.14)
        needed = max(needed, tall + Inches(0.2))
    h = needed

    for i, (ic, kicker, name, lines) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        rect(s, x, y, cw, h, fill=fill, line=line)
        rect(s, x, y, cw, Inches(0.05), fill=GREEN)
        if ic:
            icon(s, ic, x + Inches(0.3), y + Inches(0.34), Inches(0.46))
        if kicker:
            text(s, x + cw - Inches(1.62), y + Inches(0.4), Inches(1.4), Inches(0.24),
                 [(kicker, dict(font=DISPLAY, size=14, color=CLOUD))], align=PP_ALIGN.RIGHT)
        th = block_h(name, inner / 914400, 21, spacing=1.1, font=DISPLAY)
        text(s, x + Inches(0.3), y + Inches(1.02), inner, th,
             [(name, dict(font=DISPLAY, size=21, color=FOREST))], line_spacing=1.1)
        yy = y + Inches(1.02) + th + Inches(0.22)
        for ln in lines:
            rect(s, x + Inches(0.3), yy + Inches(0.07), Inches(0.07), Inches(0.07),
                 fill=GREEN, shape=MSO_SHAPE.OVAL)
            lh = block_h(ln, bullet_w, 11.5)
            text(s, x + Inches(0.48), yy, cw - Inches(0.78), lh,
                 [(ln, dict(size=11.5, color=INK_800))], line_spacing=1.28)
            yy += lh + Inches(0.14)
    return y + h


def tiles(s, y, items, *, cols=3, h=Inches(1.38), gap=Inches(0.28)):
    """items: (icon, label). Icon left, label right. For checklists and inventories."""
    tw = _cols(cols, gap)
    rows = (len(items) + cols - 1) // cols
    for i, (ic, label) in enumerate(items):
        c, r = i % cols, i // cols
        x = MARGIN + c * (tw + gap)
        ty = y + r * (h + Inches(0.26))
        rect(s, x, ty, tw, h, fill=MINT_TINT, line=MINT)
        icon(s, ic, x + Inches(0.32), ty + (h - Inches(0.56)) // 2, Inches(0.56))
        text(s, x + Inches(1.06), ty + (h - Inches(0.3)) // 2, tw - Inches(1.3), Inches(0.3),
             [(label, dict(size=15, bold=True, color=FOREST))])
    return y + rows * h + (rows - 1) * Inches(0.26)


def stat_row(s, y, stats, *, h=Inches(1.62), dark=False, gap=Inches(0.34)):
    """items: (big, mid, small). Numbers first: the reason to keep reading."""
    cw = _cols(len(stats), gap)
    for i, (big, mid, small) in enumerate(stats):
        x = MARGIN + i * (cw + gap)
        rect(s, x, y, cw, h, fill=FOREST_900 if dark else MIST,
             line=GREEN_600 if dark else CLOUD)
        rect(s, x, y, Inches(0.05), h, fill=GREEN)
        text(s, x + Inches(0.3), y + Inches(0.2), cw - Inches(0.55), Inches(0.42),
             [(big, dict(size=27, bold=True, color=WHITE if dark else FOREST))])   # body face: Playfair's old-style figures read 0 as o
        mh = block_h(mid, (cw - Inches(0.55)) / 914400, 12.5)
        text(s, x + Inches(0.3), y + Inches(0.68), cw - Inches(0.55), mh,
             [(mid, dict(size=12.5, color=MIST_INK if dark else INK_800))], line_spacing=1.25)
        if small:
            text(s, x + Inches(0.3), y + h - Inches(0.4), cw - Inches(0.55), Inches(0.24),
                 [(small, dict(size=10, color=GREEN_300 if dark else INK_MUTED))])
    return y + h


# ── sequences ────────────────────────────────────────────────────────────────
def timeline(s, y, stages, *, h=Inches(3.02), gap=Inches(0.3)):
    """items: (icon, horizon, priority_label, priority, moves_label, moves).

    Like icon_cards, `h` is a minimum and the body text is measured. The fixed
    offsets this used to carry let a two-line priority run into the OUTPUT label.
    """
    cw = _cols(len(stages), gap)
    inner = cw - Inches(0.6)
    iw = inner / 914400

    needed = h
    for _, _, _, priority, _, moves in stages:
        tall = (Inches(1.46) + Inches(0.24) + block_h(priority, iw, 12.5, spacing=1.26)
                + Inches(0.18) + Inches(0.22) + block_h(moves, iw, 11.5) + Inches(0.24))
        needed = max(needed, tall)
    h = needed

    ph_max = max(block_h(st[3], iw, 12.5, spacing=1.26) for st in stages)   # second labels line up
    for i, (ic, horizon, plabel, priority, mlabel, moves) in enumerate(stages):
        x = MARGIN + i * (cw + gap)
        live = i == 0
        rect(s, x, y, cw, h, fill=MIST if live else WHITE, line=CLOUD)
        rect(s, x, y, cw, Inches(0.05), fill=GREEN if live else GREEN_300)
        icon(s, ic, x + Inches(0.3), y + Inches(0.34), Inches(0.5))
        text(s, x + Inches(0.3), y + Inches(1.0), inner, Inches(0.34),
             [(horizon, dict(font=DISPLAY, size=21, color=FOREST))])
        text(s, x + Inches(0.3), y + Inches(1.46), inner, Inches(0.22),
             [(plabel.upper(), dict(size=9, bold=True, color=GREEN_600, spacing=1.4))])
        ph = block_h(priority, iw, 12.5, spacing=1.26)
        text(s, x + Inches(0.3), y + Inches(1.7), inner, ph,
             [(priority, dict(size=12.5, color=INK_800))], line_spacing=1.26)
        my = y + Inches(1.7) + ph_max + Inches(0.18)
        text(s, x + Inches(0.3), my, inner, Inches(0.22),
             [(mlabel.upper(), dict(size=9, bold=True, color=GREEN_600, spacing=1.4))])
        text(s, x + Inches(0.3), my + Inches(0.22), inner, block_h(moves, iw, 11.5),
             [(moves, dict(size=11.5, color=INK_MUTED))], line_spacing=1.28)
    return y + h


def steps_row(s, y, steps, *, numbered=True):
    """Short numbered steps across the slide. For process, not for detail."""
    gap = Inches(0.3)
    cw = _cols(len(steps), gap)
    tallest = y
    for i, (lead, body_) in enumerate(steps):
        x = MARGIN + i * (cw + gap)
        if numbered:
            rect(s, x, y + Inches(0.02), Inches(0.3), Inches(0.3), fill=MINT,
                 shape=MSO_SHAPE.OVAL)
            text(s, x, y + Inches(0.07), Inches(0.3), Inches(0.22),
                 [(str(i + 1), dict(size=12, bold=True, color=GREEN_600))], align=PP_ALIGN.CENTER)
        tx = x + (Inches(0.44) if numbered else Inches(0))
        text(s, tx, y, cw - (tx - x), Inches(0.28),
             [(lead, dict(size=13.5, bold=True, color=FOREST))])
        bh = block_h(body_, (cw - (tx - x)) / 914400, 11.5)
        text(s, tx, y + Inches(0.32), cw - (tx - x), bh,
             [(body_, dict(size=11.5, color=INK_MUTED))], line_spacing=1.3)
        tallest = max(tallest, y + Inches(0.32) + bh)
    return tallest


def hub_spoke(s, y, left, right, *, lockup=None,
              left_label=None, right_label=None, card_h=Inches(0.58), gap=Inches(0.1),
              top=None, bottom=None, caption=None):
    """Two labelled columns either side of the Servantium mark. items: (name, icon).
       top / bottom: optional (title, text, icon) boxes that feed into the hub from above and
       come out of it below, each joined by an arrow. Every spoke lands on the hub's edge."""
    col_w = Inches(3.46)
    rows = max(len(left), len(right))
    band_h = card_h * rows + gap * (rows - 1)
    lx, rx = MARGIN, MARGIN + BODY_W - col_w
    hub_w, hub_h = Inches(3.0), Inches(1.12)
    hub_x = MARGIN + (BODY_W - hub_w) // 2
    box_w, arrow_gap = Inches(6.2), Inches(0.42)

    def feed_h(item):
        t, txt, ic = item
        inner = box_w - (Inches(0.72) if ic else Inches(0.22)) - Inches(0.16)
        return max(Inches(0.66), Inches(0.36) + (block_h(txt, inner / 914400, 10.5) if txt else 0) + Inches(0.12))

    def feed(bx_y, item, dark):
        t, txt, ic = item
        box_h = feed_h(item)
        b = rect(s, MARGIN + (BODY_W - box_w) // 2, bx_y, box_w, box_h, fill=MINT_TINT if not dark else WHITE,
                 line=GREEN_300, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _round(b, Inches(0.1))
        x0 = MARGIN + (BODY_W - box_w) // 2
        if ic:
            icon(s, ic, x0 + Inches(0.22), bx_y + (box_h - Inches(0.34)) // 2, Inches(0.34))
        tx = x0 + (Inches(0.72) if ic else Inches(0.22))
        text(s, tx, bx_y + Inches(0.1), box_w - (tx - x0) - Inches(0.16), Inches(0.24),
             [(t, dict(size=12.5, bold=True, color=FOREST))])
        if txt:
            iw = box_w - (tx - x0) - Inches(0.16)
            text(s, tx, bx_y + Inches(0.36), iw, block_h(txt, iw / 914400, 10.5),
                 [(txt, dict(size=10.5, color=INK_MUTED))], line_spacing=1.15)

    yy = y
    if top:
        feed(yy, top, False)
        yy += feed_h(top) + arrow_gap
    band_top = yy
    hub_y = band_top + (band_h - hub_h) // 2

    def column(x, label, items, right_align):
        if label:
            text(s, x, band_top - Inches(0.34), col_w, Inches(0.24),
                 [(label.upper(), dict(size=10, bold=True, color=GREEN_600, spacing=1.5))],
                 align=PP_ALIGN.RIGHT if right_align else PP_ALIGN.LEFT)
        for j, (nm, ic) in enumerate(items):
            cy = band_top + j * (card_h + gap)
            rect(s, x, cy, col_w, card_h, fill=WHITE, line=CLOUD,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            if ic:     # an icon name, or a logo file (fitted into the same square)
                _mark(s, ic, x + Inches(0.22), cy + (card_h - Inches(0.36)) // 2, Inches(0.36))
            text(s, x + Inches(0.7), cy + (card_h - Inches(0.24)) // 2, col_w - Inches(0.94),
                 Inches(0.24), [(nm, dict(size=13, bold=True, color=INK_800))])

    # spokes first, so the cards and the hub sit on top of their ends
    def land(cy):          # where a spoke meets the hub's side: its own height, kept on the box
        return min(max(cy, hub_y + Inches(0.2)), hub_y + hub_h - Inches(0.2))
    for j in range(len(left)):
        cy = band_top + j * (card_h + gap) + card_h // 2
        _line(s, lx + col_w, cy, hub_x, land(cy), GREEN_300, w=Pt(1.5), curve=True)
    for j in range(len(right)):
        cy = band_top + j * (card_h + gap) + card_h // 2
        _line(s, rx, cy, hub_x + hub_w, land(cy), GREEN_300, w=Pt(1.5), curve=True)
    column(lx, left_label, left, False)
    column(rx, right_label, right, True)
    if top:
        _arrow(_line(s, MARGIN + BODY_W // 2, band_top - arrow_gap, MARGIN + BODY_W // 2, hub_y, GREEN_600, w=Pt(1.5)))
    _round(rect(s, hub_x, hub_y, hub_w, hub_h, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.14))
    if caption:
        # the hub names what it is (e.g. the use case factory) under the mark
        s.shapes.add_picture(str(lockup or LOGO_ON_DARK), hub_x + Inches(0.47), hub_y + Inches(0.24),
                             width=Inches(2.06))
        text(s, hub_x, hub_y + Inches(0.76), hub_w, Inches(0.22),
             [(caption.upper(), dict(size=9.5, bold=True, color=GREEN_300, spacing=1.5))], align=PP_ALIGN.CENTER)
    else:
        s.shapes.add_picture(str(lockup or LOGO_ON_DARK), hub_x + Inches(0.47), hub_y + Inches(0.4),
                             width=Inches(2.06))
    end = band_top + band_h
    if bottom:
        by = end + arrow_gap
        _arrow(_line(s, MARGIN + BODY_W // 2, hub_y + hub_h, MARGIN + BODY_W // 2, by, GREEN_600, w=Pt(1.5)))
        feed(by, bottom, True)
        end = by + feed_h(bottom)
    return end


# ── data ─────────────────────────────────────────────────────────────────────
# Pills are status marks: green means done/on track in every brand, so it is the status green.
_PILL = {'green': (SUCCESS, WHITE), 'forest': (FOREST, WHITE), 'mint': (MINT, FOREST), 'amber': (AMBER, INK),
         'coral': (CORAL, WHITE), 'slate': (SLATE, WHITE), 'cloud': (CLOUD, INK_800)}


def table(s, y, head, rows, widths, *, row_h=Inches(0.56), zebra=True, pills=None):
    """pills: {column index: {cell value: tone}} draws those cells as colored status pills
       (priority, wave, status) instead of plain text."""
    # Long tables tighten their rows to fit above the footer instead of overflowing.
    row_h = max(Inches(0.4), min(row_h, (SAFE_BOTTOM - y - Inches(0.42)) // max(len(rows), 1)))
    pills = pills or {}
    rect(s, MARGIN, y, BODY_W, Inches(0.42), fill=FOREST)
    cx = MARGIN
    for lbl, cw in zip(head, widths):
        text(s, cx + Inches(0.22), y + Inches(0.12), cw - Inches(0.3), Inches(0.2),
             [(lbl.upper(), dict(size=9.5, bold=True, color=GREEN_300, spacing=1.3))])
        cx += cw
    y += Inches(0.42)
    for r, row in enumerate(rows):
        if zebra and r % 2:
            rect(s, MARGIN, y, BODY_W, row_h, fill=MIST)
        rect(s, MARGIN, y, BODY_W, Emu(9525), fill=CLOUD)
        cx = MARGIN
        for c, (val, cw) in enumerate(zip(row, widths)):
            first = c == 0
            tone = pills.get(c, {}).get(val)
            if tone and val:
                bg, fg = _PILL.get(tone, _PILL['cloud'])
                pw_ = Inches(text_w_in(val, 10, bold=True) + 0.34)
                pill = rect(s, cx + Inches(0.2), y + (row_h - Inches(0.28)) // 2, pw_, Inches(0.28), fill=bg,
                            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                _round(pill, Inches(0.14))
                text(s, cx + Inches(0.2), y + (row_h - Inches(0.28)) // 2 + Inches(0.05), pw_, Inches(0.2),
                     [(val, dict(size=10, bold=True, color=fg))], align=PP_ALIGN.CENTER)
            else:
                text(s, cx + Inches(0.22), y + (row_h - Inches(0.2)) // 2, cw - Inches(0.3), Inches(0.26),
                     [(val, dict(size=12, bold=first, color=FOREST if first else INK_800))])
            cx += cw
        y += row_h
    return y


def compare(s, y, left, right, *, h=Inches(3.5)):
    """Two panels, ours on the right. (label, [line, ...], icon). Rows are ruled and set
       large so the panels use their height instead of floating three small lines."""
    gap = Inches(0.4)
    cw = (BODY_W - gap) // 2
    for i, (label, lines, ic, ours) in enumerate([(*left, False), (*right, True)]):
        x = MARGIN + i * (cw + gap)
        rect(s, x, y, cw, h, fill=MINT_TINT if ours else MIST, line=MINT if ours else CLOUD)
        rect(s, x, y, cw, Inches(0.06), fill=GREEN if ours else CLOUD)
        badge = Inches(0.62)
        rect(s, x + Inches(0.36), y + Inches(0.36), badge, badge, fill=FOREST if ours else CLOUD,
             shape=MSO_SHAPE.OVAL)
        icon(s, ic or ('check' if ours else 'x'), x + Inches(0.36) + Inches(0.14), y + Inches(0.5),
             Inches(0.34), color=GREEN_300 if ours else INK_MUTED)
        text(s, x + Inches(1.18), y + Inches(0.46), cw - Inches(1.5), Inches(0.4),
             [(label, dict(font=DISPLAY, size=24, color=FOREST if ours else INK_MUTED))])
        n = max(len(lines), 1)
        top = y + Inches(1.3)
        row = (y + h - Inches(0.25) - top) // n
        for j, ln in enumerate(lines):
            ry = top + j * row
            rect(s, x + Inches(0.36), ry, cw - Inches(0.72), Emu(9525), fill=MINT if ours else CLOUD)
            icon(s, 'check' if ours else 'x', x + Inches(0.36), ry + (row - Inches(0.26)) // 2,
                 Inches(0.26), color=GREEN_600 if ours else INK_MUTED)
            text(s, x + Inches(0.8), ry + (row - Inches(0.3)) // 2, cw - Inches(1.2), Inches(0.3),
                 [(ln, dict(size=15, color=INK if ours else INK_MUTED))])
    return y + h


def pull_quote(s, y, quote, attribution, *, h=Inches(2.4)):
    rect(s, MARGIN, y, BODY_W, h, fill=FOREST)
    rect(s, MARGIN, y, Inches(0.05), h, fill=GREEN)
    icon(s, 'quote', MARGIN + Inches(0.5), y + Inches(0.4), Inches(0.42), color=GREEN_300)
    qh = block_h(quote, (BODY_W - Inches(2.0)) / 914400, 24, spacing=1.2, font=DISPLAY)
    text(s, MARGIN + Inches(1.2), y + Inches(0.44), BODY_W - Inches(2.0), qh,
         [(quote, dict(font=DISPLAY, size=24, color=WHITE))], line_spacing=1.24)
    text(s, MARGIN + Inches(1.2), y + Inches(0.5) + qh + Inches(0.18), BODY_W - Inches(2.0),
         Inches(0.24), [(attribution, dict(size=12, color=GREEN_300))])
    return y + h


# ── screenshots ──────────────────────────────────────────────────────────────
def hero_shot(s, y, img, caption, steps, *, folder='crop'):
    """Big shot right, numbered steps left. Fits by HEIGHT so tall crops behave."""
    top = y + Inches(0.1)
    h_max = SAFE_BOTTOM - top - Inches(0.4)
    pic = picture(s, Path(img), MARGIN, top, h=h_max)
    pic.left = MARGIN + BODY_W - pic.width
    lw = pic.left - MARGIN - Inches(0.42)
    text(s, pic.left, top + pic.height + Inches(0.12), pic.width, Inches(0.28),
         [(caption, dict(size=10.5, color=INK_MUTED))])
    yy = y + Inches(0.16)
    for i, (lead, body_) in enumerate(steps, 1):
        rect(s, MARGIN, yy + Inches(0.02), Inches(0.3), Inches(0.3), fill=MINT,
             shape=MSO_SHAPE.OVAL)
        text(s, MARGIN, yy + Inches(0.07), Inches(0.3), Inches(0.22),
             [(str(i), dict(size=12, bold=True, color=GREEN_600))], align=PP_ALIGN.CENTER)
        text(s, MARGIN + Inches(0.44), yy, lw - Inches(0.44), Inches(0.28),
             [(lead, dict(size=13.5, bold=True, color=FOREST))])
        bh = block_h(body_, (lw - Inches(0.44)) / 914400, 11.5)
        text(s, MARGIN + Inches(0.44), yy + Inches(0.32), lw - Inches(0.44), bh,
             [(body_, dict(size=11.5, color=INK_MUTED))], line_spacing=1.3)
        yy += Inches(0.32) + bh + Inches(0.3)
    return max(yy, top + pic.height + Inches(0.4))


def shot_row(s, y, items, *, folder='crop', cap_size=11.5):
    """Two or three shots, top aligned, captions on a shared baseline."""
    gap = Inches(0.42) if len(items) == 2 else Inches(0.3)
    cw = _cols(len(items), gap)
    top = y + Inches(0.08)
    h_max = SAFE_BOTTOM - top - Inches(0.86)
    placed = []
    for i, (img, lead, cap) in enumerate(items):
        pic = picture(s, Path(img), MARGIN + i * (cw + gap), top, h=h_max)
        pic.left = MARGIN + i * (cw + gap) + (cw - pic.width) // 2
        placed.append(pic)
    base = top + max(p.height for p in placed) + Inches(0.22)
    for i, (img, lead, cap) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        text(s, x, base, cw, Inches(0.28), [(lead, dict(size=14, bold=True, color=FOREST))])
        text(s, x, base + Inches(0.32), cw, Inches(0.5),
             [(cap, dict(size=cap_size, color=INK_MUTED))], line_spacing=1.3)
    return base + Inches(0.9)


# ── blocks first built for real decks, promoted here ─────────────────────────
# Each came out of a real client deck (a proposal and a status report) and was worth keeping.

def chevron_stages(s, y, stages, *, h=Inches(0.74), overlap=Inches(0.22)):
    """Interlocking stage chevrons with a card under each. For one object examined
       at successive stages. stages: (stage, question, supporting line)."""
    n = len(stages)
    cw = (BODY_W - overlap * (n - 1)) // n
    step = cw + overlap
    inner = cw - Inches(0.56)
    iw = inner / 914400
    card_h = Inches(0)
    for _, q, line in stages:
        tall = (Inches(0.34) + block_h(q, iw, 17, spacing=1.14, font=DISPLAY)
                + Inches(0.26) + block_h(line, iw, 11.5) + Inches(0.36))
        card_h = max(card_h, tall)
    cy = y + h + Inches(0.26)
    for i, (stage, q, line) in enumerate(stages):
        x = MARGIN + i * step
        chw = cw + overlap if i < n - 1 else cw
        shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        ch = s.shapes.add_shape(shape, x, y, chw, h)
        ch.fill.solid(); ch.fill.fore_color.rgb = FOREST
        ch.line.fill.background(); ch.shadow.inherit = False
        ch.text_frame.clear()
        tx = x if i == 0 else x + overlap
        text(s, tx, y + (h - Inches(0.26)) // 2, cw - overlap, Inches(0.26),
             [(stage.upper(), dict(size=11, bold=True, color=GREEN_300, spacing=1.5))],
             align=PP_ALIGN.CENTER)
        rect(s, x, cy, cw, card_h, fill=MIST if i % 2 == 0 else MINT_TINT, line=CLOUD)
        qh = block_h(q, iw, 17, spacing=1.14, font=DISPLAY)
        text(s, x + Inches(0.28), cy + Inches(0.34), inner, qh,
             [(q, dict(font=DISPLAY, size=17, color=FOREST))], line_spacing=1.14)
        text(s, x + Inches(0.28), cy + Inches(0.34) + qh + Inches(0.26), inner,
             block_h(line, iw, 11.5), [(line, dict(size=11.5, color=INK_800))], line_spacing=1.28)
    return cy + card_h


def two_panel(s, y, left, right):
    """Two labelled panels, height measured from the longer list.
       Each: (icon, name, tag, tone, [line, ...]); tone is GREEN or AMBER."""
    gap = Inches(0.4)
    pw = (BODY_W - gap) // 2
    inner = pw - Inches(0.86)
    ph = Inches(0)
    for _, _, _, _, lines in (left, right):
        tall = Inches(1.5)
        for ln in lines:
            tall += block_h(ln, inner / 914400, 12.5) + Inches(0.16)
        ph = max(ph, tall + Inches(0.26))
    for i, (ic, name, tag, tone, lines) in enumerate((left, right)):
        x = MARGIN + i * (pw + gap)
        warm = tone == AMBER
        rect(s, x, y, pw, ph, fill=MINT_TINT if warm else MIST, line=MINT if warm else CLOUD)
        rect(s, x, y, pw, Inches(0.05), fill=tone)
        if ic:
            icon(s, ic, x + Inches(0.34), y + Inches(0.36), Inches(0.48), color=tone)
        if tag:
            text(s, x + pw - Inches(2.0), y + Inches(0.44), Inches(1.7), Inches(0.24),
                 [(tag.upper(), dict(size=9.5, bold=True, color=tone, spacing=1.3))], align=PP_ALIGN.RIGHT)
        text(s, x + Inches(0.34), y + Inches(1.02), pw - Inches(0.68), Inches(0.34),
             [(name, dict(font=DISPLAY, size=22, color=FOREST))])
        yy = y + Inches(1.5)
        for ln in lines:
            rect(s, x + Inches(0.34), yy + Inches(0.08), Inches(0.07), Inches(0.07), fill=tone, shape=MSO_SHAPE.OVAL)
            lh = block_h(ln, inner / 914400, 12.5)
            text(s, x + Inches(0.52), yy, inner, lh, [(ln, dict(size=12.5, color=INK_800))], line_spacing=1.28)
            yy += lh + Inches(0.16)
    return y + ph


def milestones(s, y, items, *, card_h=Inches(2.1)):
    """A dotted track with a card per milestone. Done is a filled dot, ahead is hollow.
       items: (icon, when, title, line, done)."""
    n = len(items)
    gap = Inches(0.28)
    cw = (BODY_W - gap * (n - 1)) // n
    inner = (cw - Inches(0.52)) / 914400
    # a title that wraps pushes its line down instead of printing over it
    th = max(block_h(t[2], inner, 19, spacing=1.1, font=DISPLAY) for t in items)
    need = Inches(0.82) + th + Inches(0.12) + max(block_h(t[3], inner, 11.5) for t in items) + Inches(0.3)
    mh = max(card_h, need)
    rect(s, MARGIN, y + Inches(0.12), BODY_W, Emu(19050), fill=GREEN_300)
    for i, (ic, when, what, detail, done) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        rect(s, x, y, Inches(0.26), Inches(0.26), fill=GREEN if done else WHITE,
             line=None if done else GREEN_300, line_w=Pt(1.6), shape=MSO_SHAPE.OVAL)
        cy = y + Inches(0.46)
        rect(s, x, cy, cw, mh, fill=MIST if done else MINT_TINT, line=CLOUD)
        rect(s, x, cy, cw, Inches(0.05), fill=GREEN if done else GREEN_300)
        if ic:
            icon(s, ic, x + Inches(0.26), cy + Inches(0.26), Inches(0.4))
        text(s, x + cw - Inches(1.7), cy + Inches(0.3), Inches(1.44), Inches(0.24),
             [(when.upper(), dict(size=9.5, bold=True, color=GREEN_600, spacing=1.2))], align=PP_ALIGN.RIGHT)
        text(s, x + Inches(0.26), cy + Inches(0.82), cw - Inches(0.52), th,
             [(what, dict(font=DISPLAY, size=19, color=FOREST))], line_spacing=1.1)
        dh = block_h(detail, inner, 11.5)
        text(s, x + Inches(0.26), cy + Inches(0.82) + th + Inches(0.12), cw - Inches(0.52), dh,
             [(detail, dict(size=11.5, color=INK_MUTED))], line_spacing=1.28)
    return y + Inches(0.46) + mh


def chain(s, y, items, *, card_h=Inches(2.6)):
    """Cards joined by arrows: each artifact inherits from the one before.
       items: (icon, title, line)."""
    arrow_w = Inches(0.42)
    n = len(items)
    cw = (BODY_W - arrow_w * (n - 1)) // n
    inner = (cw - Inches(0.44)) / 914400
    tallest = max(block_h(t[2], inner, 11.5) for t in items)
    ch = max(card_h, Inches(0.62 + 0.4 + 0.34 + 0.16 + 0.6) + tallest)
    for i, (ic, name, line) in enumerate(items):
        x = MARGIN + i * (cw + arrow_w)
        rect(s, x, y, cw, ch, fill=MIST if i % 2 == 0 else MINT_TINT, line=CLOUD)
        rect(s, x, y, cw, Inches(0.05), fill=GREEN)
        lh = block_h(line, inner, 11.5)
        blk = Inches(0.62) + Inches(0.4) + Inches(0.34) + Inches(0.16) + lh
        top = y + (ch - blk) // 2
        icon(s, ic, x + (cw - Inches(0.62)) // 2, top, Inches(0.62))
        text(s, x + Inches(0.2), top + Inches(1.02), cw - Inches(0.4), Inches(0.34),
             [(name, dict(font=DISPLAY, size=19, color=FOREST))], align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.22), top + Inches(1.52), cw - Inches(0.44), lh,
             [(line, dict(size=11.5, color=INK_800))], align=PP_ALIGN.CENTER, line_spacing=1.28)
        if i < n - 1:
            icon(s, 'arrow-right', x + cw + (arrow_w - Inches(0.26)) // 2, y + ch // 2 - Inches(0.13),
                 Inches(0.26), color=GREEN_600)
    return y + ch


def logo_row(s, y, items, *, card_h=Inches(3.0), connect=True):
    """Partner cards with real logos, fitted to one box so wide and square marks read
       level. items: (logo_path or 'house', title, line). 'house' (or the brand id,
       or 'servantium' in older specs) draws the dark card with the brand's white logo."""
    from PIL import Image
    from .kit import LOGO_ON_DARK
    n = len(items)
    gap = Inches(0.36)
    cw = (BODY_W - gap * (n - 1)) // n
    inner = cw - Inches(0.56)
    for i, (logo, name, line) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        house = logo in HOUSE_KEYS
        rect(s, x, y, cw, card_h, fill=FOREST if house else WHITE, line=None if house else CLOUD)
        rect(s, x, y, cw, Inches(0.05), fill=GREEN)
        sh = block_h(line, inner / 914400, 11.5)
        blk = Inches(0.78) + Inches(0.42) + Inches(0.34) + Inches(0.14) + sh
        top = y + (card_h - blk) // 2
        path = LOGO_ON_DARK if house else Path(logo)
        im = Image.open(path)
        box_w, box_h = Inches(2.0), Inches(0.78)
        sc = min(box_w / im.width, box_h / im.height)
        w_, h_ = int(im.width * sc), int(im.height * sc)
        s.shapes.add_picture(str(path), x + (cw - w_) // 2, top + (box_h - h_) // 2, width=w_, height=h_)
        text(s, x + Inches(0.28), top + Inches(1.2), inner, Inches(0.34),
             [(name, dict(font=DISPLAY, size=19, color=WHITE if house else FOREST))], align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.28), top + Inches(1.68), inner, sh,
             [(line, dict(size=11.5, color=MIST_INK if house else INK_MUTED))], align=PP_ALIGN.CENTER,
             line_spacing=1.28)
        if connect and i < n - 1:
            icon(s, 'swap_horiz', x + cw + (gap - Inches(0.38)) // 2, y + card_h // 2 - Inches(0.19),
                 Inches(0.38), color=GREEN_600)
    return y + card_h


def people(s, y, items, *, workdir='.', logos_label=None, min_h=Inches(2.9)):
    """Team bios: a round photo, name, role, two short credential lines, and the logos of
       where they worked before. items: (photo, name, role, [line, ...], [logo, ...]).
       Photos are cropped to a circle into the workdir. Logos are sized to the same visual
       area, not the same height, so a heavy wordmark doesn't shout over a small icon.
       The cards grow to fit the longest bio, and the logo strip sits below the text."""
    from PIL import Image, ImageDraw
    n = len(items)
    gap = Inches(0.4)
    cw = _cols(n, gap)
    ph = Inches(1.5)
    tx = Inches(0.36) + ph + Inches(0.36)
    tw = cw - tx - Inches(0.3)

    def text_bottom(lines):
        yy = Inches(1.34)
        for ln in lines:
            yy += block_h(ln, tw / 914400, 12) + Inches(0.1)
        return yy

    body = max(max(text_bottom(it[3]) for it in items), Inches(0.4) + ph)
    has_logos = any(it[4] for it in items)
    card_h = max(min_h, body + (Inches(0.98) if has_logos else Inches(0.3)))
    for i, (photo, name, role, lines, logos) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        rect(s, x, y, cw, card_h, fill=WHITE, line=CLOUD)
        rect(s, x, y, cw, Inches(0.05), fill=GREEN)
        im = Image.open(photo).convert('RGBA')
        side = min(im.size)
        im = im.crop(((im.width - side) // 2, 0, (im.width - side) // 2 + side, side)).resize((600, 600))
        # A cut-out (transparent background) sits on a mint disc; the circle keeps the
        # photo's own transparency instead of turning it black.
        disc = Image.new('RGBA', (600, 600), (0, 0, 0, 0))
        ImageDraw.Draw(disc).ellipse((0, 0, 599, 599), fill=tuple(int(str(MINT)[i:i + 2], 16) for i in (0, 2, 4)) + (255,))
        im = Image.alpha_composite(disc, im)
        mask = Image.new('L', (600, 600), 0)
        ImageDraw.Draw(mask).ellipse((0, 0, 599, 599), fill=255)
        from PIL import ImageChops
        im.putalpha(ImageChops.multiply(im.getchannel('A'), mask))
        out = Path(workdir) / f'person-{abs(hash(str(photo))) % 10**8}.png'
        out.parent.mkdir(parents=True, exist_ok=True)
        im.save(out)
        s.shapes.add_picture(str(out), x + Inches(0.36), y + Inches(0.4), width=ph, height=ph)
        text(s, x + tx, y + Inches(0.42), tw, Inches(0.26),
             [(role.upper(), dict(size=10.5, bold=True, color=GREEN_600, spacing=1.4))])
        text(s, x + tx, y + Inches(0.74), tw, Inches(0.5),
             [(name, dict(font=DISPLAY, size=24, color=FOREST))])
        yy = y + Inches(1.34)
        for ln in lines:
            lh = block_h(ln, tw / 914400, 12)
            text(s, x + tx, yy, tw, lh, [(ln, dict(size=12, color=INK_800))], line_spacing=1.28)
            yy += lh + Inches(0.1)
        if logos:
            ly = y + card_h - Inches(0.7)
            rect(s, x + Inches(0.36), ly - Inches(0.14), cw - Inches(0.72), Emu(9525), fill=CLOUD)
            text(s, x + Inches(0.36), ly + Inches(0.14), Inches(1.9), Inches(0.24),
                 [((logos_label or f'Before {BRAND_SHORT}').upper(), dict(size=9.5, bold=True, color=INK_MUTED, spacing=1.2))])
            lx = x + Inches(2.3)
            area, max_h = 0.36, 0.42          # square inches; inches
            for logo in logos:
                li = Image.open(logo)
                aspect = li.width / li.height
                h_in = min(max_h, (area / aspect) ** 0.5)
                h_ = Inches(h_in)
                w_ = int(h_ * aspect)
                s.shapes.add_picture(str(logo), lx, ly + (Inches(max_h) - h_) // 2 + Inches(0.04), width=w_, height=h_)
                lx += w_ + Inches(0.4)
    return y + card_h


# ── plans, flows and layers (added for plan decks) ───────────────────────────
_BAR = {'green': (GREEN, WHITE), 'forest': (FOREST, WHITE), 'amber': (AMBER, WHITE),
        'mint': (MINT, FOREST), 'slate': (SLATE, WHITE)}


def _mark(s, name, x, y, size):
    """An icon name, or a path to a logo image (partner logos live in the deck's
       workspace, never in this repo). Logos are fitted into the square."""
    if isinstance(name, str) and name.lower().endswith(('.png', '.jpg', '.jpeg')):
        from PIL import Image
        w, h = Image.open(name).size
        scale = min(size / w, size / h)
        pw, ph = int(w * scale), int(h * scale)
        s.shapes.add_picture(name, x + (size - pw) // 2, y + (size - ph) // 2, width=pw, height=ph)
    else:
        icon(s, name, x, y, size)


def _round(shape, radius=Inches(0.12)):
    """Fixed corner radius. The default is a share of the short side, so tall cards
       came out nearly oval."""
    shape.adjustments[0] = min(0.5, radius / max(1, min(shape.width, shape.height)))
    return shape


def _line(s, x1, y1, x2, y2, color, w=Pt(1.25), dash=False, curve=False):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    x1, y1, x2, y2 = (int(round(v)) for v in (x1, y1, x2, y2))     # EMUs are whole numbers
    if y1 == y2 or x1 == x2:
        curve = False           # a curve with no height (or width) is degenerate; draw it straight
    c = s.shapes.add_connector(MSO_CONNECTOR.CURVE if curve else MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = w
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def gantt(s, y, periods, rows, gates=(), *, label_w=Inches(2.7), max_h=Inches(4.4)):
    """A plan as a Gantt chart: workstream rows, period columns, bars laid over the grid,
       and gates as diamonds on the column boundary they close.
       rows: (name, note, [(from, to, label, tone), ...]) with 1-based inclusive periods.
       gates: (after_period, label)."""
    head_h, gate_h = Inches(0.44), Inches(0.62) if gates else 0
    row_h = min(Inches(0.62), (max_h - head_h - gate_h) // max(len(rows), 1))
    grid_x = MARGIN + label_w
    grid_w = BODY_W - label_w
    col_w = grid_w // len(periods)
    body_h = row_h * len(rows)

    rect(s, MARGIN, y, BODY_W, head_h, fill=FOREST)
    text(s, MARGIN + Inches(0.2), y + Inches(0.14), label_w - Inches(0.3), Inches(0.2),
         [('WORKSTREAM', dict(size=9.5, bold=True, color=GREEN_300, spacing=1.3))])
    for i, p in enumerate(periods):
        text(s, grid_x + i * col_w, y + Inches(0.14), col_w, Inches(0.2),
             [(p.upper(), dict(size=9, bold=True, color=GREEN_300, spacing=1.0))], align=PP_ALIGN.CENTER)
    top = y + head_h
    for r, (name, note, _) in enumerate(rows):
        ry = top + r * row_h
        if r % 2:
            rect(s, MARGIN, ry, BODY_W, row_h, fill=MIST)
        rect(s, MARGIN, ry + row_h - Emu(9525), BODY_W, Emu(9525), fill=CLOUD)
        nh = Inches(0.22) + (Inches(0.18) if note else 0)
        text(s, MARGIN + Inches(0.2), ry + (row_h - nh) // 2, label_w - Inches(0.3), Inches(0.22),
             [(name, dict(size=12, bold=True, color=FOREST))])
        if note:
            text(s, MARGIN + Inches(0.2), ry + (row_h - nh) // 2 + Inches(0.22), label_w - Inches(0.3),
                 Inches(0.18), [(note, dict(size=9.5, color=INK_MUTED))])
    for i in range(1, len(periods)):
        rect(s, grid_x + i * col_w, top, Emu(6350), body_h, fill=CLOUD)
    rect(s, grid_x, top, Emu(9525), body_h, fill=CLOUD)

    if gates:       # the dashed gate lines go under the bars, so they never cross a bar's label
        for after, _ in gates:
            gx = grid_x + int(after * col_w)
            _line(s, gx, top, gx, top + body_h + Inches(0.1), GREEN_600, w=Pt(1.25), dash=True)

    for r, (_, _, bars) in enumerate(rows):
        ry = top + r * row_h
        for (a, b, label, tone) in bars:
            fill, ink = _BAR.get(tone or 'green', _BAR['green'])
            bx = grid_x + (a - 1) * col_w + Inches(0.06)
            bw = (b - a + 1) * col_w - Inches(0.12)
            bh = row_h - Inches(0.2)
            rect(s, bx, ry + Inches(0.1), bw, bh, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, bx + Inches(0.12), ry + Inches(0.1) + (bh - Inches(0.2)) // 2, bw - Inches(0.2),
                 Inches(0.2), [(label, dict(size=10.5, bold=True, color=ink))])

    if gates:
        gy = top + body_h
        for after, label in gates:
            gx = grid_x + int(after * col_w)
            d = Inches(0.2)
            rect(s, gx - d // 2, gy + Inches(0.06), d, d, fill=GREEN, shape=MSO_SHAPE.DIAMOND)
            lw = Inches(text_w_in(label, 9.5, bold=True) + 0.1)
            lx = min(max(gx - lw // 2, MARGIN), MARGIN + BODY_W - lw)   # centered, kept on the slide
            text(s, lx, gy + Inches(0.3), lw, Inches(0.2),
                 [(label, dict(size=9.5, bold=True, color=FOREST))], align=PP_ALIGN.CENTER)
    return top + body_h + gate_h


def flow(s, y, steps, active, inputs, node, output):
    """The pipeline picture: stage tabs, the inputs that feed this stage, Servantium in
       the middle, and the document it drafts. steps: [label]; active: 1-based;
       inputs: (icon, title, meta, tag, carried); node: (title, verb, chip);
       output: (label, ref, title, [(key, text, open)], foot). With no steps, the picture
       starts at the top: for a flow that runs across every stage (checks on all inputs)."""
    steps = steps or []
    n = max(len(steps), 1)
    gap = Inches(0.08)
    tw = (BODY_W - gap * (n - 1)) // n
    for i, lbl in enumerate(steps):
        x = MARGIN + i * (tw + gap)
        on, done = i + 1 == active, i + 1 < active
        if on:
            rect(s, x, y, tw, Inches(0.5), fill=MINT_TINT)
        text(s, x + Inches(0.14), y + Inches(0.14), tw - Inches(0.2), Inches(0.24),
             [[(f'{i + 1:02d}  ', dict(size=9.5, color=GREEN_600 if (on or done) else INK_MUTED)),
               (lbl, dict(size=13, bold=on, color=FOREST if (on or done) else INK_MUTED))]])
        rect(s, x, y + Inches(0.48), tw, Emu(19050 if (on or done) else 12700),
             fill=GREEN_600 if (on or done) else CLOUD)

    top = y + (Inches(0.95) if steps else Inches(0.32))
    in_w, card_h, cgap = Inches(3.95), Inches(0.8), Inches(0.14)
    out_w = Inches(3.75)
    out_x = MARGIN + BODY_W - out_w
    node_w, node_h = Inches(2.8), Inches(1.95 if len(node) > 3 and node[3] else 1.75)
    node_x = MARGIN + in_w + (out_x - MARGIN - in_w - node_w) // 2

    rows = output[3]
    out_h = Inches(0.95) + Inches(0.42) * len(rows) + (Inches(0.5) if output[4] else Inches(0.1))
    in_h = card_h * len(inputs) + cgap * (len(inputs) - 1)
    band = max(out_h, in_h, node_h)
    node_y = top + (band - node_h) // 2
    in_y = top + (band - in_h) // 2

    text(s, MARGIN, in_y - Inches(0.3), Inches(1), Inches(0.2),
         [('IN', dict(size=9, color=INK_MUTED, spacing=1.5))])
    for j, (ic, t, meta, tag, carried) in enumerate(inputs):
        cy = in_y + j * (card_h + cgap)
        _round(rect(s, MARGIN, cy, in_w, card_h, fill=MINT_TINT if carried else WHITE,
                    line=GREEN_300 if carried else CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE))
        rect(s, MARGIN + Inches(0.16), cy + Inches(0.18), Inches(0.44), Inches(0.44), fill=MINT,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _mark(s, ic or 'description', MARGIN + Inches(0.2), cy + Inches(0.22), Inches(0.36))
        tsize = 12.5 if est_lines(t, (in_w - Inches(1.9)) / 914400, 12.5, bold=True) <= 1 else 11
        text(s, MARGIN + Inches(0.76), cy + Inches(0.15), in_w - Inches(1.9), Inches(0.24),
             [(t, dict(size=tsize, bold=True, color=INK))])
        msize = 10.5 if est_lines(meta, (in_w - Inches(1.9)) / 914400, 10.5) <= 1 else 9.5
        text(s, MARGIN + Inches(0.76), cy + Inches(0.45), in_w - Inches(1.9), Inches(0.2),
             [(meta, dict(size=msize, color=INK_MUTED))])
        if tag:
            tag_w = Inches(0.9)
            tx = MARGIN + in_w - tag_w - Inches(0.14)
            if carried:
                rect(s, tx, cy + Inches(0.26), tag_w, Inches(0.28), fill=GREEN_600,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, tx, cy + Inches(0.31), tag_w, Inches(0.2),
                 [(tag.upper(), dict(size=8.5, bold=True, color=WHITE if carried else INK_MUTED, spacing=1.0))],
                 align=PP_ALIGN.CENTER)
        _line(s, MARGIN + in_w, cy + card_h // 2, node_x, node_y + node_h // 2, GREEN_300, curve=True)

    _round(rect(s, node_x, node_y, node_w, node_h, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE),
           Inches(0.18))
    ntitle, verb, chip = node[:3]
    cycle = node[3] if len(node) > 3 else None
    if ntitle.strip().lower() in HOUSE_KEYS:
        lw_ = Inches(1.9)
        s.shapes.add_picture(str(LOGO_ON_DARK), node_x + (node_w - lw_) // 2, node_y + Inches(0.24), width=lw_)
    else:
        text(s, node_x, node_y + Inches(0.3), node_w, Inches(0.34),
             [(ntitle, dict(font=DISPLAY, size=20, color=WHITE))], align=PP_ALIGN.CENTER)
    text(s, node_x + Inches(0.15), node_y + Inches(0.74), node_w - Inches(0.3), Inches(0.22),
         [(verb, dict(size=11.5, color=MIST_INK))], align=PP_ALIGN.CENTER)
    if cycle:
        _cycle(s, node_x + node_w // 2, node_y + Inches(1.36), cycle)
    elif chip:
        cw = Inches(2.2)
        rect(s, node_x + (node_w - cw) // 2, node_y + Inches(1.18), cw, Inches(0.3), fill=None,
             line=GREEN_300, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, node_x + (node_w - cw) // 2, node_y + Inches(1.24), cw, Inches(0.2),
             [(chip.upper(), dict(size=8.5, color=GREEN_300, spacing=1.2))], align=PP_ALIGN.CENTER)
    out_y = top + (band - out_h) // 2
    _line(s, node_x + node_w, node_y + node_h // 2, out_x, node_y + node_h // 2, GREEN_600, w=Pt(1.5))

    label, ref, otitle, orows, foot = output
    text(s, out_x, out_y - Inches(0.3), Inches(1), Inches(0.2),
         [('OUT', dict(size=9, color=INK_MUTED, spacing=1.5))])
    card = _round(rect(s, out_x, out_y, out_w, out_h, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE))
    from .kit import _soft_shadow
    _soft_shadow(card)
    ix, iw = out_x + Inches(0.22), out_w - Inches(0.44)
    text(s, ix, out_y + Inches(0.2), iw, Inches(0.2),
         [(label.upper(), dict(size=9, color=GREEN_600, spacing=1.3))])
    if ref:
        text(s, ix, out_y + Inches(0.2), iw, Inches(0.2),
             [(ref, dict(size=9.5, color=INK_MUTED))], align=PP_ALIGN.RIGHT)
    osize = 17 if est_lines(otitle, iw / 914400, 17, font=DISPLAY) <= 1 else 14   # one line, or it runs into the rows
    text(s, ix, out_y + Inches(0.46), iw, Inches(0.34),
         [(otitle, dict(font=DISPLAY, size=osize, color=INK))])
    ry = out_y + Inches(0.95)
    for key, val, is_open in orows:
        rect(s, ix, ry, iw, Emu(9525), fill=CLOUD)
        text(s, ix, ry + Inches(0.11), Inches(0.9), Inches(0.2),
             [(key, dict(size=10.5, bold=True, color=AMBER if is_open else GREEN_600))])
        text(s, ix + Inches(0.95), ry + Inches(0.11), iw - Inches(0.95), Inches(0.2),
             [(val, dict(size=11, color=INK))])
        ry += Inches(0.42)
    if foot:
        rect(s, ix, ry + Inches(0.02), iw, Emu(9525), fill=CLOUD)
        text(s, ix, ry + Inches(0.14), iw, Inches(0.2), [(foot, dict(size=9.5, color=GREEN_600))])
    return top + band


def stack(s, y, tiers, foundation, *, label_w=Inches(2.3)):
    """Layers, top to bottom, on a foundation. tiers: (label, note, [(name, icon)]);
       foundation: (title, text, [chips]). Connector ticks drop from each tier to the
       foundation, so the picture reads as everything standing on one record."""
    tier_h, tgap = (Inches(0.82), Inches(0.3)) if len(tiers) < 3 else (Inches(0.64), Inches(0.22))
    x0, w0 = MARGIN + label_w, BODY_W - label_w
    gap = Inches(0.18)
    boxes = []                       # per tier: [(x, w)], so ticks only land on a box below
    for (_, _, items) in tiers:
        n = len(items)
        pw = (w0 - gap * (n - 1)) // n
        boxes.append([(x0 + i * (pw + gap), pw) for i in range(n)])
    yy = y
    for t, (label, note, items) in enumerate(tiers):
        # a long label wraps; the note sits under however many lines it took
        # tracked caps run wider than the measure assumes, so measure against a narrower box
        lh = block_h(label.upper(), (label_w - Inches(0.2)) / 914400 * 0.82, 10, bold=True, spacing=1.15)
        text(s, MARGIN, yy + Inches(0.18), label_w - Inches(0.2), lh,
             [(label.upper(), dict(size=10, bold=True, color=GREEN_600, spacing=1.3))])
        if note:
            text(s, MARGIN, yy + Inches(0.22) + lh, label_w - Inches(0.2), Inches(0.2),
                 [(note, dict(size=10, color=INK_MUTED))])
        last = t == len(tiers) - 1
        for (px, pw), (nm, ic) in zip(boxes[t], items):
            cx_ = px + pw // 2
            # a tick drops only where it meets something: the foundation under the last tier,
            # or a box (not a gap) in the tier below
            if last or any(bx + Inches(0.15) <= cx_ <= bx + bw - Inches(0.15) for bx, bw in boxes[t + 1]):
                rect(s, cx_, yy + tier_h, Emu(12700), tgap, fill=GREEN_300)
            _round(rect(s, px, yy, pw, tier_h, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE))
            _mark(s, ic or 'apps', px + Inches(0.2), yy + (tier_h - Inches(0.38)) // 2, Inches(0.38))
            text(s, px + Inches(0.7), yy + (tier_h - Inches(0.24)) // 2, pw - Inches(0.8), Inches(0.24),
                 [(nm, dict(size=12.5, bold=True, color=INK_800))])
        yy += tier_h + tgap

    ftitle, ftext, chips = foundation
    fh = Inches(1.5)
    rect(s, MARGIN, yy, BODY_W, fh, fill=FOREST)
    rect(s, MARGIN, yy, BODY_W, Inches(0.05), fill=GREEN)
    s.shapes.add_picture(str(LOGO_ON_DARK), MARGIN + Inches(0.3), yy + Inches(0.3), width=Inches(1.7))
    text(s, MARGIN + Inches(0.3), yy + Inches(0.78), label_w - Inches(0.1), Inches(0.5),
         [(ftitle, dict(size=11, bold=True, color=GREEN_300))], line_spacing=1.2)
    if ftext:
        text(s, x0, yy + Inches(0.24), w0 - Inches(0.3), Inches(0.24),
             [(ftext, dict(font=DISPLAY, size=16, color=WHITE))])
    cx, cy = x0, yy + Inches(0.8)
    for c in chips:
        cw = Inches(0.3) + Inches(text_w_in(c, 10.5, bold=True))
        rect(s, cx, cy, cw, Inches(0.36), fill=None, line=GREEN_300, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, cx, cy + Inches(0.08), cw, Inches(0.2),
             [(c, dict(size=10.5, bold=True, color=MIST_INK))], align=PP_ALIGN.CENTER)
        cx += cw + Inches(0.12)
    return yy + fh


def shipped(s, y, groups, *, max_h=Inches(4.5)):
    """A wall of what shipped, grouped. groups: (title, [(icon, label), ...]). Volume is
       the point, so items are compact chips and the groups sit side by side."""
    n = len(groups)
    gap = Inches(0.3)
    cw = (BODY_W - gap * (n - 1)) // n
    chip_h, cgap = Inches(0.46), Inches(0.1)
    tallest = max(len(g[1]) for g in groups)
    h = Inches(0.62) + tallest * (chip_h + cgap) + Inches(0.2)
    for i, (gt, items) in enumerate(groups):
        x = MARGIN + i * (cw + gap)
        rect(s, x, y, cw, h, fill=MIST, line=CLOUD)
        rect(s, x, y, cw, Inches(0.05), fill=GREEN)
        text(s, x + Inches(0.24), y + Inches(0.2), cw - Inches(0.48), Inches(0.26),
             [(gt.upper(), dict(size=11, bold=True, color=GREEN_600, spacing=1.4))])
        for j, (ic, label) in enumerate(items):
            cy = y + Inches(0.62) + j * (chip_h + cgap)
            _round(rect(s, x + Inches(0.18), cy, cw - Inches(0.36), chip_h, fill=WHITE, line=CLOUD,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.08))
            icon(s, ic or 'check_circle', x + Inches(0.3), cy + (chip_h - Inches(0.24)) // 2, Inches(0.24))
            text(s, x + Inches(0.66), cy + (chip_h - Inches(0.22)) // 2, cw - Inches(0.96), Inches(0.22),
                 [(label, dict(size=11.5, bold=True, color=INK_800))])
    return y + h


def split_hero(prs, left, right, *, tag=None, title_=None, workdir=None):
    """Full-bleed marketing slide cut corner to corner into two triangles, a photo behind
       each, one big result per side. left/right: (image, label, stat, headline, line).
       The photos are composited into one picture; every word stays editable text."""
    from PIL import Image, ImageDraw
    from .kit import blank, W as SW, H as SH
    s = blank(prs)
    PW, PH = 2400, 1350

    def cover(path):
        im = Image.open(path).convert('RGB')
        r = max(PW / im.width, PH / im.height)
        im = im.resize((int(im.width * r) + 1, int(im.height * r) + 1))
        l, t_ = (im.width - PW) // 2, (im.height - PH) // 2
        return im.crop((l, t_, l + PW, t_ + PH))

    def tint(im, rgb, a):
        return Image.blend(im, Image.new('RGB', im.size, rgb), a)

    fr = tuple(int(str(FOREST)[i:i + 2], 16) for i in (0, 2, 4))
    ft = tuple(int(str(FOREST_900)[i:i + 2], 16) for i in (0, 2, 4))
    a = tint(cover(left[0]), fr, 0.62)
    b = tint(cover(right[0]), ft, 0.58)
    mask = Image.new('L', (PW, PH), 0)
    ImageDraw.Draw(mask).polygon([(0, 0), (PW, 0), (0, PH)], fill=255)
    comp = Image.composite(a, b, mask)
    d = ImageDraw.Draw(comp)
    g = tuple(int(str(GREEN)[i:i + 2], 16) for i in (0, 2, 4))
    d.line([(PW, 0), (0, PH)], fill=g, width=8)
    out = Path(workdir or '.') / f'split-{abs(hash((str(left[0]), str(right[0])))) % 10**8}.jpg'
    out.parent.mkdir(parents=True, exist_ok=True)
    comp.save(out, quality=90)
    s.shapes.add_picture(str(out), 0, 0, width=SW, height=SH)

    def side(x, y, w, item, align):
        _, label, stat, headline, line = item
        text(s, x, y, w, Inches(0.26), [(label.upper(), dict(size=11, bold=True, color=GREEN_300, spacing=1.6))], align=align)
        # big numbers in the body face: Playfair's old-style zero made "60%" read "6o%"
        text(s, x, y + Inches(0.4), w, Inches(1.1), [(stat, dict(size=64, color=WHITE))], align=align)
        text(s, x, y + Inches(1.62), w, Inches(0.4), [(headline, dict(font=DISPLAY, size=22, color=WHITE))], align=align)
        if line:
            text(s, x, y + Inches(2.12), w, Inches(0.3), [(line, dict(size=13, color=MIST_INK))], align=align)

    top_y = Inches(0.75)
    if title_:
        # a slide title in the top-left, inside the left photo where the diagonal can't cross it
        text(s, Inches(0.7), Inches(0.52), Inches(7.0), Inches(0.5),
             [(title_, dict(font=DISPLAY, size=26, color=WHITE))])
        rect(s, Inches(0.7), Inches(1.12), Inches(0.9), Emu(28575), fill=GREEN)
        top_y = Inches(1.5)
    side(Inches(0.7), top_y, Inches(6.2), left, PP_ALIGN.LEFT)
    side(SW - Inches(6.9), SH - Inches(3.35), Inches(6.2), right, PP_ALIGN.RIGHT)
    if tag:
        rect(s, SW - Inches(4.3), Inches(0.5), Inches(3.6), Inches(0.36), fill=None, line=GREEN_300,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, SW - Inches(4.3), Inches(0.57), Inches(3.6), Inches(0.22),
             [(tag.upper(), dict(size=9.5, bold=True, color=GREEN_300, spacing=1.2))], align=PP_ALIGN.CENTER)
    return s


# ── workflows: swimlanes ─────────────────────────────────────────────────────
def _arrow(conn, dashed=False):
    """python-pptx has no arrowhead API; add the DrawingML tail end."""
    from lxml import etree
    ln = conn.line._get_or_add_ln()
    for tag in ('a:tailEnd',):
        for old in ln.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd'):
            ln.remove(old)
    tail = etree.SubElement(ln, '{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd')
    tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return conn


def _wire(s, pts, color, dashed=False, w=Pt(1.5)):
    """A route through points, drawn as straight segments; the last one carries the arrow."""
    from pptx.enum.shapes import MSO_CONNECTOR
    last = None
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        if (x1, y1) == (x2, y2):
            continue
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
        c.line.color.rgb, c.line.width = color, w
        if dashed:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        last = c
    if last is not None:
        _arrow(last, dashed)


def swimlane(s, y, lanes, steps, loops=(), phases=(), clock=None, *, label_w=Inches(1.3), max_h=None):
    """A process across lanes, left to right.

    lanes: [(name, note)]. steps: dicts with id, lanes (list of lane indexes; two or more
    adjacent lanes = a joint step), title, text, kind (task|call|optional|approval|end),
    target (time budget chip), num, icon, branch (an optional detour: sits below the main
    row between its neighbours, with a solid skip line past it). loops: (from_id, to_id,
    label) drawn dashed. phases: (label, first_id, last_id) bands over the columns. clock:
    a label on the band row (the total time budget). Boxes size to their content.
    """
    max_h = max_h or (SAFE_BOTTOM - y)
    main = [st for st in steps if not st.get('branch')]
    col_of = {st['id']: i for i, st in enumerate(main)}
    n = len(main)
    grid_x = MARGIN + label_w
    col_w = (BODY_W - label_w) // n
    box_w = col_w - Inches(0.18)
    tw_in = (box_w - Inches(0.2)) / 914400
    HEAD, TSZ, BSZ = Inches(0.4), 10, 8.5

    def need(st):
        th = block_h(st['title'], tw_in, TSZ, spacing=1.08, bold=True)
        bh = block_h(st['text'], tw_in, BSZ, spacing=1.12) if st.get('text') else 0
        return HEAD + th + (Inches(0.03) + bh if bh else 0) + Inches(0.12)
    box_h = max([Inches(0.9)] + [need(st) for st in main if st.get('kind') != 'end'])
    br_h = max([Inches(0.7)] + [need(st) for st in steps if st.get('branch')])
    pad = Inches(0.12)
    band_h = Inches(0.4) if (phases or clock) else 0
    has_branch = {ln for st in steps if st.get('branch') for ln in st['lanes']}

    def heights(p):
        return [box_h + p * 2 + (br_h + p if i in has_branch else 0) for i in range(len(lanes))]
    lane_h = heights(pad)
    while band_h + sum(lane_h) > max_h and pad > Inches(0.04):   # squeeze padding first
        pad -= Inches(0.01)
        lane_h = heights(pad)
    lane_y = []
    yy = y + band_h
    for h in lane_h:
        lane_y.append(yy); yy += h
    bottom = yy

    # Phase bands and the clock.
    if band_h:
        rect(s, MARGIN, y, BODY_W, band_h - Inches(0.06), fill=FOREST)
        if clock:
            text(s, MARGIN + Inches(0.16), y + Inches(0.1), label_w - Inches(0.2), Inches(0.22),
                 [(clock, dict(size=9.5, bold=True, color=GREEN_300))])
        for label, a, b in phases:
            if a not in col_of or b not in col_of:
                continue
            x1 = grid_x + col_of[a] * col_w
            x2 = grid_x + (col_of[b] + 1) * col_w
            rect(s, x1 + Inches(0.04), y + Inches(0.06), x2 - x1 - Inches(0.08), band_h - Inches(0.18),
                 fill=FOREST_900, line=GREEN_600, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, x1, y + Inches(0.1), x2 - x1, Inches(0.22),
                 [(label.upper(), dict(size=9.5, bold=True, color=WHITE, spacing=1.4))], align=PP_ALIGN.CENTER)

    # Lanes.
    for i, ((name, note), ly, lh) in enumerate(zip(lanes, lane_y, lane_h)):
        rect(s, MARGIN, ly, BODY_W, lh, fill=MIST if i % 2 == 0 else WHITE, line=CLOUD)
        rect(s, MARGIN, ly, label_w - Inches(0.1), lh, fill=MINT_TINT, line=CLOUD)
        rect(s, MARGIN, ly, Inches(0.06), lh, fill=GREEN)
        th = block_h(name, (label_w - Inches(0.4)) / 914400, 12, bold=True) + (Inches(0.36) if note else 0)
        text(s, MARGIN + Inches(0.2), ly + (lh - th) // 2, label_w - Inches(0.4), Inches(0.3),
             [(name, dict(size=12, bold=True, color=FOREST))], line_spacing=1.1)
        if note:
            text(s, MARGIN + Inches(0.2), ly + (lh - th) // 2 + th - Inches(0.34), label_w - Inches(0.4),
                 Inches(0.3), [(note, dict(size=9, color=INK_MUTED))], line_spacing=1.1)

    # Geometry of each step.
    geo = {}
    for i, st in enumerate(steps):
        ls = sorted(st['lanes'])
        if st.get('branch'):
            prev = next(p for p in reversed(steps[:i]) if not p.get('branch'))
            nxt = next(p for p in steps[i + 1:] if not p.get('branch'))
            cx = grid_x + (col_of[prev['id']] + col_of[nxt['id']] + 1) * col_w // 2
            x = cx - box_w // 2
            top = lane_y[ls[0]] + pad + box_h + pad
            h = br_h
        else:
            x = grid_x + col_of[st['id']] * col_w + (col_w - box_w) // 2
            top = lane_y[ls[0]] + pad
            h = (lane_y[ls[-1]] + pad + box_h) - top
        geo[st['id']] = (x, top, box_w, h)

    # Wires first, so boxes sit on top of them.
    def mid_r(g, at_y=None): return (g[0] + g[2], at_y if at_y is not None else g[1] + g[3] // 2)
    def mid_l(g, at_y=None): return (g[0], at_y if at_y is not None else g[1] + g[3] // 2)
    def row_y(st):         # where a wire meets a step: its own middle, or the shared row of a joint step
        g = geo[st['id']]
        return g[1] + min(g[3], box_h) // 2
    wires = []
    for i, st in enumerate(main[:-1]):
        nx = main[i + 1]
        wires.append((st, nx, False))
    for i, st in enumerate(steps):
        if st.get('branch'):
            prev = next(p for p in reversed(steps[:i]) if not p.get('branch'))
            nxt = next(p for p in steps[i + 1:] if not p.get('branch'))
            wires += [(prev, st, True), (st, nxt, True)]
    by_id = {st['id']: st for st in steps}
    for a, b, dashed in wires:
        ga, gb = geo[a['id']], geo[b['id']]
        if b.get('branch'):                     # down into the detour
            sx = ga[0] + ga[2] - Inches(0.25)
            pts = [(sx, ga[1] + ga[3]), (sx, gb[1] + gb[3] // 2), (gb[0], gb[1] + gb[3] // 2)]
        elif a.get('branch'):                   # back up to the main row
            ex = gb[0] + Inches(0.25)
            pts = [(ga[0] + ga[2], ga[1] + ga[3] // 2), (ex, ga[1] + ga[3] // 2), (ex, gb[1] + gb[3])]
        else:
            ya = row_y(a) if len(a['lanes']) == 1 else (row_y(b) if len(b['lanes']) == 1 and ga[1] <= row_y(b) <= ga[1] + ga[3] else ga[1] + ga[3] // 2)
            yb = row_y(b) if len(b['lanes']) == 1 else (ya if gb[1] <= ya <= gb[1] + gb[3] else gb[1] + gb[3] // 2)
            sx, ex = ga[0] + ga[2], gb[0]
            if ya == yb:
                pts = [(sx, ya), (ex, yb)]
            else:
                mx = (sx + ex) // 2
                pts = [(sx, ya), (mx, ya), (mx, yb), (ex, yb)]
        _wire(s, pts, GREEN_600, dashed=dashed)
    for fid, tid, label in loops:
        ga, gb = geo[fid], geo[tid]
        fx = ga[0] + ga[2] // 2
        tx = gb[0] + gb[2] // 2
        if gb[1] < ga[1]:                        # target lane above: up, across, up into its bottom
            gap_y = gb[1] + gb[3] + (ga[1] - gb[1] - gb[3]) // 2
            pts = [(fx, ga[1]), (fx, gap_y), (tx, gap_y), (tx, gb[1] + gb[3])]
            ly = gap_y - Inches(0.22)
        else:
            gap_y = ga[1] + ga[3] + Inches(0.18)
            pts = [(fx, ga[1] + ga[3]), (fx, gap_y), (tx, gap_y), (tx, gb[1])]
            ly = gap_y + Inches(0.02)
        _wire(s, pts, AMBER, dashed=True, w=Pt(1.25))
        if label:
            lw = Inches(text_w_in(label, 8.5, bold=True) + 0.16)
            lx = max(fx, tx) + Inches(0.06)
            rect(s, lx, gap_y - Inches(0.1), lw, Inches(0.2), fill=WHITE)
            text(s, lx + Inches(0.08), gap_y - Inches(0.08), lw, Inches(0.18),
                 [(label, dict(size=8.5, bold=True, color=AMBER))])

    # Boxes.
    KIND_ICON = {'call': 'groups', 'approval': 'verified', 'end': 'send'}
    for st in steps:
        x, top, w, h = geo[st['id']]
        kind = st.get('kind', 'task')
        if kind == 'end':
            ph = Inches(0.66)
            py = top + (min(h, box_h) - ph) // 2
            _round(rect(s, x, py, w, ph, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.33))
            text(s, x, py + Inches(0.13), w, Inches(0.2),
                 [(st['title'], dict(size=TSZ, bold=True, color=WHITE))], align=PP_ALIGN.CENTER)
            if st.get('text'):
                text(s, x, py + Inches(0.36), w, Inches(0.2),
                     [(st['text'], dict(size=BSZ, color=GREEN_300))], align=PP_ALIGN.CENTER)
            continue
        line_c = AMBER if kind == 'approval' else (GREEN_600 if kind == 'optional' else CLOUD)
        box = _round(rect(s, x, top, w, h, fill=WHITE, line=line_c, line_w=Pt(1.25),
                          shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.08))
        if kind == 'optional':
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            rect(s, x, top, w, Inches(0.05), fill=AMBER if kind == 'approval' else GREEN)
        ix = x + Inches(0.1)
        hx = ix
        if st.get('num') is not None:
            rect(s, ix, top + Inches(0.1), Inches(0.24), Inches(0.24), fill=FOREST, shape=MSO_SHAPE.OVAL)
            text(s, ix, top + Inches(0.135), Inches(0.24), Inches(0.18),
                 [(str(st['num']), dict(size=8.5, bold=True, color=WHITE))], align=PP_ALIGN.CENTER)
            hx += Inches(0.3)
        if st.get('target'):
            cw = Inches(text_w_in(st['target'], 8, bold=True) + 0.2)
            _round(rect(s, hx, top + Inches(0.12), cw, Inches(0.2), fill=MINT,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.1))
            text(s, hx, top + Inches(0.14), cw, Inches(0.16),
                 [(st['target'], dict(size=8, bold=True, color=FOREST))], align=PP_ALIGN.CENTER)
        if kind == 'optional':
            text(s, x + w - Inches(0.9), top + Inches(0.15), Inches(0.8), Inches(0.16),
                 [('OPTIONAL', dict(size=7.5, bold=True, color=GREEN_600, spacing=1.2))], align=PP_ALIGN.RIGHT)
        else:
            ic = st.get('icon') or KIND_ICON.get(kind)
            if ic:
                icon(s, ic, x + w - Inches(0.3), top + Inches(0.11), Inches(0.21),
                     color=AMBER if kind == 'approval' else GREEN_600)
        th = block_h(st['title'], tw_in, TSZ, spacing=1.08, bold=True)
        ty = top + HEAD
        text(s, ix, ty, w - Inches(0.2), th, [(st['title'], dict(size=TSZ, bold=True, color=FOREST))],
             line_spacing=1.08)
        if st.get('text'):
            bh = block_h(st['text'], tw_in, BSZ, spacing=1.12)
            text(s, ix, ty + th + Inches(0.03), w - Inches(0.2), bh,
                 [(st['text'], dict(size=BSZ, color=INK_800))], line_spacing=1.12)
    return bottom


# ── consulting pictures: adoption curve, approach matrix, zoom, team ─────────
# Added for a consulting pitch whose feedback was "too many rows of boxes": each of
# these has its own shape, so a deck can change the picture from slide to slide.

def _freeform(s, pts, fill, line=None, line_w=Pt(1)):
    """A closed polygon from slide-coordinate points (EMU)."""
    ff = s.shapes.build_freeform(int(pts[0][0]), int(pts[0][1]), scale=1.0)
    ff.add_line_segments([(int(x), int(y)) for x, y in pts[1:]], close=True)
    shp = ff.convert_to_shape()
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


ROGERS = [('Innovators', '2.5%', -3.0, -2.0), ('Early adopters', '13.5%', -2.0, -1.0),
          ('Early majority', '34%', -1.0, 0.0), ('Late majority', '34%', 0.0, 1.0),
          ('Laggards', '16%', 1.0, 3.0)]


def _arc_arrow(s, cx, cy, rx, ry, t0, t1, color, w=Pt(1.25), n=24):
    """An open elliptical arc from angle t0 to t1 (radians, y up), arrowhead at the end."""
    import math
    pts = [(cx + rx * math.cos(t0 + (t1 - t0) * k / n), cy - ry * math.sin(t0 + (t1 - t0) * k / n)) for k in range(n + 1)]
    ff = s.shapes.build_freeform(int(pts[0][0]), int(pts[0][1]), scale=1.0)
    ff.add_line_segments([(int(x), int(y)) for x, y in pts[1:]], close=False)
    shp = ff.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = color
    shp.line.width = w
    shp.shadow.inherit = False
    return _arrow(shp)


def _cycle(s, cx, cy, steps):
    """Two steps that go round and round (AI drafts -> you approve -> AI drafts ...): two pills
       with an arc over the top one way and under the bottom the other."""
    import math
    a, b = steps[0], steps[1]
    pw, ph, gap = Inches(1.02), Inches(0.3), Inches(0.26)
    lx = cx - gap // 2 - pw
    rx_ = cx + gap // 2
    for x, label, filled in ((lx, a, True), (rx_, b, False)):
        pill = rect(s, x, cy - ph // 2, pw, ph, fill=GREEN_300 if filled else None, line=GREEN_300,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _round(pill, ph // 2)
        size = 8.0
        room = (pw - Inches(0.16)) / 914400
        while size > 6.0 and text_w_in(label.upper(), size, bold=True) + len(label) * 0.8 / 72 > room:
            size -= 0.25
        text(s, x, cy - Inches(0.08), pw, Inches(0.18),
             [(label.upper(), dict(size=size, bold=True, color=FOREST if filled else GREEN_300, spacing=0.8))],
             align=PP_ALIGN.CENTER)
    # over the top, left pill to right pill; under the bottom, back again
    rx2 = (rx_ + pw // 2) - (lx + pw // 2)
    _arc_arrow(s, cx, cy - ph // 2, rx2 // 2, Inches(0.2), math.pi, 0.12, GREEN_300)
    _arc_arrow(s, cx, cy + ph // 2, rx2 // 2, Inches(0.2), 0.0, -math.pi + 0.12, GREEN_300)


def adoption(s, y, here, *, labels=None, note=None, h=Inches(3.5)):
    """The technology adoption curve (Rogers' segments) with Moore's chasm cut between
       early adopters and the early majority, and a marker saying who is at the chasm now.
       The shapes are real vector polygons, so the curve stays crisp and editable."""
    import math
    names = labels or [r[0] for r in ROGERS]
    gap = Inches(0.42)                       # the chasm
    label_band = Inches(0.92)
    top_pad = Inches(0.86)                   # room for the marker above the curve
    x0, w = MARGIN, BODY_W - gap
    base = y + h - label_band
    ch = base - (y + top_pad)
    sx = lambda u: x0 + (u + 3) / 6 * w + (gap if u > -1 + 1e-9 else 0)
    sy = lambda u: base - math.exp(-u * u / 2) * ch
    fills = [FOREST, GREEN_600, MINT, CLOUD, MIST]
    lines = [None, None, GREEN, None, None]
    for i, (_, _, a, b) in enumerate(ROGERS):
        n = 36
        us = [a + (b - a) * k / n for k in range(n + 1)]
        # the chasm side of each segment sits exactly on its boundary
        pts = [(sx(u) if not (abs(u + 1) < 1e-9 and i == 2) else x0 + (u + 3) / 6 * w + gap, sy(u)) for u in us]
        pts += [(pts[-1][0], base), (pts[0][0], base)]
        _freeform(s, pts, fills[i], lines[i], Pt(1.25))
    rect(s, x0, base, w + gap, Emu(12700), fill=CLOUD)
    # the chasm: a dashed drop and its name
    cx = int(x0 + 2 / 6 * w + gap // 2)
    _line(s, cx, sy(-1) - Inches(0.1), cx, base, INK_MUTED, w=Pt(1), dash=True)
    text(s, cx - Inches(0.9), base + Inches(0.64), Inches(1.8), Inches(0.24),
         [('THE CHASM', dict(size=9, bold=True, color=INK_MUTED, spacing=1.4))], align=PP_ALIGN.CENTER)
    # who is here: a forest pill above the chasm with a pointer down to it
    pw = Inches(max(2.2, text_w_in(here, 12, bold=True) + 0.5))
    py = y + Inches(0.06)
    pill = rect(s, cx - pw // 2, py, pw, Inches(0.46), fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _round(pill, Inches(0.23))
    text(s, cx - pw // 2, py + Inches(0.11), pw, Inches(0.26),
         [(here, dict(size=12, bold=True, color=WHITE))], align=PP_ALIGN.CENTER)
    _freeform(s, [(cx - Inches(0.11), py + Inches(0.45)), (cx + Inches(0.11), py + Inches(0.45)),
                  (cx, py + Inches(0.62))], FOREST)
    _line(s, cx, py + Inches(0.62), cx, sy(-1) - Inches(0.1), FOREST, w=Pt(1.25))
    if note:
        text(s, cx + pw // 2 + Inches(0.2), py + Inches(0.1), Inches(4.2), Inches(0.4),
             [(note, dict(size=11.5, color=INK_MUTED))])
    # segment names and shares under the baseline
    for i, (name, share, a, b) in enumerate(ROGERS):
        lx, rx = sx(a) if i != 2 else x0 + 2 / 6 * w + gap, sx(b)
        bw = max(rx - lx, Inches(1.3))
        mid = (lx + rx) // 2
        text(s, mid - bw // 2, base + Inches(0.1), bw, Inches(0.26),
             [(names[i], dict(size=12, bold=True, color=FOREST if i < 2 else INK_800))], align=PP_ALIGN.CENTER)
        text(s, mid - bw // 2, base + Inches(0.38), bw, Inches(0.22),
             [(share, dict(size=10.5, color=INK_MUTED))], align=PP_ALIGN.CENTER)
    return y + h


_CELL = {'plain': (MIST, None, INK_800, False), 'record': (MINT_TINT, None, FOREST, True),
         'ai': (WHITE, GREEN, FOREST, False), 'dark': (FOREST, None, WHITE, True)}


def approach(s, y, stages, rows, *, active=None, label_w=Inches(1.35)):
    """Phases across, layers down: a chevron per stage, then one row per layer (who does
       it, what moves forward, what you get). The consulting methodology picture; it also
       draws a value chain with layers. rows: (label, [cell per stage], tone)."""
    n = len(stages)
    gap = Inches(0.08)
    gx = MARGIN + label_w
    cw = (BODY_W - label_w - gap * (n - 1)) // n
    hh = Inches(0.52)
    for i, st in enumerate(stages):
        x = gx + i * (cw + gap)
        on = active is not None and i + 1 == active
        ext = Inches(0.12) if i < n - 1 else 0          # overlap the next stage; the last one stops at the margin
        shp = s.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, y, cw + ext, hh)
        shp.fill.solid(); shp.fill.fore_color.rgb = GREEN_600 if on else FOREST
        shp.line.fill.background(); shp.shadow.inherit = False
        shp.adjustments[0] = 0.32 if i == 0 else 0.42
        inset = Inches(0.16) if i else Inches(0.04)
        text(s, x + inset, y + Inches(0.14), cw - inset - Inches(0.06), Inches(0.26),
             [(st, dict(size=12.5, bold=True, color=WHITE))], align=PP_ALIGN.CENTER)
    yy = y + hh + Inches(0.12)
    inner = cw - Inches(0.2)
    for label, cells, tone in rows:
        bg, edge, fg, bold = _CELL.get(tone or 'plain', _CELL['plain'])
        rh = max([Inches(0.64)] + [block_h(c, inner / 914400, 11, bold=bold) + Inches(0.26) for c in cells])
        text(s, MARGIN, yy + rh // 2 - Inches(0.12), label_w - Inches(0.15), Inches(0.26),
             [(label.upper(), dict(size=9.5, bold=True, color=INK_MUTED, spacing=1.2))])
        for i, c in enumerate(cells):
            x = gx + i * (cw + gap)
            box = rect(s, x, yy, cw, rh, fill=bg, line=edge)
            if edge is not None:
                box.line.dash_style = 4      # dash
                box.line.width = Pt(1)
            ch = block_h(c, inner / 914400, 11, bold=bold)
            text(s, x + Inches(0.1), yy + (rh - ch) // 2, inner, ch,
                 [(c, dict(size=11, bold=bold, color=fg))], align=PP_ALIGN.CENTER, line_spacing=1.15)
        yy += rh + Inches(0.08)
    return yy - Inches(0.08)


def zoom(s, y, steps, active, title_, rows, *, meta=None, foot=None):
    """One step of a process, magnified: the steps stacked on the left with the active one
       lit, a light wedge opening from it into a panel that shows what happens inside that
       step. rows: (key, title, text)."""
    n = len(steps)
    pill_w, pill_h, pgap = Inches(2.35), Inches(0.46), Inches(0.14)
    px = MARGIN + Inches(3.55)
    pw = MARGIN + BODY_W - px
    inner = pw - Inches(0.6) - Inches(1.0)
    head_h = Inches(0.95) if meta else Inches(0.7)
    rhs = [max(Inches(0.62), Inches(0.3) + Inches(0.26) + (block_h(t, inner / 914400, 11.5) if t else 0))
           for _, _, t in rows]
    panel_h = head_h + sum(rhs) + (Inches(0.5) if foot else Inches(0.2))
    stack_h = n * pill_h + (n - 1) * pgap
    top = y
    sy0 = top + (panel_h - stack_h) // 2
    ay = sy0 + (active - 1) * (pill_h + pgap)
    # the wedge first, so the pill and panel sit on it
    _freeform(s, [(MARGIN + pill_w, ay), (px, top), (px, top + panel_h), (MARGIN + pill_w, ay + pill_h)], MINT)
    for i, st in enumerate(steps):
        on = i + 1 == active
        yy = sy0 + i * (pill_h + pgap)
        b = rect(s, MARGIN, yy, pill_w, pill_h, fill=FOREST if on else WHITE, line=None if on else CLOUD,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _round(b, Inches(0.08))
        text(s, MARGIN + Inches(0.2), yy + Inches(0.12), pill_w - Inches(0.3), Inches(0.24),
             [[(f'{i + 1:02d}  ', dict(size=10.5, bold=True, color=GREEN_300 if on else GREEN_600)),
               (st, dict(size=13, bold=on, color=WHITE if on else INK_800))]])
    card = rect(s, px, top, pw, panel_h, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _round(card, Inches(0.1))
    rect(s, px, top, pw, Inches(0.05), fill=GREEN)
    text(s, px + Inches(0.3), top + Inches(0.24), pw - Inches(0.6), Inches(0.36),
         [(title_, dict(font=DISPLAY, size=20, color=FOREST))])
    if meta:
        text(s, px + Inches(0.3), top + Inches(0.62), pw - Inches(0.6), Inches(0.24),
             [(meta, dict(size=11.5, color=INK_MUTED))])
    yy = top + head_h
    for (key, t1, t2), rh in zip(rows, rhs):
        rect(s, px + Inches(0.3), yy, pw - Inches(0.6), Emu(9525), fill=CLOUD)
        if key:
            text(s, px + Inches(0.3), yy + Inches(0.16), Inches(0.9), Inches(0.24),
                 [(key, dict(size=11, bold=True, color=GREEN_600))])
        text(s, px + Inches(1.3), yy + Inches(0.14), inner, Inches(0.26),
             [(t1, dict(size=13, bold=True, color=INK))])
        if t2:
            text(s, px + Inches(1.3), yy + Inches(0.42), inner, block_h(t2, inner / 914400, 11.5),
                 [(t2, dict(size=11.5, color=INK_MUTED))], line_spacing=1.2)
        yy += rh
    if foot:
        text(s, px + Inches(0.3), top + panel_h - Inches(0.38), pw - Inches(0.6), Inches(0.24),
             [(foot, dict(size=11, color=GREEN_600))])
    return top + max(panel_h, stack_h)


# fill, label, title, text: light steps carry forest type, dark steps white
_STAIR = {'mint_tint': (MINT_TINT, GREEN_600, FOREST, INK_800), 'mint': (MINT, GREEN_600, FOREST, INK_800),
          'jade': (GREEN_300, FOREST, FOREST, INK), 'green': (GREEN, WHITE, WHITE, WHITE),
          'forest': (FOREST, GREEN_300, WHITE, MIST_INK)}
_STAIR_TONES = {3: ['mint', 'green', 'forest'], 4: ['mint_tint', 'mint', 'green', 'forest'],
                5: ['mint_tint', 'mint', 'jade', 'green', 'forest']}


def staircase(s, y, steps, *, here=None, here_label='Start here', h_min=Inches(1.45), rise=Inches(0.34), numbered=True):
    """What opens at each step, rising left to right and darkening from mint to forest:
       the maturity staircase, for access or capability that widens one proven step at a
       time. The tallest text sets the first step; each step after it rises by `rise`.
       steps: (label, title, text, when); here: the 1-based step to flag."""
    n = len(steps)
    gap = Inches(0.06)
    sw = (BODY_W - gap * (n - 1)) // n
    pad = Inches(0.18)
    inner = sw - 2 * pad
    w_in = inner / 914400

    def need(step):
        label, t1, t2, _ = step
        return (pad + (Inches(0.3) if label else 0) + block_h(t1, w_in, 16, spacing=1.08, font=DISPLAY)
                + (Inches(0.08) + block_h(t2, w_in, 11, spacing=1.2) if t2 else 0) + pad)
    base_h = max([h_min] + [need(st) - rise * i for i, st in enumerate(steps)])
    flag_h = Inches(0.46)
    if here:          # room above the flagged step for its marker
        y += max(0, flag_h - (n - here) * rise)
    base = y + base_h + rise * (n - 1)
    for i, (label, t1, t2, when) in enumerate(steps):
        fill_, lab_c, t1_c, t2_c = _STAIR[_STAIR_TONES[n][i]]
        x = MARGIN + i * (sw + gap)
        h = base_h + rise * i
        top = base - h
        rect(s, x, top, sw, h, fill=fill_)
        ty = top + pad
        if label:
            text(s, x + pad, ty, inner, Inches(0.22),
                 [[*([(f'{i + 1:02d}  ', dict(size=10, bold=True, color=lab_c))] if numbered else []),
                   (label.upper(), dict(size=10, bold=True, color=lab_c, spacing=1.2))]])
            ty += Inches(0.3)
        th = block_h(t1, w_in, 16, spacing=1.08, font=DISPLAY)
        text(s, x + pad, ty, inner, th, [(t1, dict(font=DISPLAY, size=16, color=t1_c))], line_spacing=1.08)
        ty += th + Inches(0.08)
        if t2:
            text(s, x + pad, ty, inner, block_h(t2, w_in, 11, spacing=1.2),
                 [(t2, dict(size=11, color=t2_c))], line_spacing=1.2)
        if when:
            text(s, x, base + Inches(0.12), sw, Inches(0.22),
                 [(when.upper(), dict(size=9.5, bold=True, color=INK_MUTED, spacing=1.2))], align=PP_ALIGN.CENTER)
        if here and i + 1 == here:
            fw = Inches(min(2.2, 0.5 + 0.085 * len(here_label)))
            fx = x + pad
            fy = top - flag_h
            _round(rect(s, fx, fy, fw, Inches(0.32), fill=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.08))
            text(s, fx, fy + Inches(0.07), fw, Inches(0.2),
                 [(here_label.upper(), dict(size=9.5, bold=True, color=INK, spacing=1.1))], align=PP_ALIGN.CENTER)
            tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, fx + Inches(0.2), fy + Inches(0.31),
                                     Inches(0.18), Inches(0.12))
            tri.rotation = 180
            tri.fill.solid(); tri.fill.fore_color.rgb = AMBER
            tri.line.fill.background(); tri.shadow.inherit = False
    rect(s, MARGIN, base, BODY_W, Emu(12700), fill=CLOUD)
    return base + (Inches(0.42) if any(st[3] for st in steps) else Inches(0.02))


_WHEN_TONES = [(AMBER, INK), (GREEN, WHITE), (FOREST, WHITE), (SLATE, WHITE)]


def middleout(s, y, center, left, right, *, center_w=Inches(2.7), gap=Inches(0.46)):
    """A system in the middle of a stack, built outward: what is live in the center, and
       what gets built on each side, each item with when it lands, every item joined to the
       center by an arrow pointing out. The founder's picture of a sales-to-services platform:
       build left toward sales and right toward services at the same time.
       center: (title, [items], note); left/right: (label, [(icon, name, text, when)])."""
    ctitle, citems, cnote = center
    side_w = (BODY_W - center_w - 2 * gap) // 2
    cx = MARGIN + side_w + gap
    xs = {'left': MARGIN, 'right': cx + center_w + gap}
    order = []                                   # pill tones follow the order phases first appear
    for _, items in (left, right):
        for it in items:
            if it[3] and it[3] not in order:
                order.append(it[3])
    tone = {w: _WHEN_TONES[min(i, len(_WHEN_TONES) - 1)] for i, w in enumerate(order)}
    head_h = Inches(0.4)
    cgap = Inches(0.12)
    pill_w = lambda when: Inches(text_w_in(when.upper(), 8.5, bold=True) + 0.34) if when else 0
    text_w = lambda when: side_w - Inches(0.92) - pill_w(when) - Inches(0.16)

    def need(it):       # names and lines are measured, so a wrapped line grows the card
        _, name, sub, when = it
        tw = text_w(when) / 914400
        return (Inches(0.26) + block_h(name, tw, 12.5, spacing=1.15, bold=True)
                + (Inches(0.04) + block_h(sub, tw, 10.5, spacing=1.2) if sub else 0))
    card_h = max([Inches(0.76)] + [need(it) for _, items in (left, right) for it in items])
    n = max(len(left[1]), len(right[1]))
    stack_h = n * card_h + (n - 1) * cgap
    # the center box grows to hold its own list and note; it never spills past its bottom edge
    note_h = (block_h(cnote, (center_w - Inches(0.56)) / 914400, 11, spacing=1.25) + Inches(0.18)) if cnote else 0
    stack_h = max(stack_h, Inches(0.78) + len(citems) * Inches(0.36) + note_h + Inches(0.22))
    top = y + head_h

    # the center: what is live
    _round(rect(s, cx, top, center_w, stack_h, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.16))
    text(s, cx, y + Inches(0.08), center_w, Inches(0.22),
         [('IN THE MIDDLE', dict(size=9.5, bold=True, color=GREEN_600, spacing=1.4))], align=PP_ALIGN.CENTER)
    text(s, cx + Inches(0.28), top + Inches(0.24), center_w - Inches(0.56), Inches(0.4),
         [(ctitle, dict(font=DISPLAY, size=21, color=WHITE))])
    iy = top + Inches(0.78)
    for it in citems:
        icon(s, 'check_circle', cx + Inches(0.28), iy + Inches(0.02), Inches(0.24), color=GREEN_300)
        text(s, cx + Inches(0.62), iy, center_w - Inches(0.86), Inches(0.26),
             [(it, dict(size=12.5, color=WHITE))])
        iy += Inches(0.36)
    if cnote:
        nh = block_h(cnote, (center_w - Inches(0.56)) / 914400, 11, spacing=1.25)
        text(s, cx + Inches(0.28), top + stack_h - nh - Inches(0.22), center_w - Inches(0.56), nh,
             [(cnote, dict(size=11, color=GREEN_300))], line_spacing=1.25)

    for side, (label, items) in (('left', left), ('right', right)):
        x = xs[side]
        text(s, x, y + Inches(0.08), side_w, Inches(0.22),
             [(label.upper(), dict(size=9.5, bold=True, color=GREEN_600, spacing=1.4))],
             align=PP_ALIGN.LEFT if side == 'left' else PP_ALIGN.RIGHT)
        inner_x = x + side_w if side == 'left' else x          # the edge that faces the center
        edge = cx if side == 'left' else cx + center_w
        sy = top + (stack_h - (len(items) * card_h + (len(items) - 1) * cgap)) // 2
        for j, (ic, name, sub, when) in enumerate(items):
            ty = sy + j * (card_h + cgap)
            _round(rect(s, x, ty, side_w, card_h, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE))
            rect(s, x + Inches(0.16), ty + Inches(0.16), Inches(0.44), Inches(0.44), fill=MINT,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            _mark(s, ic or 'add_circle', x + Inches(0.2), ty + Inches(0.2), Inches(0.36))
            tw = text_w(when)
            nh = block_h(name, tw / 914400, 12.5, spacing=1.15, bold=True)
            sh = block_h(sub, tw / 914400, 10.5, spacing=1.2) if sub else 0
            ty0 = ty + (card_h - nh - (Inches(0.04) + sh if sub else 0)) // 2
            text(s, x + Inches(0.76), ty0, tw, nh, [(name, dict(size=12.5, bold=True, color=INK))], line_spacing=1.15)
            if sub:
                text(s, x + Inches(0.76), ty0 + nh + Inches(0.04), tw, sh,
                     [(sub, dict(size=10.5, color=INK_MUTED))], line_spacing=1.2)
            if when:
                bg, ink = tone[when]
                pw = pill_w(when)
                px = x + side_w - pw - Inches(0.16)
                _round(rect(s, px, ty + (card_h - Inches(0.28)) // 2, pw, Inches(0.28), fill=bg,
                            shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.14))
                text(s, px, ty + (card_h - Inches(0.28)) // 2 + Inches(0.06), pw, Inches(0.18),
                     [(when.upper(), dict(size=8.5, bold=True, color=ink, spacing=0.8))], align=PP_ALIGN.CENTER)
            cy = ty + card_h // 2
            _arrow(_line(s, edge, cy, inner_x, cy, GREEN_300, w=Pt(1.5)))
    return top + stack_h


_MAP_GROUP = {'live': (FOREST, WHITE), 'plain': (MIST, INK_MUTED)}


def processmap(s, y, stages, groups=(), *, loop=None):
    """A whole business on one page: numbered stops along a spine, each with what Servantium
       covers at that step (live today, what we build, what comes later), the tools it
       replaces, optional bands that group the steps, and a loop from the last step back to
       the first. The deck version of the "one-person revenue stack" map.
       stages: dicts (title, text, live, build, later, replaces); groups: (label, from, to, tone)."""
    n = len(stages)
    gap = Inches(0.1)
    cw = (BODY_W - gap * (n - 1)) // n
    xs = [MARGIN + i * (cw + gap) for i in range(n)]
    yy = y
    if groups:
        for label, a, b2, tone in groups:
            x0, x1 = xs[a - 1], xs[b2 - 1] + cw
            fill_, ink = _MAP_GROUP.get(tone, (MINT, FOREST))
            _round(rect(s, x0, yy, x1 - x0, Inches(0.28), fill=fill_, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.06))
            text(s, x0, yy + Inches(0.06), x1 - x0, Inches(0.18),
                 [(label.upper(), dict(size=8.5, bold=True, color=ink, spacing=1.2))], align=PP_ALIGN.CENTER)
        yy += Inches(0.42)
    stop = Inches(0.34)
    spine_y = yy + stop // 2
    _line(s, xs[0] + cw // 2, spine_y, xs[-1] + cw // 2, spine_y, GREEN, w=Pt(4))
    inner = cw - Inches(0.12)
    w_in = inner / 914400
    title_h = max(block_h(st['title'], w_in, 14, spacing=1.1, font=DISPLAY) for st in stages)
    text_h = max((block_h(st.get('text') or '', w_in, 9.5, spacing=1.2) if st.get('text') else 0) for st in stages)
    for i, st in enumerate(stages):
        cx = xs[i] + cw // 2
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - stop // 2, yy, stop, stop)
        dot.fill.solid(); dot.fill.fore_color.rgb = FOREST
        dot.line.color.rgb = GREEN_300; dot.line.width = Pt(2); dot.shadow.inherit = False
        text(s, cx - stop // 2, yy + Inches(0.07), stop, Inches(0.2),
             [(str(i + 1), dict(size=11, bold=True, color=WHITE))], align=PP_ALIGN.CENTER)
        text(s, xs[i] + Inches(0.06), yy + stop + Inches(0.08), inner, title_h,
             [(st['title'], dict(font=DISPLAY, size=14, color=FOREST))], align=PP_ALIGN.CENTER, line_spacing=1.1)
        if st.get('text'):
            text(s, xs[i] + Inches(0.06), yy + stop + Inches(0.1) + title_h, inner, text_h,
                 [(st['text'], dict(size=9.5, color=INK_MUTED))], align=PP_ALIGN.CENTER, line_spacing=1.2)
    chips_y = yy + stop + Inches(0.18) + title_h + text_h + Inches(0.06)
    chip_gap = Inches(0.07)
    bottoms = []
    any_state = set()
    for i, st in enumerate(stages):
        cy = chips_y
        items = [(t, 'live') for t in st.get('live', [])] + [(t, 'build') for t in st.get('build', [])] + \
                [(t, 'later') for t in st.get('later', [])]
        for t, state in items:
            any_state.add(state)
            th = block_h(t, (cw - Inches(0.2)) / 914400, 9.5, spacing=1.15, bold=state == 'live')
            ch = max(Inches(0.32), th + Inches(0.14))
            if state == 'live':
                box = rect(s, xs[i], cy, cw, ch, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE); ink = WHITE
            elif state == 'build':
                box = rect(s, xs[i], cy, cw, ch, fill=WHITE, line=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE); ink = INK
            else:
                box = rect(s, xs[i], cy, cw, ch, fill=MIST, line=INK_MUTED, shape=MSO_SHAPE.ROUNDED_RECTANGLE); ink = INK_MUTED
                box.line.dash_style = 4
            _round(box, Inches(0.07))
            text(s, xs[i] + Inches(0.1), cy + (ch - th) // 2, cw - Inches(0.2), th,
                 [(t, dict(size=9.5, bold=state == 'live', color=ink))], align=PP_ALIGN.CENTER, line_spacing=1.15)
            cy += ch + chip_gap
        bottoms.append(cy - chip_gap)
    body_bottom = max(bottoms)
    if any(st.get('replaces') for st in stages):
        ly = body_bottom + Inches(0.12)
        logo = Inches(0.26)
        for i, st in enumerate(stages):
            logos = st.get('replaces') or []
            if not logos:
                continue
            text(s, xs[i], ly, cw, Inches(0.16), [('REPLACES', dict(size=7.5, bold=True, color=INK_MUTED, spacing=1.0))],
                 align=PP_ALIGN.CENTER)
            row_w = len(logos) * logo + (len(logos) - 1) * Inches(0.08)
            lx = xs[i] + (cw - row_w) // 2
            for p in logos:
                _mark(s, p, lx, ly + Inches(0.2), logo)
                lx += logo + Inches(0.08)
        body_bottom = ly + Inches(0.2) + logo
    if loop:
        ly = body_bottom + Inches(0.26)
        x_last, x_first = xs[-1] + cw // 2, xs[0] + cw // 2
        _wire(s, [(x_last, bottoms[-1] if not stages[-1].get('replaces') else body_bottom), (x_last, ly),
                  (x_first, ly), (x_first, bottoms[0] if not stages[0].get('replaces') else body_bottom)],
              GREEN_600, dashed=True, w=Pt(1.5))
        lw = Inches(text_w_in(loop, 10, bold=True) + 0.3)
        rect(s, MARGIN + (BODY_W - lw) // 2, ly - Inches(0.12), lw, Inches(0.24), fill=WHITE)
        text(s, MARGIN + (BODY_W - lw) // 2, ly - Inches(0.1), lw, Inches(0.2),
             [(loop, dict(size=10, bold=True, color=GREEN_600))], align=PP_ALIGN.CENTER)
        body_bottom = ly + Inches(0.12)
    if len(any_state) > 1:
        gy = body_bottom + Inches(0.14)
        gx = MARGIN
        for state, label in (('live', 'Live today'), ('build', 'We build'), ('later', 'Later')):
            if state not in any_state:
                continue
            sw = rect(s, gx, gy + Inches(0.03), Inches(0.3), Inches(0.16),
                      fill=FOREST if state == 'live' else (WHITE if state == 'build' else MIST),
                      line=None if state == 'live' else (GREEN if state == 'build' else INK_MUTED),
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            if state == 'later':
                sw.line.dash_style = 4
            text(s, gx + Inches(0.38), gy, Inches(1.2), Inches(0.2), [(label, dict(size=9.5, color=INK_MUTED))])
            gx += Inches(0.38) + Inches(text_w_in(label, 9.5) + 0.3)
        body_bottom = gy + Inches(0.22)
    return body_bottom


def board(s, y, columns, *, sources=None, sources_label='Sources'):
    """A pipeline board: where every signal lands and how it moves. An optional row of
       sources feeds the first column; each column is a stage with example cards (title,
       meta, score, tag). columns: (title, [(title, meta, score, tag)])."""
    yy = y
    n = len(columns)
    gap = Inches(0.12)
    cw = (BODY_W - gap * (n - 1)) // n
    if sources:
        text(s, MARGIN, yy + Inches(0.06), Inches(1.0), Inches(0.2),
             [(sources_label.upper(), dict(size=9, bold=True, color=GREEN_600, spacing=1.3))])
        sx = MARGIN + Inches(1.0)
        centers = []
        for src in sources:
            w = Inches(text_w_in(src, 10.5, bold=True) + 0.36)
            _round(rect(s, sx, yy, w, Inches(0.32), fill=MINT_TINT, line=GREEN_300, shape=MSO_SHAPE.ROUNDED_RECTANGLE),
                   Inches(0.16))
            text(s, sx, yy + Inches(0.07), w, Inches(0.2), [(src, dict(size=10.5, bold=True, color=FOREST))],
                 align=PP_ALIGN.CENTER)
            centers.append(sx + w // 2)
            sx += w + Inches(0.1)
        # every source joins one rail, and the rail feeds the first column
        rail = yy + Inches(0.48)
        target = MARGIN + cw // 2
        for c in centers:
            _line(s, c, yy + Inches(0.32), c, rail, GREEN_300, w=Pt(1.25))
        _line(s, min(centers + [target]), rail, max(centers + [target]), rail, GREEN_300, w=Pt(1.25))
        _arrow(_line(s, target, rail, target, yy + Inches(0.72), GREEN_600, w=Pt(1.5)))
        yy += Inches(0.72)
    head_h = Inches(0.42)
    card_w = cw - Inches(0.2)
    iw = (card_w - Inches(0.24)) / 914400

    def card_h(c):
        t, meta, score, tag = c
        return (Inches(0.12) + block_h(t, iw, 11, spacing=1.15, bold=True)
                + (Inches(0.04) + block_h(meta, iw, 9.5, spacing=1.2) if meta else 0)
                + (Inches(0.32) if (score or tag) else 0) + Inches(0.12))
    col_h = head_h + max([sum(card_h(c) + Inches(0.1) for c in cards) for _, cards in columns] + [Inches(0.6)]) + Inches(0.06)
    for i, (title, cards) in enumerate(columns):
        x = MARGIN + i * (cw + gap)
        _round(rect(s, x, yy, cw, col_h, fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.1))
        text(s, x + Inches(0.14), yy + Inches(0.12), cw - Inches(0.6), Inches(0.22),
             [(title, dict(size=12, bold=True, color=FOREST))])
        text(s, x + cw - Inches(0.44), yy + Inches(0.13), Inches(0.3), Inches(0.2),
             [(str(len(cards)), dict(size=10.5, color=INK_MUTED))], align=PP_ALIGN.RIGHT)
        if i < n - 1:
            _arrow(_line(s, x + cw, yy + Inches(0.23), x + cw + gap, yy + Inches(0.23), INK_MUTED, w=Pt(1)))
        cy = yy + head_h
        for c in cards:
            t, meta, score, tag = c
            h = card_h(c)
            card = rect(s, x + Inches(0.1), cy, card_w, h, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            _round(card, Inches(0.08))
            ty = cy + Inches(0.12)
            th = block_h(t, iw, 11, spacing=1.15, bold=True)
            text(s, x + Inches(0.22), ty, card_w - Inches(0.24), th, [(t, dict(size=11, bold=True, color=INK))],
                 line_spacing=1.15)
            ty += th
            if meta:
                mh = block_h(meta, iw, 9.5, spacing=1.2)
                text(s, x + Inches(0.22), ty + Inches(0.04), card_w - Inches(0.24), mh,
                     [(meta, dict(size=9.5, color=INK_MUTED))], line_spacing=1.2)
                ty += Inches(0.04) + mh
            px = x + Inches(0.22)
            for val, dark in ((score, True), (tag, False)):
                if not val:
                    continue
                pw = Inches(text_w_in(val, 9, bold=True) + 0.24)
                _round(rect(s, px, ty + Inches(0.08), pw, Inches(0.22), fill=FOREST if dark else MINT,
                            shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.11))
                text(s, px, ty + Inches(0.115), pw, Inches(0.16),
                     [(val, dict(size=9, bold=True, color=WHITE if dark else FOREST))], align=PP_ALIGN.CENTER)
                px += pw + Inches(0.06)
            cy += h + Inches(0.1)
    return yy + col_h


def _balanced_w(txt, width_in, size, *, font=None):
    """The narrowest width that keeps the same number of lines, so a wrapped line is not
       left with one word on it ("the first / sprint.")."""
    kw = {'font': font} if font else {}
    lines = est_lines(txt, width_in, size, **kw)
    if lines < 2:
        return width_in
    lo, hi = width_in * 0.45, width_in
    for _ in range(14):
        mid = (lo + hi) / 2
        if est_lines(txt, mid, size, **kw) > lines:
            lo = mid
        else:
            hi = mid
    return hi + 0.04


def questions(s, y, items, *, note=None, panel_w=Inches(2.9), min_h=Inches(2.7)):
    """A pause for the room: a forest panel with a large question mark, then two to four
       questions, each with the reason we ask. These are the spots in a pitch where the
       client does the talking. items: (question, why)."""
    gap = Inches(0.42)
    qx = MARGIN + panel_w + gap
    qw = BODY_W - panel_w - gap
    num_w = Inches(0.66)
    tw = qw - num_w
    rows = []
    for q, why in items:
        qh = block_h(q, tw / 914400, 19, spacing=1.12, font=DISPLAY)
        wh = block_h(why, tw / 914400, 12, spacing=1.25) if why else 0
        rows.append((qh, wh, max(Inches(0.9), Inches(0.26) + qh + (Inches(0.08) + wh if why else 0) + Inches(0.26))))
    total = sum(r[2] for r in rows)
    h = max(total, min_h)
    _round(rect(s, MARGIN, y, panel_w, h, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.14))
    note_w = _balanced_w(note, (panel_w - Inches(0.6)) / 914400, 14) if note else 0
    nh = block_h(note, note_w, 14, spacing=1.25) if note else 0
    note_area = nh + Inches(0.5) if note else 0
    mark_h = Inches(1.6)            # the glyph's box at 104 pt; centered in the space above the note
    text(s, MARGIN, y + max(Inches(0.1), (h - note_area - mark_h) // 2), panel_w, mark_h,
         [('?', dict(font=DISPLAY, size=104, color=GREEN_300))], align=PP_ALIGN.CENTER)
    if note:
        nw = Inches(note_w)
        text(s, MARGIN + (panel_w - nw) // 2, y + h - nh - Inches(0.34), nw, nh,
             [(note, dict(size=14, color=WHITE))], align=PP_ALIGN.CENTER, line_spacing=1.25)
    yy = y + (h - total) // 2
    for i, ((q, why), (qh, wh, rh)) in enumerate(zip(items, rows)):
        if i:
            rect(s, qx, yy, qw, Emu(9525), fill=CLOUD)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, qx, yy + Inches(0.26), Inches(0.42), Inches(0.42))
        dot.fill.solid(); dot.fill.fore_color.rgb = MINT
        dot.line.fill.background(); dot.shadow.inherit = False
        text(s, qx, yy + Inches(0.35), Inches(0.42), Inches(0.24),
             [(str(i + 1), dict(size=13, bold=True, color=FOREST))], align=PP_ALIGN.CENTER)
        text(s, qx + num_w, yy + Inches(0.26), tw, qh, [(q, dict(font=DISPLAY, size=19, color=FOREST))],
             line_spacing=1.12)
        if why:
            text(s, qx + num_w, yy + Inches(0.26) + qh + Inches(0.08), tw, wh,
                 [(why, dict(size=12, color=INK_MUTED))], line_spacing=1.25)
        yy += rh
    return y + h


def _trim_logo(path, workdir='.'):
    """A copy of the logo cropped to its ink: transparent or near-white margins inside the
       file made some marks render far smaller than others at the same size."""
    from PIL import Image, ImageChops
    im = Image.open(path).convert('RGBA')
    box = im.getchannel('A').getbbox()
    if box == (0, 0, im.width, im.height) or box is None:
        bg = Image.new('RGB', im.size, (255, 255, 255))
        diff = ImageChops.difference(im.convert('RGB'), bg).convert('L').point(lambda v: 255 if v > 24 else 0)
        box = diff.getbbox() or box
    if not box or box == (0, 0, im.width, im.height):
        return str(path)
    out = Path(workdir) / f'logo-{Path(path).stem}-trim.png'
    out.parent.mkdir(parents=True, exist_ok=True)
    im.crop(box).save(out)
    return str(out)


def _fit_logo(path, area_in2=0.30, max_h_in=0.34, max_w_in=1.15):
    """Logo size in inches by visual area, so wordmarks and square marks weigh the same."""
    from PIL import Image
    im = Image.open(path)
    aspect = im.width / im.height
    h_in = min(max_h_in, (area_in2 / aspect) ** 0.5)
    w_in = h_in * aspect
    if w_in > max_w_in:
        w_in, h_in = max_w_in, max_w_in / aspect
    return w_in, h_in


def team_slide(prs, *, eyebrow_text, title_, points, credentials, people, credentials_label=None,
               workdir='.', proof=None, stat=None, logos=None, logos_label=None):
    """Who we are, as Servantium and as people. Left: the title, how we work and where
       we've done it. Right: a forest panel with each person's cut-out photo rising out of
       a jade circle, their name and role, and the logos of where they worked before on
       white chips. Photos must be PNG cut-outs (transparent background)."""
    from PIL import Image, ImageDraw, ImageFilter
    from .kit import blank, W as SW, H as SH
    s = blank(prs)
    panel_x = int(SW * 0.5)
    pw_in, ph_in = (SW - panel_x) / 914400, SH / 914400
    DPI = 220
    PW, PH = int(pw_in * DPI), int(ph_in * DPI)
    rgb = lambda c: tuple(int(str(c)[i:i + 2], 16) for i in (0, 2, 4))
    bg = Image.new('RGB', (PW, PH), rgb(FOREST))
    glow = Image.new('L', (PW, PH), 0)
    ImageDraw.Draw(glow).ellipse((-PW * 0.4, -PH * 0.5, PW * 0.9, PH * 0.6), fill=255)
    glow = glow.filter(ImageFilter.GaussianBlur(PW // 6))
    bg = Image.composite(Image.new('RGB', (PW, PH), rgb(FOREST_900)), bg, glow)
    k = len(people)
    slot = PW / k
    circ = int(min(slot * 0.74, PH * 0.4))
    cy = int(PH * 0.43)
    canvas = bg.convert('RGBA')
    for i, p in enumerate(people):
        cx = int(slot * (i + 0.5))
        ring = Image.new('RGBA', (PW, PH), (0, 0, 0, 0))
        ImageDraw.Draw(ring).ellipse((cx - circ // 2, cy - circ // 2, cx + circ // 2, cy + circ // 2),
                                     fill=rgb(GREEN_300) + (72,))
        canvas = Image.alpha_composite(canvas, ring)
        ph = Image.open(p['photo']).convert('RGBA')
        bbox = ph.getchannel('A').getbbox() or (0, 0, ph.width, ph.height)
        ph = ph.crop(bbox)
        scale = circ * 1.04 / ph.width
        ph = ph.resize((int(ph.width * scale), int(ph.height * scale)), Image.LANCZOS)
        # sit the cut-out in the circle: clipped to the circle below its middle, free above it
        px0, py0 = cx - ph.width // 2, cy + circ // 2 - ph.height
        mask = Image.new('L', ph.size, 0)
        d = ImageDraw.Draw(mask)
        d.rectangle((0, 0, ph.width, cy - py0), fill=255)
        d.ellipse((cx - circ // 2 - px0, cy - circ // 2 - py0, cx + circ // 2 - px0, cy + circ // 2 - py0), fill=255)
        from PIL import ImageChops
        ph.putalpha(ImageChops.multiply(ph.getchannel('A'), mask))
        layer = Image.new('RGBA', (PW, PH), (0, 0, 0, 0))
        layer.paste(ph, (px0, py0), ph)
        canvas = Image.alpha_composite(canvas, layer)
    out = Path(workdir) / f'team-{abs(hash(tuple(p["photo"] for p in people))) % 10**8}.png'
    out.parent.mkdir(parents=True, exist_ok=True)
    canvas.convert('RGB').save(out)
    s.shapes.add_picture(str(out), panel_x, 0, width=SW - panel_x, height=SH)
    # names, roles and logos, as editable text on the panel
    base_y = Emu(int((cy + circ / 2) / DPI * 914400)) + Inches(0.28)
    slot_w = (SW - panel_x) // k
    for i, p in enumerate(people):
        x = panel_x + i * slot_w
        text(s, x, base_y, slot_w, Inches(0.42), [(p['name'], dict(font=DISPLAY, size=21, color=WHITE))],
             align=PP_ALIGN.CENTER)
        text(s, x, base_y + Inches(0.48), slot_w, Inches(0.24),
             [(p['role'].upper(), dict(size=10, bold=True, color=GREEN_300, spacing=1.4))], align=PP_ALIGN.CENTER)
        plogos = [] if logos else (p.get('logos') or [])
        if plogos:
            logos_ = plogos
            chip_h, cgap = Inches(0.48), Inches(0.12)
            logos_ = [_trim_logo(l, workdir) for l in logos_]
            sizes = [_fit_logo(l) for l in logos_]
            chip_ws = [Inches(max(1.0, w + 0.36)) for w, _ in sizes]
            total = sum(chip_ws) + cgap * (len(logos_) - 1)
            lx = x + (slot_w - total) // 2
            ly = base_y + Inches(0.92)
            for l, (w_in, h_in), cwid in zip(logos_, sizes, chip_ws):
                chip = rect(s, lx, ly, cwid, chip_h, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                _round(chip, Inches(0.1))
                s.shapes.add_picture(str(l), lx + (cwid - Inches(w_in)) // 2, ly + (chip_h - Inches(h_in)) // 2,
                                     width=Inches(w_in), height=Inches(h_in))
                lx += cwid + cgap
            text(s, x, ly + chip_h + Inches(0.1), slot_w, Inches(0.22),
                 [(f'BEFORE {BRAND_SHORT.upper()}', dict(size=8.5, bold=True, color=ON_FOREST_MUTED, spacing=1.3))],
                 align=PP_ALIGN.CENTER)
    panel_w = SW - panel_x
    if logos:
        # one strip for the team: every company once, sized by visual weight
        trimmed = [_trim_logo(l, workdir) for l in logos]
        sizes = [_fit_logo(l, area_in2=0.2, max_h_in=0.3, max_w_in=0.95) for l in trimmed]
        chip_h, cgap = Inches(0.46), Inches(0.1)
        chip_ws = [Inches(w + 0.26) for w, _ in sizes]
        total = sum(chip_ws) + cgap * (len(trimmed) - 1)
        lx = panel_x + (panel_w - total) // 2
        ly = base_y + Inches(0.98)
        for l, (w_in, h_in), cwid in zip(trimmed, sizes, chip_ws):
            chip = rect(s, lx, ly, cwid, chip_h, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            _round(chip, Inches(0.1))
            s.shapes.add_picture(str(l), lx + (cwid - Inches(w_in)) // 2, ly + (chip_h - Inches(h_in)) // 2,
                                 width=Inches(w_in), height=Inches(h_in))
            lx += cwid + cgap
        text(s, panel_x, ly + chip_h + Inches(0.12), panel_w, Inches(0.22),
             [((logos_label or 'Where we have worked').upper(), dict(size=8.5, bold=True, color=ON_FOREST_MUTED, spacing=1.3))],
             align=PP_ALIGN.CENTER)
    if stat:
        value, label = stat
        text(s, panel_x, Inches(0.42), panel_w, Inches(0.7),
             [(value, dict(size=44, bold=True, color=WHITE))], align=PP_ALIGN.CENTER)
        text(s, panel_x, Inches(1.16), panel_w, Inches(0.24),
             [(label.upper(), dict(size=10, bold=True, color=GREEN_300, spacing=1.5))], align=PP_ALIGN.CENTER)
    # left: eyebrow, title, how we work, where we've done it
    lw = panel_x - MARGIN - Inches(0.5)
    text(s, MARGIN, Inches(0.62), lw, Inches(0.26),
         [(eyebrow_text.upper(), dict(size=9.5, bold=True, color=GREEN_600, spacing=1.6))])
    th = block_h(title_, lw / 914400, 30, spacing=1.08, font=DISPLAY)
    text(s, MARGIN, Inches(0.96), lw, th, [(title_, dict(font=DISPLAY, size=30, color=FOREST))], line_spacing=1.06)
    yy = Inches(0.96) + th + Inches(0.42)
    if proof:
        # compact: how we work in one row, then where we've done it as a 2 x 2 grid
        n = max(1, len(points))
        cgap = Inches(0.28)
        cw_ = (lw - cgap * (n - 1)) // n
        tallest = 0
        for i, pt in enumerate(points):
            x = MARGIN + i * (cw_ + cgap)
            icon(s, pt.get('icon') or 'check_circle', x, yy, Inches(0.34))
            text(s, x, yy + Inches(0.44), cw_, Inches(0.26), [(pt['lead'], dict(size=13, bold=True, color=INK))])
            h2 = block_h(pt['text'], cw_ / 914400, 11)
            text(s, x, yy + Inches(0.74), cw_, h2, [(pt['text'], dict(size=11, color=INK_MUTED))], line_spacing=1.2)
            tallest = max(tallest, Inches(0.74) + h2)
        yy += tallest + Inches(0.36)
        rect(s, MARGIN, yy, lw, Emu(9525), fill=CLOUD)
        text(s, MARGIN, yy + Inches(0.16), lw, Inches(0.22),
             [((credentials_label or 'Where we have done this').upper(), dict(size=9, bold=True, color=INK_MUTED, spacing=1.3))])
        gy = yy + Inches(0.5)
        gw = (lw - Inches(0.3)) // 2
        row_h = max([Inches(0.8)] + [Inches(0.34) + block_h(p_['text'], gw / 914400, 11) + Inches(0.12) for p_ in proof])
        for i, p_ in enumerate(proof):
            x = MARGIN + (i % 2) * (gw + Inches(0.3))
            y_ = gy + (i // 2) * row_h
            rect(s, x, y_ + Inches(0.03), Inches(0.05), row_h - Inches(0.2), fill=GREEN)
            text(s, x + Inches(0.18), y_, gw - Inches(0.18), Inches(0.28),
                 [(p_['lead'], dict(size=13.5, bold=True, color=FOREST))])
            text(s, x + Inches(0.18), y_ + Inches(0.32), gw - Inches(0.18), block_h(p_['text'], (gw - Inches(0.18)) / 914400, 11),
                 [(p_['text'], dict(size=11, color=INK_800))], line_spacing=1.2)
        points = []           # drawn above
        credentials = None    # the grid replaces the one-line version
    for pt in points:
        icon(s, pt.get('icon') or 'check_circle', MARGIN, yy + Inches(0.02), Inches(0.4))
        text(s, MARGIN + Inches(0.6), yy, lw - Inches(0.6), Inches(0.28),
             [(pt['lead'], dict(size=14, bold=True, color=INK))])
        th2 = block_h(pt['text'], (lw - Inches(0.6)) / 914400, 12)
        text(s, MARGIN + Inches(0.6), yy + Inches(0.32), lw - Inches(0.6), th2,
             [(pt['text'], dict(size=12, color=INK_MUTED))], line_spacing=1.22)
        yy += Inches(0.32) + th2 + Inches(0.3)
    if credentials:
        cy_ = FOOTER_Y - Inches(1.28)
        rect(s, MARGIN, cy_, lw, Emu(9525), fill=CLOUD)
        text(s, MARGIN, cy_ + Inches(0.16), lw, Inches(0.22),
             [((credentials_label or 'Where we have done this').upper(), dict(size=9, bold=True, color=INK_MUTED, spacing=1.3))])
        text(s, MARGIN, cy_ + Inches(0.44), lw, block_h(credentials, lw / 914400, 12.5),
             [(credentials, dict(size=12.5, color=INK_800))], line_spacing=1.25)
    return s


# ── proposal and plan compositions (2026-10-20) ──────────────────────────────
# Ten shapes added so a 25-slide proposal never repeats one: an inventory of use cases
# (catalog), choices on an axis (spectrum), doors (options), a comparison grid (matrix),
# spend against a budget cap (cap), a cell-level plan (weekgrid), lanes by period
# (roadmap), price tiers (tiers), questions and answers (faq) and the math behind a price
# (ledger). Coordinates go through _r and _t so every EMU is a whole number.
from pptx.enum.dml import MSO_LINE_DASH_STYLE as _DASH

_EMU = 914400
_GRID = {**_BAR, 'cloud': (CLOUD, INK_800), 'teal': (TEAL, WHITE), 'coral': (CORAL, WHITE)}


def _in(v):
    return v / _EMU


def _r(s, x, y, w, h, **kw):
    return rect(s, int(x), int(y), int(max(w, 1)), int(max(h, 1)), **kw)


def _t(s, x, y, w, h, runs, **kw):
    return text(s, int(x), int(y), int(max(w, 1)), int(max(h, 1)), runs, **kw)


def _pill(s, x, y, label, tone='forest', *, size=8.5, h=Inches(0.24), align='left'):
    """A rounded label; returns its width. align says which edge x is: left, center, right."""
    bg, ink = _PILL.get(tone, _PILL['forest'])
    w = Inches(text_w_in(label.upper(), size, bold=True, ) + 0.28)
    px = x - w if align == 'right' else (x - w // 2 if align == 'center' else x)
    _round(_r(s, px, y, w, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE), h // 2)
    th = Inches(size * 1.25 / 72)
    _t(s, px, y + (h - th) // 2, w, th, [(label.upper(), dict(size=size, bold=True, color=ink, spacing=0.8))],
       align=PP_ALIGN.CENTER)
    return w


def catalog(s, y, groups, *, gap=Inches(0.12)):
    """An inventory: one column per group (a workstream, a team), a dark header with its
       status pill, and the items found inside it as chips.
       groups: (title, tag, tone, [items])."""
    n = len(groups)
    cw = (BODY_W - gap * (n - 1)) // n
    tw = cw - Inches(0.28)
    iw = cw - Inches(0.42)

    def head_h(g):
        return (Inches(0.16) + block_h(g[0], _in(tw), 12.5, spacing=1.12, bold=True)
                + (Inches(0.4) if g[1] else Inches(0.06)) + Inches(0.1))

    def item_h(t):
        return max(Inches(0.38), block_h(t, _in(iw), 10.5, spacing=1.18) + Inches(0.16))

    hh = max(head_h(g) for g in groups)
    body = max(sum(item_h(t) + Inches(0.08) for t in g[3]) for g in groups)
    col_h = hh + Inches(0.1) + body + Inches(0.04)
    for i, (title_, tag, tone, items) in enumerate(groups):
        x = MARGIN + i * (cw + gap)
        _round(_r(s, x, y, cw, col_h, fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.1))
        _round(_r(s, x, y, cw, hh, fill=FOREST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.1))
        th = block_h(title_, _in(tw), 12.5, spacing=1.12, bold=True)
        _t(s, x + Inches(0.14), y + Inches(0.16), tw, th, [(title_, dict(size=12.5, bold=True, color=WHITE))],
           line_spacing=1.12)
        if tag:
            _pill(s, x + Inches(0.14), y + hh - Inches(0.36), tag, 'mint' if tone in (None, 'forest') else tone)
        iy = y + hh + Inches(0.1)
        for t in items:
            h = item_h(t)
            _round(_r(s, x + Inches(0.08), iy, cw - Inches(0.16), h, fill=WHITE, line=CLOUD,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.06))
            _r(s, x + Inches(0.15), iy + Inches(0.1), Inches(0.04), h - Inches(0.2), fill=GREEN)
            th = block_h(t, _in(iw), 10.5, spacing=1.18)
            _t(s, x + Inches(0.27), iy + (h - th) // 2, iw, th, [(t, dict(size=10.5, color=INK))], line_spacing=1.18)
            iy += h + Inches(0.08)
    return y + col_h


def spectrum(s, y, ends, stops, *, pick_label='Our approach'):
    """Choices laid on one axis between two extremes, with the one we recommend lit.
       ends: (left label, right label); stops: (title, text, pick)."""
    n = len(stops)
    lab = dict(size=9.5, bold=True, color=INK_MUTED, spacing=1.3)
    _t(s, MARGIN, y, Inches(4), Inches(0.22), [(ends[0].upper(), lab)])
    _t(s, MARGIN + BODY_W - Inches(4), y, Inches(4), Inches(0.22), [(ends[1].upper(), lab)], align=PP_ALIGN.RIGHT)
    ax = y + Inches(0.52)
    mid = MARGIN + BODY_W // 2
    _arrow(_line(s, mid, ax, MARGIN, ax, INK_MUTED, w=Pt(1.75)))
    _arrow(_line(s, mid, ax, MARGIN + BODY_W, ax, INK_MUTED, w=Pt(1.75)))
    slot = BODY_W // n
    cw = slot - Inches(0.4)
    tw = cw - Inches(0.56)
    tag_h = Inches(0.38)

    def need(st):
        t, body, _ = st
        return (Inches(0.3) + tag_h + block_h(t, _in(tw), 21, spacing=1.08, font=DISPLAY)
                + (Inches(0.12) + block_h(body, _in(tw), 12.5, spacing=1.28) if body else 0) + Inches(0.32))
    ch = max(need(st) for st in stops)
    top = ax + Inches(0.5)
    for i, (t, body, pick) in enumerate(stops):
        cx = MARGIN + i * slot + slot // 2
        x = cx - cw // 2
        d = Inches(0.34) if pick else Inches(0.24)
        _line(s, cx, ax, cx, top, GREEN if pick else CLOUD, w=Pt(1.5))
        dot = _r(s, cx - d // 2, ax - d // 2, d, d, fill=GREEN if pick else WHITE,
                 line=WHITE if pick else INK_MUTED, line_w=Pt(2 if pick else 1.25), shape=MSO_SHAPE.OVAL)
        _round(_r(s, x, top, cw, ch, fill=FOREST if pick else WHITE, line=None if pick else CLOUD,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.14))
        ty = top + Inches(0.3)
        if pick:
            _pill(s, x + Inches(0.28), ty, pick_label, 'green')
        ty += tag_h
        th = block_h(t, _in(tw), 21, spacing=1.08, font=DISPLAY)
        _t(s, x + Inches(0.28), ty, tw, th, [(t, dict(font=DISPLAY, size=21, color=WHITE if pick else FOREST))],
           line_spacing=1.06)
        if body:
            bh = block_h(body, _in(tw), 12.5, spacing=1.28)
            _t(s, x + Inches(0.28), ty + th + Inches(0.12), tw, bh,
               [(body, dict(size=12.5, color=MIST_INK if pick else INK_800))], line_spacing=1.28)
    return top + ch


def options(s, y, items, *, pick_label='Recommended', door_w=Inches(3.2), gap=Inches(0.55)):
    """The ways in, drawn as doors: label, name, length and what each one gets you. The one
       we recommend is lit. items: (label, title, meta, [lines], pick)."""
    n = len(items)
    dw = min(door_w, (BODY_W - gap * (n - 1)) // n)
    total = n * dw + (n - 1) * gap
    x0 = MARGIN + (BODY_W - total) // 2
    pad = Inches(0.36)
    tw = dw - 2 * pad
    arch = int(dw * 0.34)                      # the radius of the rounded top
    lw = tw - Inches(0.32)

    def content(it):
        _, t, meta, lines, _ = it
        h = block_h(t, _in(tw), 24, spacing=1.08, font=DISPLAY) + Inches(0.14)
        h += Inches(0.44) if meta else 0
        h += sum(block_h(ln, _in(lw), 11.5, spacing=1.22) + Inches(0.12) for ln in lines)
        return h
    body_top = arch + Inches(0.12)             # the label sits inside the arch, the rest below it
    dh = body_top + max(content(it) for it in items) + Inches(0.3)
    top = y + Inches(0.16)
    for i, (label, t, meta, lines, pick) in enumerate(items):
        x = x0 + i * (dw + gap)
        door = _r(s, x, top, dw, dh, fill=FOREST if pick else WHITE, line=None if pick else GREEN_300,
                  line_w=Pt(1.5), shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        door.adjustments[0] = min(0.5, arch / max(1, min(dw, dh)))
        door.adjustments[1] = 0.0
        if pick:
            _pill(s, x + dw // 2, y, pick_label, 'green', align='center')
        if label:
            _t(s, x, top + arch - Inches(0.36), dw, Inches(0.22),
               [(label.upper(), dict(size=10, bold=True, color=GREEN_300 if pick else GREEN_600, spacing=1.6))],
               align=PP_ALIGN.CENTER)
        yy = top + body_top
        th = block_h(t, _in(tw), 24, spacing=1.08, font=DISPLAY)
        _t(s, x + pad, yy, tw, th, [(t, dict(font=DISPLAY, size=24, color=WHITE if pick else FOREST))],
           line_spacing=1.06)
        yy += th + Inches(0.14)
        if meta:
            _pill(s, x + pad, yy, meta, 'mint' if pick else 'forest', size=9, h=Inches(0.28))
            yy += Inches(0.44)
        for ln in lines:
            lh = block_h(ln, _in(lw), 11.5, spacing=1.22)
            icon(s, 'check', x + pad, yy + Inches(0.01), Inches(0.2), color=GREEN_300 if pick else GREEN)
            _t(s, x + pad + Inches(0.32), yy, lw, lh, [(ln, dict(size=11.5, color=MIST_INK if pick else INK_800))],
               line_spacing=1.22)
            yy += lh + Inches(0.12)
    return top + dh


def matrix(s, y, columns, rows, *, highlight=None, pick_label='Recommended', label_w=Inches(2.35)):
    """Options compared on the same questions: a column per option, a row per question, each
       cell a short answer with an optional four-step meter (0-4). The recommended column is
       outlined. rows: (label, [cells], [levels] or None); highlight is 1-based."""
    n = len(columns)
    cw = (BODY_W - label_w) // n
    tw = cw - Inches(0.4)
    head_h = Inches(0.56)
    hi = highlight - 1 if highlight else None
    top = y + (Inches(0.18) if hi is not None else 0)

    def row_h(r):
        label, cells, levels = r
        ch = max(block_h(c, _in(tw), 11.5, spacing=1.22) for c in cells) + (Inches(0.24) if levels else 0)
        return max(ch, block_h(label, _in(label_w - Inches(0.3)), 12, spacing=1.15, bold=True)) + Inches(0.3)
    hs = [row_h(r) for r in rows]
    total = head_h + sum(hs)
    if hi is not None:
        hx = MARGIN + label_w + hi * cw
        _round(_r(s, hx + Inches(0.04), top - Inches(0.06), cw - Inches(0.08), total + Inches(0.12), fill=MINT_TINT,
                  line=GREEN, line_w=Pt(1.75), shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.12))
        _pill(s, hx + cw // 2, top - Inches(0.18), pick_label, 'green', align='center')
    for i, c in enumerate(columns):
        _t(s, MARGIN + label_w + i * cw + Inches(0.2), top + Inches(0.14), tw, Inches(0.34),
           [(c, dict(font=DISPLAY, size=18, color=FOREST))], align=PP_ALIGN.CENTER)
    _r(s, MARGIN, top + head_h - Emu(12700), BODY_W, Emu(12700), fill=FOREST)
    yy = top + head_h
    for r, (label, cells, levels) in enumerate(rows):
        h = hs[r]
        _t(s, MARGIN, yy + Inches(0.15), label_w - Inches(0.3), h - Inches(0.3),
           [(label, dict(size=12, bold=True, color=FOREST))], line_spacing=1.15)
        for i, c in enumerate(cells):
            x = MARGIN + label_w + i * cw + Inches(0.2)
            cy = yy + Inches(0.15)
            if levels:
                lv = levels[i]
                mw = 4 * Inches(0.22) - Inches(0.04)
                mx = x + (tw - mw) // 2
                for k in range(4):
                    _round(_r(s, mx + k * Inches(0.22), cy + Inches(0.03), Inches(0.18), Inches(0.1),
                              fill=GREEN if k < lv else CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.05))
                cy += Inches(0.24)
            _t(s, x, cy, tw, block_h(c, _in(tw), 11.5, spacing=1.22),
               [(c, dict(size=11.5, color=INK, bold=i == hi))], align=PP_ALIGN.CENTER, line_spacing=1.22)
        if r < len(rows) - 1:
            _r(s, MARGIN, yy + h - Emu(9525), BODY_W, Emu(9525), fill=CLOUD)
        yy += h
    return yy + (Inches(0.06) if hi is not None else 0)


def cap(s, y, left, right, cap_, *, h=Inches(3.7), unknown_share=0.3):
    """Two ways to buy the same work, as spend stacked against a budget cap. The left stack
       grows phase by phase (phases not priced yet drawn dashed with a question mark), the
       right is one block under the cap; each side lists what it means beside the chart.
       left/right: (label, tone 'warn'|'good', [(label, amount, value or None)], [lines]);
       cap_: (value, label)."""
    cap_v, cap_label = cap_
    unk = cap_v * unknown_share
    tot = lambda side: sum(b[2] if b[2] is not None else unk for b in side[2])
    top_v = max(tot(left), tot(right), cap_v) * 1.06
    label_h = Inches(0.46)
    plot_h = h - label_h
    base = y + plot_h
    scale = plot_h / top_v
    sw, sgap = Inches(1.85), Inches(0.55)
    chart_w = 2 * sw + sgap + Inches(0.5)
    cyy = base - int(cap_v * scale)
    for k, (label, tone, blocks, _) in enumerate((left, right)):
        x = MARGIN + Inches(0.25) + k * (sw + sgap)
        by = base
        for bl, amt, val in blocks:
            bh = int((val if val is not None else unk) * scale)
            by -= bh
            if val is None:
                shp = _r(s, x, by + Inches(0.03), sw, bh - Inches(0.06), fill=WHITE, line=INK_MUTED, line_w=Pt(1.25),
                         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                shp.line.dash_style = _DASH.DASH
                _round(shp, Inches(0.06))
                _t(s, x, by + (bh - Inches(0.5)) // 2, sw, Inches(0.5),
                   [[('?  ', dict(font=DISPLAY, size=20, color=INK_MUTED)), (bl, dict(size=10.5, bold=True, color=INK_MUTED))]],
                   align=PP_ALIGN.CENTER)
            else:
                fill = FOREST if tone == 'good' else SLATE
                _round(_r(s, x, by + Inches(0.03), sw, bh - Inches(0.06), fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE),
                       Inches(0.06))
                two = amt and bh > Inches(0.62)
                th = Inches(0.46) if two else Inches(0.22)
                runs = [[(bl, dict(size=11, bold=True, color=WHITE))]] + ([[(amt, dict(size=13, color=GREEN_300 if tone == 'good' else MIST_INK))]] if two else [])
                if amt and not two:
                    runs = [[(bl + '  ', dict(size=11, bold=True, color=WHITE)), (amt, dict(size=11, color=MIST_INK))]]
                _t(s, x, by + (bh - th) // 2, sw, th, runs, align=PP_ALIGN.CENTER)
        _r(s, x - Inches(0.1), base, sw + Inches(0.2), Emu(12700), fill=INK_MUTED)
    _line(s, MARGIN, cyy, MARGIN + chart_w, cyy, GREEN_600, w=Pt(1.75), dash=True)
    _t(s, MARGIN, cyy - Inches(0.28), chart_w, Inches(0.22),
       [(cap_label.upper(), dict(size=9.5, bold=True, color=GREEN_600, spacing=1.2))], align=PP_ALIGN.RIGHT)
    # what each side means, beside the chart; each heading carries its stack's colour as the key
    lx = MARGIN + chart_w + Inches(0.45)
    lw_all = BODY_W - chart_w - Inches(0.45)
    colw = (lw_all - Inches(0.4)) // 2
    tw = colw - Inches(0.86)
    need = lambda lines: Inches(0.86) + sum(block_h(ln, _in(tw), 13, spacing=1.22) + Inches(0.2) for ln in lines) + Inches(0.12)
    ph = max(need(side[3]) for side in (left, right))
    py = y + max(0, (plot_h - ph) // 2)
    for k, (label, tone, _, lines) in enumerate((left, right)):
        x = lx + k * (colw + Inches(0.4))
        good = tone == 'good'
        _round(_r(s, x, py, colw, ph, fill=MINT_TINT if good else MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.12))
        _round(_r(s, x + Inches(0.26), py + Inches(0.3), Inches(0.22), Inches(0.22), fill=FOREST if good else SLATE,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.05))
        _t(s, x + Inches(0.6), py + Inches(0.24), colw - Inches(0.8), Inches(0.32),
           [(label, dict(font=DISPLAY, size=18, color=FOREST))])
        yy = py + Inches(0.86)
        for ln in lines:
            lh = block_h(ln, _in(tw), 13, spacing=1.22)
            icon(s, 'check_circle' if good else 'cancel', x + Inches(0.26), yy, Inches(0.26), color=GREEN if good else CORAL)
            _t(s, x + Inches(0.62), yy, tw, lh, [(ln, dict(size=13, color=INK_800))], line_spacing=1.22)
            yy += lh + Inches(0.2)
    return y + h


def weekgrid(s, y, periods, stages, rows, *, marks=(), label_w=Inches(2.3), max_h=Inches(4.4)):
    """A short plan at cell level: one row per workstream, one square per period, each square
       coloured by the stage it is in and the stage named where it starts. A legend sits on
       top; optional marks (steering, readout) sit on their own line above the grid.
       stages: (key, label, short, tone); rows: (name, note, [(from, to, stage, label or None)]);
       marks: (at, label). A span's own label replaces the stage name (an integration row)."""
    tone = {k: _GRID.get(t or 'green', _GRID['green']) for k, _, _, t in stages}
    names = {k: (label, short) for k, label, short, _ in stages}
    legend_h = Inches(0.32)
    lx = MARGIN + label_w
    for k, label, _, _ in stages:
        bg, _ = tone[k]
        _round(_r(s, lx, y + Inches(0.05), Inches(0.2), Inches(0.2), fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.04))
        lw = Inches(text_w_in(label, 10.5) + 0.06)
        _t(s, lx + Inches(0.28), y + Inches(0.04), lw, Inches(0.22), [(label, dict(size=10.5, color=INK_800))])
        lx += Inches(0.28) + lw + Inches(0.28)
    grid_x = MARGIN + label_w
    cw = (BODY_W - label_w) // len(periods)
    head_y = y + legend_h + Inches(0.08)
    head_h = Inches(0.28)
    marks_h = Inches(0.34) if marks else 0
    row_h = min(Inches(0.7), (max_h - legend_h - Inches(0.08) - head_h - marks_h) // max(len(rows), 1))
    for i, p in enumerate(periods):
        _t(s, grid_x + i * cw, head_y + Inches(0.04), cw, Inches(0.2),
           [(p.upper(), dict(size=9, bold=True, color=INK_MUTED, spacing=0.8))], align=PP_ALIGN.CENTER)
    top = head_y + head_h
    if marks:
        _r(s, grid_x, top + marks_h // 2, BODY_W - label_w, Emu(9525), fill=CLOUD)
        _t(s, MARGIN, top + Inches(0.08), label_w - Inches(0.2), Inches(0.2),
           [('MILESTONES', dict(size=9, bold=True, color=GREEN_600, spacing=1.2))])
        for at, label in marks:
            mx = grid_x + int((at - 0.5) * cw)
            d = Inches(0.18)
            _r(s, mx - d // 2, top + (marks_h - d) // 2, d, d, fill=GREEN, shape=MSO_SHAPE.DIAMOND)
            mw = Inches(text_w_in(label, 9.5, bold=True) + 0.08)
            mlx = mx + Inches(0.14)
            if mlx + mw > MARGIN + BODY_W:          # no room on the right: label to the left
                mlx = mx - Inches(0.14) - mw
            _t(s, mlx, top + (marks_h - Inches(0.2)) // 2, mw, Inches(0.2), [(label, dict(size=9.5, bold=True, color=FOREST))],
               align=PP_ALIGN.RIGHT if mlx < mx else PP_ALIGN.LEFT)
        top += marks_h
    for r, (name, note, spans) in enumerate(rows):
        ry = top + r * row_h
        nh = Inches(0.22) + (Inches(0.18) if note else 0)
        _t(s, MARGIN, ry + (row_h - nh) // 2, label_w - Inches(0.2), Inches(0.22),
           [(name, dict(size=12, bold=True, color=FOREST))])
        if note:
            _t(s, MARGIN, ry + (row_h - nh) // 2 + Inches(0.22), label_w - Inches(0.2), Inches(0.18),
               [(note, dict(size=9.5, color=INK_MUTED))])
        for p in range(len(periods)):
            _round(_r(s, grid_x + p * cw + Inches(0.04), ry + Inches(0.07), cw - Inches(0.08), row_h - Inches(0.14),
                      fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.05))
        for a, b, st, own in spans:
            bg, ink = tone[st]
            for p in range(a, b + 1):
                _round(_r(s, grid_x + (p - 1) * cw + Inches(0.04), ry + Inches(0.07), cw - Inches(0.08),
                          row_h - Inches(0.14), fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.05))
            label, short = (own, None) if own else names[st]
            room = (b - a + 1) * cw - Inches(0.2)
            word = label if text_w_in(label, 9, bold=True) <= _in(room) else (short or label[:3])
            _t(s, grid_x + (a - 1) * cw + Inches(0.1), ry + (row_h - Inches(0.18)) // 2, room, Inches(0.18),
               [(word, dict(size=9, bold=True, color=ink))])
    return top + len(rows) * row_h


def roadmap(s, y, periods, lanes, *, label_w=Inches(1.8)):
    """Lane by lane and period by period: what happens in each lane each month, as cards in a
       grid. lanes: (name, tone, [[items] per period])."""
    n = len(periods)
    cw = (BODY_W - label_w) // n
    iw = cw - Inches(0.34)
    head_h = Inches(0.42)
    card = lambda t: block_h(t, _in(iw), 10, spacing=1.16) + Inches(0.16)
    lane_h = lambda cells: max([Inches(0.56)] + [sum(card(t) + Inches(0.06) for t in c) + Inches(0.12) for c in cells])
    hs = [lane_h(c) for _, _, c in lanes]
    grid_x = MARGIN + label_w
    for i, p in enumerate(periods):
        x = grid_x + i * cw + Inches(0.04)
        _round(_r(s, x, y, cw - Inches(0.08), head_h - Inches(0.06), fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.08))
        _t(s, x, y + Inches(0.1), cw - Inches(0.08), Inches(0.22), [(p, dict(size=12, bold=True, color=FOREST))],
           align=PP_ALIGN.CENTER)
    yy = y + head_h
    for li, (name, tone_, cells) in enumerate(lanes):
        h = hs[li]
        accent, _ = _GRID.get(tone_ or 'green', _GRID['green'])
        if li % 2 == 0:
            _r(s, MARGIN, yy, BODY_W, h, fill=MIST)
        _r(s, MARGIN, yy + Inches(0.06), Inches(0.06), h - Inches(0.12), fill=accent)
        _t(s, MARGIN + Inches(0.18), yy + Inches(0.14), label_w - Inches(0.3), h - Inches(0.2),
           [(name, dict(size=12, bold=True, color=FOREST))], line_spacing=1.12)
        for i, items in enumerate(cells):
            x = grid_x + i * cw + Inches(0.08)
            cy = yy + Inches(0.06)
            for t in items:
                ch = card(t)
                _round(_r(s, x, cy, cw - Inches(0.16), ch, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE),
                       Inches(0.05))
                _r(s, x, cy + Inches(0.06), Inches(0.04), ch - Inches(0.12), fill=accent)
                _t(s, x + Inches(0.12), cy + Inches(0.08), iw, ch - Inches(0.16), [(t, dict(size=10, color=INK))],
                   line_spacing=1.16)
                cy += ch + Inches(0.06)
        yy += h
    for i in range(1, n):
        _r(s, grid_x + i * cw, y + head_h, Emu(6350), yy - y - head_h, fill=CLOUD)
    return yy


def tiers(s, y, items, *, pick_label='Recommended', gap=Inches(0.36)):
    """Prices side by side: name, price, what it covers and how it is billed. The tier we
       recommend is raised and dark. items: (label, title, price, unit, meta, [lines], billing, pick)."""
    n = len(items)
    cw = min(Inches(3.7), (BODY_W - gap * (n - 1)) // n)
    x0 = MARGIN + (BODY_W - (n * cw + (n - 1) * gap)) // 2
    lift = Inches(0.26)
    pad = Inches(0.3)
    tw = cw - 2 * pad
    lw = tw - Inches(0.32)

    def need(it):
        label, t, price, unit, meta, lines, billing, _ = it
        h = (Inches(0.28) if label else 0) + block_h(t, _in(tw), 19, spacing=1.08, font=DISPLAY) + Inches(0.12)
        h += Inches(0.56) + (Inches(0.24) if unit else 0) + (Inches(0.42) if meta else 0) + Inches(0.2)
        h += sum(block_h(ln, _in(lw), 11, spacing=1.22) + Inches(0.1) for ln in lines)
        h += (Inches(0.16) + block_h(billing, _in(tw), 10.5, spacing=1.2)) if billing else 0
        return h
    ch = max(need(it) for it in items) + 2 * Inches(0.26)
    top = y + Inches(0.14)
    for i, (label, t, price, unit, meta, lines, billing, pick) in enumerate(items):
        x = x0 + i * (cw + gap)
        card_y, card_h = (top, ch + 2 * lift) if pick else (top + lift, ch)
        _round(_r(s, x, card_y, cw, card_h, fill=FOREST if pick else WHITE, line=None if pick else CLOUD,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.14))
        if pick:
            _pill(s, x + cw // 2, top - Inches(0.12), pick_label, 'green', align='center')
        ink, soft = (WHITE, MIST_INK) if pick else (FOREST, INK_800)
        yy = top + lift + Inches(0.26)
        if label:
            _t(s, x + pad, yy, tw, Inches(0.2), [(label.upper(), dict(size=9.5, bold=True,
                                                                       color=GREEN_300 if pick else GREEN_600, spacing=1.4))])
            yy += Inches(0.28)
        th = block_h(t, _in(tw), 19, spacing=1.08, font=DISPLAY)
        _t(s, x + pad, yy, tw, th, [(t, dict(font=DISPLAY, size=19, color=ink))], line_spacing=1.06)
        yy += th + Inches(0.12)
        _t(s, x + pad, yy, tw, Inches(0.5), [(price, dict(size=30, bold=True, color=ink))])
        yy += Inches(0.56)
        if unit:
            _t(s, x + pad, yy, tw, Inches(0.2), [(unit, dict(size=11, color=GREEN_300 if pick else INK_MUTED))])
            yy += Inches(0.24)
        if meta:
            _pill(s, x + pad, yy + Inches(0.06), meta, 'mint' if pick else 'forest', size=9, h=Inches(0.26))
            yy += Inches(0.42)
        _r(s, x + pad, yy + Inches(0.06), tw, Emu(9525), fill=GREEN_600 if pick else CLOUD)
        yy += Inches(0.2)
        for ln in lines:
            lh = block_h(ln, _in(lw), 11, spacing=1.22)
            icon(s, 'check', x + pad, yy + Inches(0.01), Inches(0.2), color=GREEN_300 if pick else GREEN)
            _t(s, x + pad + Inches(0.32), yy, lw, lh, [(ln, dict(size=11, color=soft))], line_spacing=1.22)
            yy += lh + Inches(0.1)
        if billing:
            bh = block_h(billing, _in(tw), 10.5, spacing=1.2)
            _t(s, x + pad, card_y + card_h - bh - Inches(0.26), tw, bh,
               [(billing, dict(size=10.5, italic=True, color=GREEN_300 if pick else INK_MUTED))], line_spacing=1.2)
    return top + ch + 2 * lift


def faq(s, y, items, *, cols=2, gap=Inches(0.55)):
    """The questions a reader will ask, each with our answer, in a two-column grid. items: (q, a)."""
    cw = (BODY_W - gap * (cols - 1)) // cols
    tw = cw - Inches(0.62)
    need = lambda q, a: (block_h(q, _in(tw), 15, spacing=1.12, font=DISPLAY) + Inches(0.08)
                         + block_h(a, _in(tw), 11.5, spacing=1.25))
    yy = y
    for r in range(0, len(items), cols):
        row = items[r:r + cols]
        rh = max(need(q, a) for q, a in row) + Inches(0.34)
        if r:
            _r(s, MARGIN, yy, BODY_W, Emu(9525), fill=CLOUD)
        for c, (q, a) in enumerate(row):
            x = MARGIN + c * (cw + gap)
            _r(s, x, yy + Inches(0.17), Inches(0.38), Inches(0.38), fill=MINT, shape=MSO_SHAPE.OVAL)
            _t(s, x, yy + Inches(0.24), Inches(0.38), Inches(0.24), [('Q', dict(size=12, bold=True, color=FOREST))],
               align=PP_ALIGN.CENTER)
            qh = block_h(q, _in(tw), 15, spacing=1.12, font=DISPLAY)
            _t(s, x + Inches(0.62), yy + Inches(0.17), tw, qh, [(q, dict(font=DISPLAY, size=15, color=FOREST))],
               line_spacing=1.1)
            ah = block_h(a, _in(tw), 11.5, spacing=1.25)
            _t(s, x + Inches(0.62), yy + Inches(0.17) + qh + Inches(0.08), tw, ah,
               [(a, dict(size=11.5, color=INK_800))], line_spacing=1.25)
        yy += rh
    return yy


def ledger(s, y, sections, *, gap=Inches(0.4)):
    """The math behind a price, set like a receipt: line items with their basis and amount,
       and a ruled total. Sections sit side by side. sections: (title, [(item, basis, amount)],
       (label, amount) or None)."""
    n = len(sections)
    cw = (BODY_W - gap * (n - 1)) // n
    pad = Inches(0.3)
    amt_w = Inches(1.3)
    tw = cw - 2 * pad - amt_w - Inches(0.12)
    line_h = lambda it: (block_h(it[0], _in(tw), 12, spacing=1.15, bold=True)
                         + (Inches(0.03) + block_h(it[1], _in(tw), 10, spacing=1.2) if it[1] else 0) + Inches(0.22))
    head_h, total_h = Inches(0.52), Inches(0.62)
    ch = head_h + max(sum(line_h(it) for it in lines) for _, lines, _ in sections) + Inches(0.1) + total_h
    for k, (title_, lines, total) in enumerate(sections):
        x = MARGIN + k * (cw + gap)
        _round(_r(s, x, y, cw, ch, fill=WHITE, line=CLOUD, shape=MSO_SHAPE.ROUNDED_RECTANGLE), Inches(0.12))
        _round(_r(s, x, y, cw, head_h, fill=MINT_TINT, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE), Inches(0.12))
        _t(s, x + pad, y + Inches(0.16), cw - 2 * pad, Inches(0.24),
           [(title_.upper(), dict(size=10.5, bold=True, color=FOREST, spacing=1.2))])
        yy = y + head_h
        for j, (item, basis, amount) in enumerate(lines):
            h = line_h((item, basis, amount))
            ih = block_h(item, _in(tw), 12, spacing=1.15, bold=True)
            _t(s, x + pad, yy + Inches(0.11), tw, ih, [(item, dict(size=12, bold=True, color=INK))], line_spacing=1.15)
            if basis:
                _t(s, x + pad, yy + Inches(0.14) + ih, tw, block_h(basis, _in(tw), 10, spacing=1.2),
                   [(basis, dict(size=10, color=INK_MUTED))], line_spacing=1.2)
            _t(s, x + cw - pad - amt_w, yy + Inches(0.11), amt_w, Inches(0.24),
               [(amount, dict(size=12, color=INK))], align=PP_ALIGN.RIGHT)
            if j < len(lines) - 1:
                ln = _line(s, x + pad, yy + h, x + cw - pad, yy + h, CLOUD, w=Pt(0.75), dash=True)
            yy += h
        if total:
            ty = y + ch - total_h
            _r(s, x + pad, ty, cw - 2 * pad, Emu(19050), fill=FOREST)
            _t(s, x + pad, ty + Inches(0.2), cw - 2 * pad - amt_w, Inches(0.26),
               [(total[0], dict(size=12.5, bold=True, color=FOREST))])
            _t(s, x + cw - pad - Inches(2.0), ty + Inches(0.15), Inches(2.0), Inches(0.32),
               [(total[1], dict(size=16, bold=True, color=FOREST))], align=PP_ALIGN.RIGHT)
    return y + ch


def ring(s, y, center, stages, *, h=Inches(3.5)):
    """A lifecycle ring: stages around a circle that goes round again, the thing that repeats
       named in the middle, and each stage spelled out in a numbered list beside the ring.
       center: (title, text); stages: (title, text) 3-6, clockwise from the top."""
    import math
    n = len(stages)
    r = min(h // 2 - Inches(0.34), Inches(1.6))
    cx = MARGIN + r + Inches(0.5)
    cy = y + h // 2
    node = Inches(0.62)
    ang = [math.pi / 2 - 2 * math.pi * i / n for i in range(n)]          # clockwise from the top
    gap = (node / 2 + Inches(0.08)) / r                                   # leave the nodes clear
    for i in range(n):
        t0, t1 = ang[i] - gap, ang[(i + 1) % n] + gap
        if t1 > t0:
            t1 -= 2 * math.pi
        _arc_arrow(s, cx, cy, r, r, t0, t1, GREEN_300, w=Pt(2))
    ctitle, ctext = center
    cw = int(r * 1.25)
    th = block_h(ctitle, _in(cw), 20, spacing=1.08, font=DISPLAY)
    xh = block_h(ctext, _in(cw), 11.5, spacing=1.22) if ctext else 0
    top = cy - (th + (Inches(0.08) + xh if ctext else 0)) // 2
    _t(s, cx - cw // 2, top, cw, th, [(ctitle, dict(font=DISPLAY, size=20, color=FOREST))],
       align=PP_ALIGN.CENTER, line_spacing=1.06)
    if ctext:
        _t(s, cx - cw // 2, top + th + Inches(0.08), cw, xh, [(ctext, dict(size=11.5, color=INK_MUTED))],
           align=PP_ALIGN.CENTER, line_spacing=1.22)
    for i, a in enumerate(ang):
        nx, ny = cx + int(r * math.cos(a)), cy - int(r * math.sin(a))
        _r(s, nx - node // 2, ny - node // 2, node, node, fill=FOREST if i == 0 else WHITE, line=GREEN, line_w=Pt(2),
           shape=MSO_SHAPE.OVAL)
        _t(s, nx - node // 2, ny - Inches(0.14), node, Inches(0.28),
           [(str(i + 1), dict(size=15, bold=True, color=WHITE if i == 0 else FOREST))], align=PP_ALIGN.CENTER)
    # the stages, spelled out beside the ring
    lx = cx + r + node // 2 + Inches(0.8)
    lw = MARGIN + BODY_W - lx
    tw = lw - Inches(0.62)
    rows = [(block_h(t, _in(tw), 15, spacing=1.12, bold=True), block_h(x or '', _in(tw), 12, spacing=1.25) if x else 0)
            for t, x in stages]
    row_h = [a + (Inches(0.06) + b if b else 0) for a, b in rows]
    space = max(Inches(0.16), (h - sum(row_h)) // max(n, 1))
    yy = y + max(0, (h - sum(row_h) - space * (n - 1)) // 2)
    for i, ((t, x), (a, b)) in enumerate(zip(stages, rows)):
        d = Inches(0.36)
        _r(s, lx, yy, d, d, fill=FOREST if i == 0 else MINT, shape=MSO_SHAPE.OVAL)
        _t(s, lx, yy + Inches(0.07), d, Inches(0.22), [(str(i + 1), dict(size=11.5, bold=True,
                                                                         color=WHITE if i == 0 else FOREST))],
           align=PP_ALIGN.CENTER)
        _t(s, lx + Inches(0.62), yy + Inches(0.02), tw, a, [(t, dict(size=15, bold=True, color=FOREST))], line_spacing=1.12)
        if x:
            _t(s, lx + Inches(0.62), yy + Inches(0.02) + a + Inches(0.06), tw, b, [(x, dict(size=12, color=INK_800))],
               line_spacing=1.25)
        yy += row_h[i] + space
    return y + h
