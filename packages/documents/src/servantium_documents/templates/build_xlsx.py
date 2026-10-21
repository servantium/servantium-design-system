"""Servantium Excel template.

GOOGLE SHEETS SHAPES THIS FILE. Sheets has no named cell styles and does not keep
Excel Table objects faithfully, so:
  * every cell carries DIRECT formatting (named styles are registered too, so Excel
    users get a Servantium cell-style gallery, but nothing depends on them);
  * tables are plain ranges with an autofilter, not ListObjects, because a Table
    style would override the header formatting in Excel and not survive in Sheets;
  * zebra striping is conditional formatting (=MOD(ROW(),2)=0), so it survives a
    sort; hand-painted stripes scramble the moment someone sorts.

THE MODELING CONVENTION IS THE STRUCTURE. Consultants and finance readers expect it
and read it without a legend:
    blue text      an input, typed by a person, safe to change
    black text     a formula, do not overwrite
    green text     pulls from another sheet
It deliberately overrides the brand palette for inputs: blue-means-input is worth
more than colour consistency.

Outputs:
    dist/Servantium Workbook Template.xlsx  -> Google: open, then File > Make a copy
    dist/Servantium.xltx                    -> Excel: File > New from template
"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.writer.theme import theme_xml

from .. import brand
from ..brand import K, hexof

BODY, DISPLAY = K.UI, K.DISPLAY
FOREST, GREEN, GREEN6 = hexof(K.FOREST), hexof(K.GREEN), hexof(K.GREEN_600)
MINT, MINT_T, MIST, CLOUD = hexof(K.MINT), hexof(K.MINT_TINT), hexof(K.MIST), hexof(K.CLOUD)
INK, MUTED, WHITE = hexof(K.INK_800), hexof(K.INK_MUTED), hexof(K.WHITE)
_tok = lambda n: K.T.COLORS[n].lstrip('#')
INPUT_BLUE = _tok('color-model-input')     # convention over palette, see module docstring
LINK_GREEN = _tok('color-model-link')      # darker than brand green so it reads as text

thin = Side(style='thin', color=CLOUD)
hair = Border(bottom=thin)
fill = lambda hx: PatternFill('solid', fgColor=hx)
F = lambda **kw: Font(name=kw.pop('name', BODY), size=kw.pop('size', 10), **kw)

NUM = {'currency': '"$"#,##0', 'currency2': '"$"#,##0.00', 'int': '#,##0',
       'pct': '0.0%', 'date': 'mmm d, yyyy', 'hours': '#,##0.0'}


def register_styles(wb):
    """A Servantium cell-style gallery for Excel. Sheets ignores the names but keeps
       the formatting, because openpyxl writes the full format onto each cell."""
    styles = {
        'SV Header': dict(font=F(bold=True, color=WHITE, size=9.5), fill=fill(FOREST),
                          alignment=Alignment(vertical='center', wrap_text=True)),
        'SV Input': dict(font=F(color=INPUT_BLUE), fill=fill(_tok('color-model-input-fill')), border=hair),
        'SV Formula': dict(font=F(color=INK), border=hair),
        'SV Link': dict(font=F(color=LINK_GREEN), border=hair),
        'SV Label': dict(font=F(bold=True, color=MUTED, size=8.5)),
        'SV Title': dict(font=F(name=DISPLAY, size=24, color=FOREST)),
    }
    for name, kw in styles.items():
        ns = NamedStyle(name=name, **kw)
        wb.add_named_style(ns)


def set_default_font(wb):
    """Empty cells and new typing inherit the Normal style's font."""
    normal = next(s for s in wb._named_styles if s.name == 'Normal')
    normal.font = F()
    wb._fonts[0] = F()


def widths(ws, cols):
    for i, w in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def header_row(ws, row, labels, col=1):
    ws.row_dimensions[row].height = 26
    for j, lbl in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lbl.upper())
        c.style = 'SV Header'
        c.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True, indent=1)


def eyebrow(ws, cell, text):
    ws[cell] = text.upper()
    ws[cell].font = F(bold=True, color=GREEN6, size=9)


def page_setup(ws, landscape=True):
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = K.FOOTER_TEXT
    ws.oddFooter.left.size, ws.oddFooter.left.font = 8, BODY
    ws.oddFooter.right.text = 'Page &P of &N'
    ws.oddFooter.right.size, ws.oddFooter.right.font = 8, BODY
    ws.sheet_view.showGridLines = False


def zebra(ws, ref):
    ws.conditional_formatting.add(ref, FormulaRule(formula=['MOD(ROW(),2)=0'], fill=fill(MIST)))


def build():
    wb = Workbook()
    wb.loaded_theme = brand.brand_theme_xml(theme_xml.encode() if isinstance(theme_xml, str) else theme_xml)
    set_default_font(wb)
    register_styles(wb)

    # ── Cover ────────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Cover'
    ws.sheet_properties.tabColor = FOREST
    widths(ws, [3, 26, 58, 3])
    logo = XLImage(str(brand.LOGO_ON_LIGHT))
    logo.width, logo.height = 188, int(188 * logo.height / logo.width)
    ws.add_image(logo, 'B2')
    eyebrow(ws, 'B7', 'Client  ·  Workbook purpose')
    ws['B8'] = 'Workbook title goes here'; ws['B8'].style = 'SV Title'
    ws.row_dimensions[8].height = 36
    ws['B9'] = 'One line on the question this workbook answers.'
    ws['B9'].font = F(size=12, color=MUTED)

    meta = [('Prepared for', '[Client name]'), ('Owner', '[Name]'), ('Version', 'v1'),
            ('Status', 'Draft'), ('Last updated', dt.date(2026, 9, 29))]
    for i, (k, v) in enumerate(meta):
        r = 11 + i
        ws.cell(row=r, column=2, value=k.upper()).style = 'SV Label'
        c = ws.cell(row=r, column=3, value=v)
        c.style = 'SV Input'
        if isinstance(v, dt.date):
            c.number_format = NUM['date']
            c.alignment = Alignment(horizontal='left')

    eyebrow(ws, 'B18', 'How this workbook works')
    legend = [('Blue text', 'An input. Typed by a person, safe to change.', 'SV Input'),
              ('Black text', 'A formula. Do not type over it.', 'SV Formula'),
              ('Green text', 'Pulls from another sheet.', 'SV Link'),
              ('Source column', 'Every number says where it came from. Inferred means "probably".', 'SV Formula')]
    for i, (k, v, st) in enumerate(legend):
        r = 19 + i
        a = ws.cell(row=r, column=2, value=k); a.style = st; a.font = F(bold=True, color=a.font.color.rgb[-6:] if a.font.color and a.font.color.rgb else INK)
        ws.cell(row=r, column=3, value=v).font = F(color=INK)

    eyebrow(ws, 'B24', 'Contents')
    for i, (sheet, purpose) in enumerate([('Inputs', 'Every assumption in one place, each with its source.'),
                                          ('Data', 'The records. One row per item, filterable.'),
                                          ('Summary', 'What the numbers say. Formulas only.')]):
        r = 25 + i
        ws.cell(row=r, column=2, value=sheet).font = F(bold=True, color=FOREST)
        ws.cell(row=r, column=3, value=purpose).font = F(color=INK)
    page_setup(ws, landscape=False)

    # ── Inputs ───────────────────────────────────────────────────────────────
    wi = wb.create_sheet('Inputs')
    wi.sheet_properties.tabColor = INPUT_BLUE
    widths(wi, [3, 34, 16, 12, 34, 30])
    eyebrow(wi, 'B2', 'Inputs')
    wi['B3'] = 'Assumptions'; wi['B3'].style = 'SV Title'; wi.row_dimensions[3].height = 34
    wi['B4'] = 'Blue cells are the only ones to edit. Everything else recalculates.'
    wi['B4'].font = F(color=MUTED)
    header_row(wi, 6, ['Assumption', 'Value', 'Unit', 'Source', 'Notes'], col=2)
    inputs = [('Standard bill rate', 185, 'USD / hr', 'Rate card, 2026', ''),
              ('Target utilization', 0.78, '%', 'Practice plan', 'Billable over available'),
              ('Hours per FTE per week', 40, 'hours', 'Policy', ''),
              ('Discount ceiling', 0.10, '%', 'Pricing policy', 'Above this needs approval')]
    for i, (name, val, unit, src, note) in enumerate(inputs):
        r = 7 + i
        wi.cell(row=r, column=2, value=name).style = 'SV Formula'
        v = wi.cell(row=r, column=3, value=val); v.style = 'SV Input'
        v.number_format = NUM['pct'] if unit == '%' else (NUM['currency'] if 'USD' in unit else NUM['int'])
        wi.cell(row=r, column=4, value=unit).style = 'SV Formula'
        wi.cell(row=r, column=5, value=src).style = 'SV Formula'
        wi.cell(row=r, column=6, value=note).style = 'SV Formula'
    wi.freeze_panes = 'B7'
    page_setup(wi)

    # ── Data ─────────────────────────────────────────────────────────────────
    wd = wb.create_sheet('Data')
    wd.sheet_properties.tabColor = GREEN
    cols = ['ID', 'Date', 'Client', 'Category', 'Hours', 'Rate', 'Amount', 'Status', 'Source']
    widths(wd, [3, 10, 14, 22, 18, 10, 11, 14, 13, 26])
    eyebrow(wd, 'B2', 'Data')
    wd['B3'] = 'Records'; wd['B3'].style = 'SV Title'; wd.row_dimensions[3].height = 34
    wd['B4'] = 'Rows marked EXAMPLE are sample data. Delete them when real data starts.'
    wd['B4'].font = F(color=MUTED)
    header_row(wd, 6, cols, col=2)
    cats = ['Discovery', 'Delivery', 'Delivery', 'Training', 'Support', 'Delivery',
            'Discovery', 'Delivery', 'Support', 'Training', 'Delivery', 'Delivery']
    status = ['Approved', 'Approved', 'Review', 'Draft', 'Approved', 'Approved',
              'Review', 'Approved', 'Draft', 'Approved', 'Review', 'Approved']
    hours = [12, 40, 32, 8, 6, 44, 10, 36, 4, 16, 28, 40]
    first, last = 7, 7 + len(cats) - 1
    for i in range(len(cats)):
        r = first + i
        vals = [f'EX-{i + 1:03d}', dt.date(2026, 7, 6) + dt.timedelta(days=7 * i),
                ['Client A', 'Client B', 'Client C'][i % 3], cats[i], hours[i]]
        for j, v in enumerate(vals):
            c = wd.cell(row=r, column=2 + j, value=v); c.style = 'SV Input'
        wd.cell(row=r, column=3).number_format = NUM['date']
        wd.cell(row=r, column=6).number_format = NUM['hours']
        rate = wd.cell(row=r, column=7, value='=Inputs!$C$7'); rate.style = 'SV Link'
        rate.number_format = NUM['currency']
        amt = wd.cell(row=r, column=8, value=f'=F{r}*G{r}'); amt.style = 'SV Formula'
        amt.number_format = NUM['currency']
        st = wd.cell(row=r, column=9, value=status[i]); st.style = 'SV Input'
        wd.cell(row=r, column=10, value='EXAMPLE').style = 'SV Formula'
    ref = f'B{first}:J{last}'
    zebra(wd, ref)
    wd.auto_filter.ref = f'B6:J{last}'
    wd.freeze_panes = 'C7'
    dv = DataValidation(type='list', formula1='"Draft,Review,Approved"', allow_blank=True)
    dv.error, dv.errorTitle = 'Choose Draft, Review or Approved.', 'Status'
    wd.add_data_validation(dv); dv.add(f'I{first}:I500')
    # Status reads at a glance: amber for anything not yet approved.
    wd.conditional_formatting.add(f'I{first}:I500', FormulaRule(
        formula=[f'AND(I{first}<>"",I{first}<>"Approved")'], font=Font(color=_tok('color-model-warning'), bold=True)))
    page_setup(wd)

    # ── Summary ──────────────────────────────────────────────────────────────
    wsu = wb.create_sheet('Summary')
    wsu.sheet_properties.tabColor = hexof(K.SLATE)
    widths(wsu, [3, 24, 16, 3, 24, 16, 3, 24, 16])
    eyebrow(wsu, 'B2', 'Summary')
    wsu['B3'] = 'What the numbers say'; wsu['B3'].style = 'SV Title'; wsu.row_dimensions[3].height = 34
    wsu['B4'] = 'Formulas only. Change the Inputs or the Data, never this sheet.'
    wsu['B4'].font = F(color=MUTED)
    D = f'Data!$H${first}:$H${last}'
    kpis = [('B', 'Total amount', f'=SUM({D})', NUM['currency']),
            ('E', 'Total hours', f'=SUM(Data!$F${first}:$F${last})', NUM['hours']),
            ('H', 'Approved share', f'=COUNTIF(Data!$I${first}:$I${last},"Approved")/COUNTA(Data!$I${first}:$I${last})',
             NUM['pct'])]
    for col, label, formula, fmt in kpis:
        ci = ord(col) - 64
        for r in (6, 7, 8):
            for cc in (ci, ci + 1):
                wsu.cell(row=r, column=cc).fill = fill(MINT_T)
        wsu.cell(row=6, column=ci).border = Border(left=Side(style='thick', color=GREEN))
        wsu.cell(row=7, column=ci).border = Border(left=Side(style='thick', color=GREEN))
        wsu.cell(row=8, column=ci).border = Border(left=Side(style='thick', color=GREEN))
        lab = wsu.cell(row=6, column=ci, value=label.upper()); lab.font = F(bold=True, color=MUTED, size=8.5)
        lab.alignment = Alignment(indent=1, vertical='bottom')
        wsu.merge_cells(start_row=7, start_column=ci, end_row=8, end_column=ci + 1)
        v = wsu.cell(row=7, column=ci, value=formula)
        v.font = F(size=22, bold=True, color=LINK_GREEN)
        v.number_format = fmt
        v.alignment = Alignment(indent=1, vertical='center')
    wsu.row_dimensions[6].height = 20
    wsu.row_dimensions[7].height = wsu.row_dimensions[8].height = 22

    header_row(wsu, 11, ['Category', 'Hours', '', 'Amount', 'Share of amount'], col=2)
    for i, cat in enumerate(['Discovery', 'Delivery', 'Training', 'Support']):
        r = 12 + i
        wsu.cell(row=r, column=2, value=cat).style = 'SV Formula'
        h = wsu.cell(row=r, column=3, value=f'=SUMIFS(Data!$F${first}:$F${last},Data!$E${first}:$E${last},B{r})')
        h.style = 'SV Link'; h.number_format = NUM['hours']
        a = wsu.cell(row=r, column=5, value=f'=SUMIFS({D},Data!$E${first}:$E${last},B{r})')
        a.style = 'SV Link'; a.number_format = NUM['currency']
        s_ = wsu.cell(row=r, column=6, value=f'=IFERROR(E{r}/$B$7,0)'); s_.style = 'SV Formula'
        s_.number_format = NUM['pct']
    zebra(wsu, 'B12:F15')
    page_setup(wsu)

    wb.active = 0
    return wb


def save(wb):
    xlsx = K.TEMPLATE_XLSX
    xltx = brand.DIST / f'{K.BRAND_SHORT}.xltx'
    wb.save(str(xlsx))
    wb.template = True            # openpyxl writes the template content type for .xltx
    wb.save(str(xltx))
    return xlsx, xltx


if __name__ == '__main__':
    a, b = save(build())
    print(f'-> {a.name}, {b.name}')
