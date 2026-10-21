"""Servantium PowerPoint template: a real master and eleven branded layouts.

WHY THIS IS NOT THE OLD TEMPLATE. Every deck the engine has produced so far sits on
python-pptx's stock "Blank" layout with zero placeholders: the brand was painted on
as loose shapes. That looks right and behaves wrong. There is no "New slide ->
Servantium" for a human, titles are not titles (blank filmstrip, no outline, nothing
for a screen reader), and Google's Theme -> Import theme finds nothing to import.

Here the brand lives in the MASTER and LAYOUTS. Placeholders carry the type styles,
so a human who never opens Python gets the same result as the engine, and the
engine can fill the same placeholders by idx.

PLACEHOLDER idx CONVENTION (the engine depends on it; do not renumber):
    title        type=title / ctrTitle
     1           main body              2   second body (two content, comparison)
    10           eyebrow                11  lede / standfirst
    13           source or note         14  picture
    15-17        big-number values      18-20  big-number labels
    21, 22       comparison headings    12  slide number

Prompt text in each placeholder is the method, not filler: "State the conclusion"
teaches an action title to whoever clicks into it.

Outputs, both from the same build:
    dist/Servantium Slides Template.pptx   -> Google: Theme > Import theme; or open and edit
    dist/Servantium.potx                   -> PowerPoint: File > New from template
"""
import copy
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn

from .. import brand
from ..brand import K, hexof

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = f'xmlns:a="{A}" xmlns:p="{P}" xmlns:r="{R}"'

W, H = int(Inches(13.333)), int(Inches(7.5))
M = int(Inches(0.62))                    # side margin
BW = W - 2 * M                           # body width
FOOT_Y = int(Inches(6.86))
E = lambda inches: int(Inches(inches))   # inches -> EMU

# ── XML helpers ──────────────────────────────────────────────────────────────
_id = [100]


def nid():
    _id[0] += 1
    return _id[0]


def solid(hx=None, scheme=None):
    if scheme:
        return f'<a:solidFill><a:schemeClr val="{scheme}"/></a:solidFill>'
    return f'<a:solidFill><a:srgbClr val="{hx}"/></a:solidFill>'


def rpr(size, *, color=None, scheme=None, bold=False, caps=False, spc=0, font='+mn-lt'):
    fill = solid(color, scheme) if (color or scheme) else ''
    b = ' b="1"' if bold else ' b="0"'
    c = ' cap="all"' if caps else ''
    s = f' spc="{spc}"' if spc else ''
    return (f'<a:defRPr sz="{int(size * 100)}"{b}{c}{s}>{fill}'
            f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/></a:defRPr>')


def lvl(n, size, *, color=None, scheme=None, bold=False, caps=False, spc=0, font='+mn-lt',
        algn='l', bullet=False, ln=None, before=0, after=0, indent_in=0.26):
    """One lvlNpPr. Bullets are green dots; spacing is in points."""
    mar = ' marL="0" indent="0"'
    bu = '<a:buNone/>'
    if bullet:
        mar = f' marL="{E(indent_in * n)}" indent="-{E(indent_in)}"'
        bu = (f'<a:buClr><a:schemeClr val="accent1"/></a:buClr>'
              f'<a:buFont typeface="Arial"/><a:buChar char="&#8226;"/>')
    lns = f'<a:lnSpc><a:spcPct val="{int(ln * 1000)}"/></a:lnSpc>' if ln else ''
    sb = f'<a:spcBef><a:spcPts val="{int(before * 100)}"/></a:spcBef>'
    sa = f'<a:spcAft><a:spcPts val="{int(after * 100)}"/></a:spcAft>'
    return (f'<a:lvl{n}pPr algn="{algn}"{mar}>{lns}{sb}{sa}{bu}'
            f'{rpr(size, color=color, scheme=scheme, bold=bold, caps=caps, spc=spc, font=font)}'
            f'</a:lvl{n}pPr>')


def xfrm(x, y, w, h):
    return f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'


def ph(kind, idx, name, x, y, w, h, prompt, styles, *, anchor='t', wrap=True, autofit=True):
    """A placeholder with its own list style and a teaching prompt."""
    i = nid()
    idx_attr = f' idx="{idx}"' if idx is not None else ''
    type_attr = f' type="{kind}"' if kind else ''
    fit = '<a:normAutofit/>' if autofit else '<a:noAutofit/>'
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="{name}"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph{type_attr}{idx_attr}/></p:nvPr></p:nvSpPr>'
            f'<p:spPr>{xfrm(x, y, w, h)}</p:spPr>'
            f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}" '
            f'wrap="{"square" if wrap else "none"}">{fit}</a:bodyPr>'
            f'<a:lstStyle>{styles}</a:lstStyle>'
            f'<a:p><a:r><a:rPr lang="en-US"/><a:t>{prompt}</a:t></a:r></a:p></p:txBody></p:sp>')


def box(x, y, w, h, fill_xml, name='Decoration', geom='rect', line_xml='<a:ln><a:noFill/></a:ln>'):
    i = nid()
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/>'
            f'</p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>'
            f'{fill_xml}{line_xml}</p:spPr>'
            f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')


def static_text(x, y, w, h, text, size, *, color=None, scheme=None, bold=False, caps=False,
                spc=0, algn='l', name='Label'):
    i = nid()
    fill = solid(color, scheme)
    b = ' b="1"' if bold else ''
    c = ' cap="all"' if caps else ''
    s = f' spc="{spc}"' if spc else ''
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="{name}"/><p:cNvSpPr txBox="1"/>'
            f'<p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>'
            f'<a:p><a:pPr algn="{algn}"/><a:r><a:rPr lang="en-US" sz="{int(size * 100)}"{b}{c}{s}>'
            f'{fill}<a:latin typeface="+mn-lt"/></a:rPr><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>')


def pic(rid, x, y, w, h, name='Logo', alpha=None):
    i = nid()
    a = f'<a:alphaModFix amt="{alpha}"/>' if alpha else ''
    return (f'<p:pic><p:nvPicPr><p:cNvPr id="{i}" name="{name}"/>'
            f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>'
            f'<p:blipFill><a:blip r:embed="{rid}">{a}</a:blip><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')


def img_size(path, width_emu=None, height_emu=None, max_h=None):
    """Scale to a width (capped at max_h) or to a height. The cap keeps a square or tall
       logo from swelling to the full width a wide lockup was given."""
    from PIL import Image
    im = Image.open(path)
    if width_emu:
        h = int(width_emu * im.height / im.width)
        if max_h and h > max_h:
            return int(max_h * im.width / im.height), max_h
        return width_emu, h
    return int(height_emu * im.width / im.height), height_emu


def sld_num(dark=False):
    return ph('sldNum', 12, 'Slide Number', W - M - E(1.0), FOOT_Y - E(0.02), E(1.0), E(0.3), '‹#›',
              lvl(1, 9.5, color=hexof(K.GREEN_300) if dark else hexof(K.INK_MUTED), algn='r'),
              anchor='ctr', autofit=False)


SPTREE_HEAD = ('<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
               '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
               '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>')


def set_sptree(cSld, inner_xml):
    old = cSld.find(qn('p:spTree'))
    new = etree.fromstring(f'<p:spTree {NS}>{SPTREE_HEAD}{inner_xml}</p:spTree>')
    cSld.replace(old, new)


def set_bg(cSld, scheme=None, hx=None):
    old = cSld.find(qn('p:bg'))
    if old is not None:
        cSld.remove(old)
    if not (scheme or hx):
        return
    bg = etree.fromstring(f'<p:bg {NS}><p:bgPr>{solid(hx, scheme)}<a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, bg)


# ── common pieces ────────────────────────────────────────────────────────────
EYEBROW_STYLE = lvl(1, 10, color=hexof(K.GREEN_600), bold=True, caps=True, spc=160)
LEDE_STYLE = lvl(1, 15, color=hexof(K.INK_MUTED), ln=128)
SOURCE_STYLE = lvl(1, 9, color=hexof(K.INK_MUTED))
TITLE_STYLE = lvl(1, 30, color=hexof(K.FOREST), font='+mj-lt', ln=104)
BODY_STYLES = (lvl(1, 16, color=hexof(K.INK_800), bullet=True, ln=125, after=9) +
               lvl(2, 14, color=hexof(K.INK_800), bullet=True, ln=125, after=6) +
               lvl(3, 12.5, color=hexof(K.INK_MUTED), bullet=True, ln=125, after=4))

TOP = E(0.62)


def header(eyebrow=True, lede=True, title_h=1.02):
    """Eyebrow, action title, optional lede. Returns (xml, y_after)."""
    x = ''
    y = TOP
    if eyebrow:
        x += ph('body', 10, 'Eyebrow', M, y, BW, E(0.26), 'SECTION · TOPIC', EYEBROW_STYLE,
                anchor='t', autofit=False)
        y += E(0.34)
    x += ph('title', None, 'Title', M, y, E(11.9), E(title_h),
            'State the conclusion in one sentence. Not the topic.', TITLE_STYLE, anchor='t')
    y += E(title_h) + E(0.1)
    if lede:
        x += ph('body', 11, 'Lede', M, y, E(11.4), E(0.42),
                'Optional standfirst. One line of context.', LEDE_STYLE, autofit=True)
        y += E(0.56)
    return x, y


def source_line():
    return sld_num() + ph('body', 13, 'Source', M, FOOT_Y - E(0.52), E(9.5), E(0.28),
              'Source: every number on the slide has one.', SOURCE_STYLE, anchor='b', autofit=True)


# ── build ────────────────────────────────────────────────────────────────────
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    master = prs.slide_masters[0]
    mpart = master.part

    # Images the master and layouts share.
    lock_green = brand.LOGO_ON_LIGHT
    lock_white = brand.LOGO_ON_DARK
    mark_white = brand.MARK_ON_DARK

    # ── master: text styles, footer, placeholders ───────────────────────────
    msld = master._element
    tx = msld.find(qn('p:txStyles'))
    new_tx = etree.fromstring(
        f'<p:txStyles {NS}>'
        f'<p:titleStyle>{TITLE_STYLE}</p:titleStyle>'
        f'<p:bodyStyle>{BODY_STYLES}'
        + ''.join(lvl(n, 12, color=hexof(K.INK_MUTED), bullet=True, ln=125) for n in (4, 5)) +
        f'</p:bodyStyle>'
        f'<p:otherStyle>{lvl(1, 14, color=hexof(K.INK_800))}</p:otherStyle></p:txStyles>')
    msld.replace(tx, new_tx)

    _, rid_green = mpart.get_or_add_image_part(str(lock_green))
    lw, lh = img_size(lock_green, width_emu=E(1.18), max_h=E(0.34))
    master_xml = (
        ph('title', None, 'Title', M, TOP + E(0.34), E(11.9), E(1.02), 'Master title', TITLE_STYLE) +
        ph('body', 1, 'Body', M, E(2.4), BW, E(3.9), 'Master body', BODY_STYLES) +
        box(M, FOOT_Y - E(0.2), BW, Emu(9525), solid(hexof(K.CLOUD)), name='Footer rule') +
        pic(rid_green, M, FOOT_Y + E(0.13) - lh // 2, lw, lh, name='Footer logo') +
        static_text(M + lw + E(0.3), FOOT_Y - E(0.02), E(3.0), E(0.3), K.COPYRIGHT, 9.5,
                    color=hexof(K.INK_MUTED), name='Copyright') +
        sld_num())
    set_sptree(msld.find(qn('p:cSld')), master_xml)
    set_bg(msld.find(qn('p:cSld')), scheme='bg1')

    layouts = list(master.slide_layouts)
    assert len(layouts) == 11, 'python-pptx default should give eleven layouts to repurpose'

    def define(layout, name, inner, *, bg_scheme=None, bg_hex=None, show_master=True, ltype='cust'):
        el = layout._element
        el.set('type', ltype)
        el.set('preserve', '1')
        el.set('showMasterSp', '1' if show_master else '0')
        cSld = el.find(qn('p:cSld'))
        cSld.set('name', name)
        set_bg(cSld, scheme=bg_scheme, hx=bg_hex)
        set_sptree(cSld, inner)

    def dark_footer(layout):
        _, rid = layout.part.get_or_add_image_part(str(lock_white))
        w_, h_ = img_size(lock_white, width_emu=E(1.18), max_h=E(0.34))
        return (pic(rid, M, FOOT_Y + E(0.13) - h_ // 2, w_, h_, name='Footer logo') +
                static_text(M + w_ + E(0.3), FOOT_Y - E(0.02), E(3.0), E(0.3), K.COPYRIGHT,
                            9.5, color=hexof(K.GREEN_300), name='Copyright') + sld_num(dark=True))

    # 0 · Cover ───────────────────────────────────────────────────────────────
    L = layouts[0]
    _, rid_lw = L.part.get_or_add_image_part(str(lock_white))
    _, rid_mk = L.part.get_or_add_image_part(str(mark_white))
    cw, ch = img_size(lock_white, width_emu=E(2.05), max_h=E(0.62))
    mw, mh = img_size(mark_white, height_emu=E(5.4))
    define(L, 'Cover', (
        pic(rid_mk, W - E(4.4), H - E(4.6), mw, mh, name='Watermark', alpha=9000) +
        box(0, H - E(0.34), W, E(0.34), solid(scheme='accent1'), name='Base band') +
        pic(rid_lw, M, E(0.55), cw, ch, name='Logo') +
        ph('body', 10, 'Eyebrow', M, E(2.3), E(9.5), E(0.26), 'CLIENT · PURPOSE',
           lvl(1, 10, color=hexof(K.GREEN_300), bold=True, caps=True, spc=160), autofit=False) +
        ph('ctrTitle', None, 'Title', M, E(2.68), E(9.6), E(1.7), 'Title of the document',
           lvl(1, 46, scheme='bg1', font='+mj-lt', ln=100), anchor='t') +
        ph('subTitle', 1, 'Subtitle', M, E(4.5), E(8.6), E(0.9), 'One line on what this is and who it is for.',
           lvl(1, 16, color=hexof(K.ON_FOREST), ln=132)) +
        ph('body', 11, 'Date and meta', M, H - E(1.3), E(8), E(0.3), 'Prepared for [Client] · [Month Year]',
           lvl(1, 11, color=hexof(K.ON_FOREST_MUTED)), autofit=False)),
        bg_scheme='tx2', show_master=False, ltype='title')

    # 1 · Section ─────────────────────────────────────────────────────────────
    L = layouts[1]
    define(L, 'Section', (
        box(0, 0, E(0.26), H, solid(scheme='accent1'), name='Section bar') +
        ph('body', 10, 'Section number', M + E(0.2), E(2.05), E(2.0), E(0.9), '01',
           lvl(1, 54, color=hexof(K.GREEN_300), font='+mj-lt'), autofit=False) +
        ph('title', None, 'Title', M + E(0.2), E(3.2), E(10.5), E(1.0), 'Section heading',
           lvl(1, 38, color=hexof(K.FOREST), font='+mj-lt', ln=100)) +
        ph('body', 11, 'Lede', M + E(0.2), E(4.35), E(9.2), E(0.7),
           'One line on what this section settles.', LEDE_STYLE) + sld_num()),
        bg_hex=hexof(K.MINT_TINT), ltype='secHead')

    # 2 · Title and content (the workhorse) ───────────────────────────────────
    L = layouts[2]
    hx, y = header()
    define(L, 'Title and Content', hx + ph('body', 1, 'Body', M, y + E(0.1), BW, FOOT_Y - E(0.7) - y,
                                           'Three points that prove the title. Nine words each.', BODY_STYLES)
           + source_line(), ltype='obj')

    # 3 · Two content ─────────────────────────────────────────────────────────
    L = layouts[3]
    hx, y = header()
    gap = E(0.5)
    cw2 = (BW - gap) // 2
    body_h = FOOT_Y - E(0.7) - y
    define(L, 'Two Content', hx +
           ph('body', 1, 'Left', M, y + E(0.1), cw2, body_h, 'Left: one idea.', BODY_STYLES) +
           ph('body', 2, 'Right', M + cw2 + gap, y + E(0.1), cw2, body_h, 'Right: its counterpart.', BODY_STYLES) +
           source_line(), ltype='twoObj')

    # 4 · Title only: the canvas the engine draws cards, charts and diagrams on
    L = layouts[4]
    hx, y = header()
    define(L, 'Title Only', hx + source_line(), ltype='titleOnly')

    # 5 · Comparison ──────────────────────────────────────────────────────────
    L = layouts[5]
    hx, y = header()
    head_h = E(0.46)
    define(L, 'Comparison', hx +
           box(M, y + E(0.1), cw2, Emu(28575), solid(hexof(K.CLOUD)), name='Left rule') +
           box(M + cw2 + gap, y + E(0.1), cw2, Emu(28575), solid(scheme='accent1'), name='Right rule') +
           ph('body', 21, 'Left heading', M, y + E(0.24), cw2, head_h, 'Today',
              lvl(1, 20, color=hexof(K.INK_MUTED), font='+mj-lt'), autofit=False) +
           ph('body', 22, 'Right heading', M + cw2 + gap, y + E(0.24), cw2, head_h, f'With {K.BRAND_SHORT}',
              lvl(1, 20, color=hexof(K.FOREST), font='+mj-lt'), autofit=False) +
           ph('body', 1, 'Left', M, y + E(0.84), cw2, body_h - E(0.74), 'What it is like now.', BODY_STYLES) +
           ph('body', 2, 'Right', M + cw2 + gap, y + E(0.84), cw2, body_h - E(0.74), 'What changes.', BODY_STYLES) +
           source_line(), ltype='twoTxTwoObj')

    # 6 · Big numbers ─────────────────────────────────────────────────────────
    L = layouts[6]
    hx, y = header()
    sg = E(0.34)
    sw = (BW - sg * 2) // 3
    sh = E(2.0)
    sy = y + E(0.3)
    cards = ''
    for i in range(3):
        x = M + i * (sw + sg)
        cards += (box(x, sy, sw, sh, solid(hexof(K.MINT_TINT)), name=f'Stat card {i + 1}',
                      line_xml=f'<a:ln w="9525">{solid(hexof(K.MINT))}</a:ln>') +
                  box(x, sy, E(0.05), sh, solid(scheme='accent1'), name=f'Stat accent {i + 1}') +
                  ph('body', 15 + i, f'Value {i + 1}', x + E(0.34), sy + E(0.28), sw - E(0.6), E(0.9), '00%',
                     lvl(1, 44, color=hexof(K.FOREST), bold=True), autofit=False) +
                  ph('body', 18 + i, f'Label {i + 1}', x + E(0.34), sy + E(1.24), sw - E(0.6), E(0.6),
                     'What the number measures, and its basis.', lvl(1, 12, color=hexof(K.INK_800), ln=124)))
    define(L, 'Big Numbers', hx + cards + source_line(), ltype='cust')

    # 7 · Quote ───────────────────────────────────────────────────────────────
    L = layouts[7]
    define(L, 'Quote', (
        box(0, 0, W, E(0.26), solid(scheme='accent1'), name='Top band') +
        static_text(M, E(1.35), E(1.2), E(1.2), '“', 120, color=hexof(K.GREEN_300), name='Quote mark') +
        ph('body', 1, 'Quote', M + E(0.9), E(2.1), E(10.6), E(2.6),
           'A sentence the customer actually said, with permission.',
           lvl(1, 30, scheme='bg1', font='+mj-lt', ln=118)) +
        ph('body', 11, 'Attribution', M + E(0.9), E(4.95), E(9.0), E(0.4), 'Name, Role, Company',
           lvl(1, 13, color=hexof(K.GREEN_300)), autofit=False) +
        dark_footer(L)),
        bg_scheme='tx2', show_master=False, ltype='cust')

    # 8 · Picture and text ────────────────────────────────────────────────────
    L = layouts[8]
    hx, y = header(lede=False)
    tw = E(4.6)
    define(L, 'Picture and Text', hx +
           ph('body', 1, 'Text', M, y + E(0.1), tw, FOOT_Y - E(0.7) - y, 'Two or three points about the picture.',
              BODY_STYLES) +
           ph('pic', 14, 'Picture', M + tw + E(0.4), y + E(0.1), BW - tw - E(0.4), FOOT_Y - E(0.75) - y,
              'Screenshot or photo. Real captures beat stock.', lvl(1, 12, color=hexof(K.INK_MUTED), algn='ctr'),
              anchor='ctr') +
           source_line(), ltype='picTx')

    # 9 · Closing ─────────────────────────────────────────────────────────────
    L = layouts[9]
    _, rid_lw = L.part.get_or_add_image_part(str(lock_white))
    define(L, 'Closing', (
        box(0, 0, W, E(0.26), solid(scheme='accent1'), name='Top band') +
        pic(rid_lw, M, E(0.72), cw, ch, name='Logo') +
        ph('title', None, 'Title', M, E(2.2), E(9.6), E(1.1), 'What we need from you.',
           lvl(1, 40, scheme='bg1', font='+mj-lt', ln=100)) +
        ph('body', 1, 'Asks', M, E(3.5), E(9.0), E(2.6), 'Decisions, owners and dates. Three to five.',
           lvl(1, 15, color=hexof(K.ON_FOREST), bullet=True, ln=128, after=10)) +
        dark_footer(L)),
        bg_scheme='tx2', show_master=False, ltype='cust')

    # 10 · Blank (footer only) ────────────────────────────────────────────────
    define(layouts[10], 'Blank', sld_num(), ltype='blank')

    # ── theme ─────────────────────────────────────────────────────────────────
    brand.apply_theme_to_package(prs.part.package)
    return prs


# ── example slides: the template teaches by showing ─────────────────────────
def fill(slide, values):
    """values: {idx or 'title': text | [paragraphs]}. Unfilled placeholders are removed
       so edit mode is clean; the layout still offers them on a new slide."""
    for shape in list(slide.placeholders):
        key = 'title' if shape.placeholder_format.type in (1, 3) else shape.placeholder_format.idx
        if key in values:
            v = values[key]
            if v is None:
                continue
            tf = shape.text_frame
            paras = v if isinstance(v, list) else [v]
            tf.text = paras[0]
            for extra in paras[1:]:
                tf.add_paragraph().text = extra
        elif shape.placeholder_format.idx != 12:
            shape._element.getparent().remove(shape._element)


def add_slide_number(slide):
    """Instantiate the sldNum placeholder the layout defines."""
    layout_ph = next((p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == 12), None)
    if layout_ph is None:
        return
    el = copy.deepcopy(layout_ph._element)
    body = el.find(qn('p:txBody'))
    for p_ in body.findall(qn('a:p')):
        body.remove(p_)
    body.append(etree.fromstring(
        f'<a:p {NS}><a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
        f'<a:rPr lang="en-US"/><a:t>‹#›</a:t></a:fld></a:p>'))
    slide.shapes._spTree.append(el)


def placeholder_capture():
    """A neutral wireframe standing in for a product screenshot. The template is
       public; a real capture can carry customer data, and one already has."""
    from PIL import Image, ImageDraw
    out = K.Path.home() / '.cache' / 'servantium-documents' / f'placeholder-capture-{K.ACTIVE.id}.png'
    out.parent.mkdir(parents=True, exist_ok=True)
    W_, H_ = 1600, 1000
    im = Image.new('RGB', (W_, H_), K.T.COLORS['color-mist-grey'])
    d = ImageDraw.Draw(im)
    rule, card, ink = K.T.COLORS['color-cloud-grey'], '#FFFFFF', K.T.COLORS['color-green']
    d.rectangle((0, 0, W_, 70), fill=card); d.line((0, 70, W_, 70), fill=rule, width=2)
    d.rectangle((0, 70, 250, H_), fill=card); d.line((250, 70, 250, H_), fill=rule, width=2)
    for i in range(6):
        d.rounded_rectangle((30, 110 + i * 56, 210, 140 + i * 56), 6, fill=rule if i else ink)
    d.rounded_rectangle((300, 110, 1550, 190), 10, fill=card, outline=rule, width=2)
    for r in range(9):
        y = 230 + r * 76
        d.rounded_rectangle((300, y, 1550, y + 60), 8, fill=card, outline=rule, width=2)
        d.rounded_rectangle((330, y + 22, 700 - (r % 3) * 90, y + 38), 5, fill=rule)
        d.rounded_rectangle((1300, y + 22, 1510, y + 38), 5, fill=ink if r % 4 == 0 else rule)
    im.save(out)
    return out


def examples(prs):
    L = {l.name: l for l in prs.slide_masters[0].slide_layouts}

    def new(name, values, number=True):
        s = prs.slides.add_slide(L[name])
        fill(s, values)
        if number:
            add_slide_number(s)
        return s

    new('Cover', {10: f'{K.BRAND_SHORT.upper()} · TEMPLATE', 'title': f'{K.BRAND_SHORT} slides template',
                  1: 'Every layout, the method behind it, and an icon library. Delete these slides and keep the layouts.',
                  11: 'Version 1 · September 2026'}, number=False)

    new('Title and Content', {
        10: 'How to use this template',
        'title': 'Write the titles first. The deck is done when they tell the story.',
        11: 'The consulting method, in five rules. Everything else in this template supports them.',
        1: ['Storyboard before you build. List the titles, agree them, then make slides.',
            'Every title is a conclusion. Read titles alone and the argument holds.',
            'One message per slide. The body only proves the title.',
            'Light slides, full notes. Nine words a line, the talk lives in the notes.',
            'Source every number. If it is inferred, say probably and give the basis.'],
        13: f'Source: {K.BRAND_SHORT} content playbook.'})

    new('Section', {10: '01', 'title': 'Section heading', 11: 'One line on what this section settles.'})

    new('Two Content', {
        10: 'Layout · Two content', 'title': 'Two ideas side by side, each proving the title.',
        1: ['Left column holds one idea.', 'Keep both columns the same length.'],
        2: ['Right column holds its counterpart.', 'Use Comparison when it is before and after.'],
        13: 'Source: example.'})

    new('Comparison', {
        10: 'Layout · Comparison', 'title': 'Before and after, with the change stated in the title.',
        21: 'Today', 22: f'With {K.BRAND_SHORT}',
        1: ['Scope lives in a document.', 'The plan is rebuilt by hand.'],
        2: ['Scope is a record.', 'The plan is derived from the quote.'],
        13: 'Source: example.'})

    s = new('Title Only', {10: 'Layout · Title only',
                           'title': 'The canvas the engine draws cards, charts and diagrams on.',
                           11: 'Use this when the visual carries the slide. The engine fills the space below.'})
    from .. import blocks as B
    B.icon_cards(s, E(2.62), [
        ('target', '01', 'Action title', ['States the conclusion', 'Not the topic']),
        ('layers', '02', 'One message', ['The body proves the title', 'Nothing else']),
        ('fact_check', '03', 'Sourced', ['Every number has a source', 'Inferred means probably']),
    ], h=Inches(2.3))

    new('Big Numbers', {10: 'Layout · Big numbers', 'title': 'Three numbers, each with its basis underneath.',
                        15: '00%', 16: '00×', 17: '00',
                        18: 'What it measures. Never an invented figure.',
                        19: 'What it measures, and the period.', 20: 'What it measures, and who reported it.',
                        13: 'Source: name the source for each number.'})

    new('Quote', {1: 'A sentence the customer actually said, used with permission.',
                  11: 'Name, Role, Company'})

    s = prs.slides.add_slide(L['Picture and Text'])
    pic_ph = next(p for p in s.placeholders if p.placeholder_format.idx == 14)
    pic_ph.insert_picture(str(placeholder_capture()))
    fill(s, {10: 'Layout · Picture and text',
             'title': 'Real screenshots, not stock, with the point in the title.',
             1: ['What to notice in the picture.', 'Why it matters.'],
             13: f'Source: {K.BRAND_SHORT} product capture, 2026.', 14: None})
    add_slide_number(s)

    new('Closing', {'title': 'What we need out of today.',
                    1: ['A decision, with an owner.', 'A date.', 'The next conversation.']})

    reference_slides(prs, L)


def reference_slides(prs, L):
    """Colour and type reference, then the icon library. Both are for humans:
       copy an icon, check a hex, without opening the repository."""
    from .. import blocks as B
    from ..kit import rect, text, icon, block_h
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    s = prs.slides.add_slide(L['Title Only'])
    fill(s, {10: 'Reference', 'title': 'Colors and type.',
             11: 'The theme carries all of this. Pick from the theme colors, not the custom picker.'})
    add_slide_number(s)
    names = ((K.BRAND_SHORT + ' green', 'Forest', 'Mint') if K.HOUSE else ('Primary', 'Dark', 'Tint'))
    sw_ = [(names[0], K.GREEN, 'Accent 1 · primary'), (names[1], K.FOREST, 'Text 2 · dark ground'),
           (names[2], K.MINT, 'Background 2 · tint'), ('Slate', K.SLATE, 'Accent 3 · data'),
           ('Teal', K.TEAL, 'Accent 4 · data'), ('Amber', K.AMBER, 'Accent 5 · caution'),
           ('Coral', K.CORAL, 'Accent 6 · risk'), ('Ink', K.INK_800, 'Body text')]
    gx, gy, gw, gh, gg = K.MARGIN, E(2.72), E(1.38), E(1.02), E(0.12)
    for i, (nm, col, role) in enumerate(sw_):
        x = gx + i * (gw + gg)
        rect(s, x, gy, gw, gh, fill=col, line=K.CLOUD)
        text(s, x, gy + gh + E(0.1), gw, E(0.22), [(nm, dict(size=10.5, bold=True, color=K.FOREST))])
        text(s, x, gy + gh + E(0.34), gw, E(0.2), [(f'#{hexof(col)}', dict(size=9.5, color=K.INK_MUTED))])
        text(s, x, gy + gh + E(0.54), gw, E(0.2), [(role, dict(size=9, color=K.INK_MUTED))])
    ty = E(4.95)
    text(s, K.MARGIN, ty, E(6), E(0.6), [(K.DISPLAY, dict(font=K.DISPLAY, size=32, color=K.FOREST))])
    text(s, K.MARGIN, ty + E(0.66), E(6), E(0.3),
         [('Titles and big numbers. Heading font in the theme.', dict(size=11.5, color=K.INK_MUTED))])
    text(s, E(7.0), ty + E(0.08), E(5.6), E(0.5), [(K.UI, dict(size=28, color=K.INK_800))])
    text(s, E(7.0), ty + E(0.66), E(5.6), E(0.3),
         [('Everything else. Body font in the theme.', dict(size=11.5, color=K.INK_MUTED))])

    # Icon library. Material Symbols Rounded (Google) and Lucide, reconciled at
    # wght 600 / 1.09x by kit. PNG, not SVG, because Google Slides does not
    # accept SVG; Microsoft users can take vectors from assets/icons/svg/.
    ICONS = ['target', 'rocket', 'flag', 'trophy', 'lightbulb', 'sparkles', 'zap', 'shield-check',
             'users', 'groups', 'user-check', 'handshake', 'building-2', 'briefcase', 'badge', 'school',
             'calendar-check', 'timer', 'hourglass', 'alarm-clock', 'refresh-cw', 'autorenew', 'repeat', 'history',
             'file-text', 'file-pen', 'clipboard-check', 'checklist', 'fact_check', 'request_quote', 'receipt_long', 'description',
             'trending-up', 'monitoring', 'query_stats', 'insights', 'pie_chart', 'gauge', 'scale', 'calculate',
             'database', 'hub', 'link', 'plug', 'workflow', 'account_tree', 'layers', 'git-merge',
             'message-square', 'forum', 'send', 'bell', 'mail', 'phone', 'video', 'globe',
             'search', 'filter', 'settings', 'lock', 'eye', 'circle-check', 'triangle-alert', 'help']
    ICONS = [n for n in ICONS if K.icon_exists(n)]
    per = 32
    for page in range(0, len(ICONS), per):
        chunk = ICONS[page:page + per]
        s = prs.slides.add_slide(L['Title Only'])
        fill(s, {10: 'Reference · Icon library',
                 'title': 'Copy an icon. Keep them one color and one size per slide.',
                 11: 'Google Material Symbols and Lucide, weight-matched so they read as one set.'})
        add_slide_number(s)
        cols, cell_w, cell_h = 8, BW // 8, E(0.86)
        y0 = E(2.62)
        for j, name in enumerate(chunk):
            c, r = j % cols, j // cols
            x = M + c * cell_w
            yy = y0 + r * cell_h
            icon(s, name, x + (cell_w - E(0.42)) // 2, yy, E(0.42), color=K.GREEN_600)
            text(s, x, yy + E(0.48), cell_w, E(0.2), [(name, dict(size=8, color=K.INK_MUTED))],
                 align=PP_ALIGN.CENTER)


def save(prs):
    pptx_path = K.TEMPLATE_PPTX
    potx_path = brand.DIST / f'{K.BRAND_SHORT}.potx'
    prs.save(str(pptx_path))
    # .potx is the same package with the template content type. PowerPoint then
    # opens it as "new from template" instead of editing the file itself.
    with zipfile.ZipFile(pptx_path) as zin, zipfile.ZipFile(potx_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(
                    b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',
                    b'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml')
            zout.writestr(item, data)
    return pptx_path, potx_path


if __name__ == '__main__':
    prs = build()
    examples(prs)
    a, b = save(prs)
    names = [l.name for l in prs.slide_masters[0].slide_layouts]
    print(f'layouts ({len(names)}): {", ".join(names)}')
    print(f'slides: {len(prs.slides)} -> {a.name}, {b.name}')
