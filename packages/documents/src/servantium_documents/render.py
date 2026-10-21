"""Spec -> PowerPoint, on the real layouts of the Servantium template.

The renderer does two jobs that used to be done by eye:

* OVERFLOW. It knows the y below every component, so content running into the
  source line or footer is an error with a measurement, not something found in the
  PNG later.
* BALANCE. A slide that ends high leaves a dead band at the bottom (every slide of a
  client status deck needed a hand fix for this). If the slack is large, the body is redrawn
  lower, by a bounded amount, so slides sit consistently.
"""
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx.util import Inches, Pt

from . import __version__, kit, blocks as B, package
from .check import Finding
from .spec import slides

TONE = {'green': kit.GREEN, 'amber': kit.AMBER, None: kit.GREEN}
GAP = Inches(0.3)
BALANCE_FROM = Inches(1.0)        # slack above this triggers a redraw
BALANCE_MAX = Inches(0.9)         # never push the body down more than this


# ── components ───────────────────────────────────────────────────────────────
def _takeaway(s, y, v, ctx):
    runs = [(v['lead'] + (' ' if v.get('text') else ''), dict(size=12.5, bold=True, color=kit.GREEN_300))]
    if v.get('text'):
        runs.append((v['text'], dict(size=12.5, color=kit.WHITE)))
    return B.strip(s, y, runs)


def _callout(s, y, v, ctx):
    tone = kit.AMBER if v.get('tone', 'amber') == 'amber' else kit.GREEN
    return B.callout(s, y, v['label'], [(v['text'], dict(size=12.5, color=kit.WHITE))],
                     tone=tone, icon_name=v.get('icon', 'help'))


def _band(s, y, v, ctx):
    runs = []
    if v.get('lead'):
        runs.append((v['lead'] + ' ', dict(size=12, bold=True, color=kit.GREEN_300)))
    if v.get('text'):
        runs.append((v['text'], dict(size=12, color=B.MIST_INK)))
    return B.band(s, y, v['heading'], runs or [(' ', dict(size=12))])


def _cards(s, y, v, ctx):
    items = [(c.get('icon'), c.get('kicker', f'{i + 1:02d}'), c['title'], c.get('lines', []))
             for i, c in enumerate(v)]
    return B.icon_cards(s, y, items, h=Inches(2.0 if all(len(c.get('lines', [])) <= 1 for c in v) else 2.5))


def _tiles(s, y, v, ctx):
    return B.tiles(s, y, [(t['icon'], t['label']) for t in v], cols=len(v), h=Inches(1.3))


def _stats(s, y, v, ctx):
    return B.stat_row(s, y, [(t['value'], t['label'], t.get('note', '')) for t in v])


def _steps(s, y, v, ctx):
    return B.steps_row(s, y, [(t['lead'], t['text']) for t in v])


def _timeline(s, y, v, ctx):
    return B.timeline(s, y, [(t['icon'], t['when'], t.get('focus_label', 'Focus'), t['focus'],
                              t.get('moves_label', 'Moves'), t['moves']) for t in v], h=Inches(2.6))


def _chevrons(s, y, v, ctx):
    return B.chevron_stages(s, y, [(t['stage'], t['question'], t['text']) for t in v])


def _milestones(s, y, v, ctx):
    return B.milestones(s, y, [(t.get('icon'), t['when'], t['title'], t['text'], t.get('done', False)) for t in v])


def _panels(s, y, v, ctx):
    p = [(t.get('icon'), t['title'], t.get('tag'), TONE[t.get('tone', 'green')], t['lines']) for t in v]
    return B.two_panel(s, y, p[0], p[1])


def _compare(s, y, v, ctx):
    return B.compare(s, y, (v['before']['label'], v['before']['lines'], v['before'].get('icon')),
                     (v['after']['label'], v['after']['lines'], v['after'].get('icon')))


def _shipped(s, y, v, ctx):
    return B.shipped(s, y, [(g['title'], [(i.get('icon'), i['label']) for i in g['items']]) for g in v['groups']])


def _chain(s, y, v, ctx):
    return B.chain(s, y, [(t['icon'], t['title'], t['text']) for t in v])


def _logos(s, y, v, ctx):
    items = [(t['logo'] if t['logo'] in B.HOUSE_KEYS else str(ctx.resolve(t['logo'])), t['title'], t['text']) for t in v]
    return B.logo_row(s, y, items)


def _people(s, y, v, ctx):
    return B.people(s, y, [(str(ctx.resolve(t['photo'])), t['name'], t['role'], t.get('lines', []),
                            [str(ctx.resolve(l)) for l in t.get('logos', [])]) for t in v],
                    workdir=ctx.workdir)


def _adoption(s, y, v, ctx):
    return B.adoption(s, y + Inches(0.1), v['here'], labels=v.get('labels'), note=v.get('note'))


def _approach(s, y, v, ctx):
    return B.approach(s, y, v['stages'], [(r['label'], r['cells'], r.get('tone')) for r in v['rows']],
                      active=v.get('active'))


def _zoom(s, y, v, ctx):
    return B.zoom(s, y, v['steps'], v['active'], v['title'],
                  [(r.get('key'), r['title'], r.get('text')) for r in v['rows']],
                  meta=v.get('meta'), foot=v.get('foot'))


def _hub(s, y, v, ctx):
    box = lambda b: (b['title'], b.get('text'), b.get('icon')) if b else None
    return B.hub_spoke(s, y + (Inches(0.0) if v.get('top') else Inches(0.34)),
                       [(x['name'], str(ctx.resolve(x['logo'])) if x.get('logo') else x.get('icon')) for x in v['left']],
                       [(x['name'], str(ctx.resolve(x['logo'])) if x.get('logo') else x.get('icon')) for x in v['right']],
                       left_label=v.get('left_label'), right_label=v.get('right_label'),
                       top=box(v.get('top')), bottom=box(v.get('bottom')), caption=v.get('caption'))


def _table(s, y, v, ctx):
    n = len(v['head'])
    widths = v.get("widths") or [kit.BODY_W / 914400 / n] * n
    scale = (kit.BODY_W / 914400) / sum(widths)
    extra = {}
    if v.get('row_h'):
        extra['row_h'] = Inches(v['row_h'])
    if v.get('pills'):
        extra['pills'] = {p['column'] - 1: p['tones'] for p in v['pills']}      # 1-based in the spec
    return B.table(s, y, v['head'], v['rows'], [Inches(w * scale) for w in widths], **extra)


def _quote(s, y, v, ctx):
    return B.pull_quote(s, y, v['text'], v['attribution'])


def _image(s, y, v, ctx):
    path = ctx.resolve(v['path'])
    h = Inches(v.get('height', 3.6))
    framed = v.get('frame', True)  # charts and drawn pictures sit flat; screenshots keep the frame
    pic = kit.picture(s, path, kit.MARGIN, y, h=h, border=framed, shadow=framed)
    if pic.width > kit.BODY_W:
        pic.width, pic.height = kit.BODY_W, int(pic.height * kit.BODY_W / pic.width)
    pic.left = kit.MARGIN + (kit.BODY_W - pic.width) // 2
    return y + pic.height


def _diagram(s, y, v, ctx):
    png = ctx.mermaid(v['mermaid'])
    h = Inches(v.get('height', 3.8))
    pic = kit.picture(s, png, kit.MARGIN, y, h=h, border=False, shadow=False)
    if pic.width > kit.BODY_W:
        pic.width, pic.height = kit.BODY_W, int(pic.height * kit.BODY_W / pic.width)
    pic.left = kit.MARGIN + (kit.BODY_W - pic.width) // 2
    return y + pic.height


def _chart(s, y, v, ctx):
    """A native, editable chart in theme colors. Values stay editable in PowerPoint;
       whether Google keeps them editable is one of the round-trip questions."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    kind = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
            'stacked': XL_CHART_TYPE.COLUMN_STACKED, 'line': XL_CHART_TYPE.LINE_MARKERS}[v['type']]
    data = CategoryChartData()
    data.categories = v['categories']
    for ser in v['series']:
        data.add_series(ser['name'], ser['values'])
    h = Inches(v.get('height', 3.4))
    gf = s.shapes.add_chart(kind, kit.MARGIN, y, kit.BODY_W, h, data)
    ch = gf.chart
    ch.has_title = False          # the action title says it; charts carry no title of their own
    ch.font.name, ch.font.size = kit.UI, Pt(11)
    ch.font.color.rgb = kit.INK_800
    palette = [kit.GREEN, kit.FOREST, kit.SLATE, kit.TEAL]
    for i, plot_series in enumerate(ch.plots[0].series):
        fmt = plot_series.format
        if v['type'] == 'line':
            fmt.line.color.rgb = palette[i]; fmt.line.width = Pt(2.5)
        else:
            fmt.fill.solid(); fmt.fill.fore_color.rgb = palette[i]
    ch.has_legend = len(v['series']) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    if v['type'] in ('column', 'bar') and len(v['series']) == 1:
        plot = ch.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(10.5)
        plot.data_labels.number_format = v.get('format') or ('0' + (v.get('unit') or ''))
        plot.data_labels.number_format_is_linked = False
        plot.gap_width = 60
    va = ch.value_axis
    if v['type'] in ('column', 'bar') and len(v['series']) == 1:
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        va.tick_label_position = XL_TICK_LABEL_POSITION.NONE   # the bars carry their own labels
    if v.get('format'):                      # the value axis prints like the labels ($#,##0, 0%)
        va.tick_labels.number_format = v['format']
        va.tick_labels.number_format_is_linked = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = kit.CLOUD
    va.format.line.fill.background()
    ch.category_axis.format.line.color.rgb = kit.CLOUD
    return y + h


def _gantt(s, y, v, ctx):
    rows = [(r['name'], r.get('note'), [(b['from'], b['to'], b['label'], b.get('tone')) for b in r['bars']])
            for r in v['rows']]
    room = _limit_for(ctx) - y - getattr(ctx, 'reserve', 0) - (Inches(0.12) if getattr(ctx, 'reserve', 0) else 0)
    return B.gantt(s, y, v['periods'], rows, [(g['after'], g['label']) for g in v.get('gates', [])],
                   max_h=max(Inches(2.0), min(Inches(4.4), room)))


def _flow(s, y, v, ctx):
    o = v['output']
    return B.flow(s, y, v.get('steps'), v.get('active', 0),
                  [(str(ctx.resolve(i['logo'])) if i.get('logo') else i.get('icon'), i['title'], i['meta'],
                    i.get('tag'), i.get('carried', False)) for i in v['inputs']],
                  (v['node']['title'], v['node']['verb'], v['node'].get('chip'), v['node'].get('cycle')),
                  (o['label'], o.get('ref'), o['title'],
                   [(r['key'], r['text'], r.get('open', False)) for r in o['rows']], o.get('foot')))


def _swimlane(s, y, v, ctx):
    keys = [ln['key'] for ln in v['lanes']]
    idx = {k: i for i, k in enumerate(keys)}
    steps = []
    for st in v['steps']:
        lanes = st.get('lanes') or [st.get('lane')]
        missing = [k for k in lanes if k not in idx]
        if missing:
            raise ValueError(f'step {st["id"]}: unknown lane {missing}; lanes are {keys}')
        steps.append({**st, 'lanes': [idx[k] for k in lanes]})
    return B.swimlane(s, y, [(ln['name'], ln.get('note')) for ln in v['lanes']], steps,
                      [(lp['from'], lp['to'], lp.get('label')) for lp in v.get('loops', [])],
                      [(ph['label'], ph['from'], ph['to']) for ph in v.get('phases', [])],
                      v.get('clock'), max_h=_limit_for(ctx) - y)


def _staircase(s, y, v, ctx):
    return B.staircase(s, y, [(st.get('label'), st['title'], st.get('text'), st.get('when')) for st in v['steps']],
                       here=v.get('here'), here_label=v.get('here_label', 'Start here'),
                       numbered=v.get('numbered', True))


def _questions(s, y, v, ctx):
    return B.questions(s, y, [(q['q'], q.get('why')) for q in v['items']], note=v.get('note'))


def _middleout(s, y, v, ctx):
    side = lambda d: (d['label'], [(str(ctx.resolve(i['logo'])) if i.get('logo') else i.get('icon'), i['name'],
                                    i.get('text'), i.get('when')) for i in d['items']])
    c = v['center']
    return B.middleout(s, y, (c['title'], c['items'], c.get('note')), side(v['left']), side(v['right']))


def _processmap(s, y, v, ctx):
    # A replaced tool is a logo file in the workspace or, when no such file exists, an icon name.
    def res(ps):
        out = []
        for p in ps or []:
            f = Path(ctx.workdir) / p
            out.append(str(ctx.resolve(p)) if ('/' in p or '.' in p or f.exists()) else str(kit.icon_png(p, color=kit.SLATE, px=128)))
        return out
    stages = [dict(st, replaces=res(st.get('replaces'))) for st in v['stages']]
    groups = [(g['label'], g['from'], g['to'], g.get('tone', 'sales')) for g in v.get('groups', [])]
    return B.processmap(s, y, stages, groups, loop=v.get('loop'))


def _board(s, y, v, ctx):
    cols = [(c['title'], [(k['title'], k.get('meta'), k.get('score'), k.get('tag')) for k in c.get('cards', [])])
            for c in v['columns']]
    return B.board(s, y, cols, sources=v.get('sources'), sources_label=v.get('sources_label', 'Sources'))


def _catalog(s, y, v, ctx):
    return B.catalog(s, y, [(g['title'], g.get('tag'), g.get('tone'), g['items']) for g in v['groups']])


def _spectrum(s, y, v, ctx):
    return B.spectrum(s, y, (v['left'], v['right']),
                      [(st['title'], st.get('text'), st.get('pick', False)) for st in v['stops']],
                      pick_label=v.get('pick_label', 'Our approach'))


def _options(s, y, v, ctx):
    return B.options(s, y, [(i.get('label'), i['title'], i.get('meta'), i.get('lines', []), i.get('pick', False))
                            for i in v['items']], pick_label=v.get('pick_label', 'Recommended'))


def _matrix(s, y, v, ctx):
    n = len(v['columns'])
    for r in v['rows']:
        if len(r['cells']) != n or (r.get('levels') and len(r['levels']) != n):
            raise ValueError(f'matrix row "{r["label"]}": {n} columns need {n} cells (and {n} levels if given)')
    return B.matrix(s, y, v['columns'], [(r['label'], r['cells'], r.get('levels')) for r in v['rows']],
                    highlight=v.get('highlight'), pick_label=v.get('pick_label', 'Recommended'))


def _cap(s, y, v, ctx):
    side = lambda d: (d['label'], d.get('tone', 'warn'),
                      [(b['label'], b.get('amount'), b.get('value')) for b in d['blocks']], d.get('lines', []))
    return B.cap(s, y, side(v['left']), side(v['right']), (v['cap']['value'], v['cap']['label']))


def _weekgrid(s, y, v, ctx):
    keys = {st['key'] for st in v['stages']}
    n = len(v['periods'])
    for r in v['rows']:
        for sp in r['spans']:
            if sp['stage'] not in keys:
                raise ValueError(f'weekgrid row "{r["name"]}": unknown stage {sp["stage"]!r}; stages are {sorted(keys)}')
            if not 1 <= sp['from'] <= sp['to'] <= n:
                raise ValueError(f'weekgrid row "{r["name"]}": span {sp["from"]}-{sp["to"]} is outside 1-{n}')
    room = _limit_for(ctx) - y - getattr(ctx, 'reserve', 0) - (Inches(0.12) if getattr(ctx, 'reserve', 0) else 0)
    return B.weekgrid(s, y, v['periods'], [(st['key'], st['label'], st.get('short'), st.get('tone')) for st in v['stages']],
                      [(r['name'], r.get('note'), [(sp['from'], sp['to'], sp['stage'], sp.get('label')) for sp in r['spans']]) for r in v['rows']],
                      marks=[(m['at'], m['label']) for m in v.get('marks', [])],
                      max_h=max(Inches(2.2), min(Inches(4.6), room)))


def _roadmap(s, y, v, ctx):
    n = len(v['periods'])
    for ln in v['lanes']:
        if len(ln['cells']) != n:
            raise ValueError(f'roadmap lane "{ln["name"]}": {n} periods need {n} cells (use [] for an empty one)')
    return B.roadmap(s, y, v['periods'], [(ln['name'], ln.get('tone'), ln['cells']) for ln in v['lanes']])


def _tiers(s, y, v, ctx):
    return B.tiers(s, y, [(i.get('label'), i['title'], i['price'], i.get('unit'), i.get('meta'), i.get('lines', []),
                           i.get('billing'), i.get('pick', False)) for i in v['items']],
                   pick_label=v.get('pick_label', 'Recommended'))


def _faq(s, y, v, ctx):
    return B.faq(s, y, [(i['q'], i['a']) for i in v['items']])


def _ledger(s, y, v, ctx):
    return B.ledger(s, y, [(sec['title'], [(l['item'], l.get('basis'), l['amount']) for l in sec['lines']],
                            (sec['total']['label'], sec['total']['amount']) if sec.get('total') else None)
                           for sec in v['sections']])


def _ring(s, y, v, ctx):
    room = _limit_for(ctx) - y - getattr(ctx, 'reserve', 0) - (Inches(0.12) if getattr(ctx, 'reserve', 0) else 0)
    c = v['center']
    return B.ring(s, y, (c['title'], c.get('text')), [(st['title'], st.get('text')) for st in v['stages']],
                  h=max(Inches(2.6), min(Inches(3.5), room)))


def _limit_for(ctx):
    return ctx.limit if getattr(ctx, 'limit', None) else B.SAFE_BOTTOM


def _stack(s, y, v, ctx):
    f = v['foundation']
    return B.stack(s, y, [(t['label'], t.get('note'),
                           [(i['name'], str(ctx.resolve(i['logo'])) if i.get('logo') else i.get('icon'))
                            for i in t['items']])
                          for t in v['tiers']], (f['title'], f.get('text'), f['chips']))


COMPONENTS = {
    'gantt': _gantt, 'flow': _flow, 'stack': _stack, 'shipped': _shipped, 'swimlane': _swimlane,
    'chart': _chart,
    'takeaway': _takeaway, 'callout': _callout, 'band': _band, 'cards': _cards, 'tiles': _tiles,
    'stats': _stats, 'steps': _steps, 'timeline': _timeline, 'chevrons': _chevrons,
    'milestones': _milestones, 'panels': _panels, 'compare': _compare, 'chain': _chain,
    'logos': _logos, 'people': _people, 'hub': _hub, 'adoption': _adoption, 'approach': _approach,
    'zoom': _zoom, 'staircase': _staircase, 'questions': _questions, 'middleout': _middleout, 'processmap': _processmap, 'board': _board, 'catalog': _catalog, 'spectrum': _spectrum, 'options': _options, 'matrix': _matrix,
    'cap': _cap, 'weekgrid': _weekgrid, 'roadmap': _roadmap, 'tiers': _tiers, 'faq': _faq, 'ledger': _ledger, 'ring': _ring,
    'table': _table, 'quote': _quote, 'image': _image, 'diagram': _diagram,
}


# ── context: paths, mermaid, findings ────────────────────────────────────────
class Context:
    def __init__(self, spec, workdir):
        self.spec = spec
        self.base = Path(spec['_path']).parent
        self.workdir = Path(workdir)
        self.findings = []

    def resolve(self, p):
        path = Path(os.path.expanduser(p))
        path = path if path.is_absolute() else self.base / path
        if not path.exists():
            raise FileNotFoundError(f'{p} not found (looked in {path.parent})')
        return path

    def mermaid(self, src):
        """A path to a .mmd file, or inline Mermaid, drawn by mermaid.py exactly as `sv build`
           draws a standalone diagram: brand theme, brand classes, brand fonts."""
        import yaml as _y
        from . import mermaid as MM
        out_dir = self.workdir / 'diagrams'
        out_dir.mkdir(parents=True, exist_ok=True)
        if src.strip().endswith('.mmd') and '\n' not in src.strip():
            text_ = self.resolve(src.strip()).read_text()
        else:
            text_ = src
        key = f'{kit.ACTIVE.id}-{abs(hash(text_)) % 10**8}'
        spec = out_dir / f'{key}.yaml'
        spec.write_text(_y.safe_dump({'mermaid': key, 'client': 'deck', 'source': text_, 'frame': False}))
        return MM.build(spec, out_dir=out_dir)['png']


def _find_mmdc():
    candidates = [os.environ.get('SV_MMDC'), shutil.which('mmdc'),
                  str(kit.REPO_PACKAGES.parent / 'node_modules' / '.bin' / 'mmdc'),
                  str(Path.home() / 'Projects/Servantium/decks/node_modules/.bin/mmdc')]
    for c in candidates:
        if c and Path(c).exists():
            return c
    raise FileNotFoundError('mermaid-cli (mmdc) not found. Install it or set SV_MMDC.')


# ── slides ───────────────────────────────────────────────────────────────────
def _eyebrow(spec, f):
    """'Status' becomes 'Acme · Status'. Workspaces whose name starts with '_' (the sandbox,
       internal decks) are folders, not audiences, so they never prefix themselves."""
    e, client = f.get('eyebrow'), spec['client']
    if client.startswith('_'):
        return e
    if e is None:
        return client
    return e if ('·' in e or e.upper() == client.upper()) else f'{client} · {e}'


def _limit(f):
    """Where content must stop: above the source line if there is one."""
    return kit.FOOTER_Y - Inches(0.64) if f.get('source') else B.SAFE_BOTTOM


# what a trailing emphasis block needs below a picture (its gap included), so a picture that
# can shrink (a Gantt's rows) leaves room for it instead of pushing it past the footer
_EMPHASIS_H = {'takeaway': Inches(0.34 + 0.62), 'callout': Inches(0.34 + 0.72),
               'band': Inches(0.34 + 1.1), 'quote': Inches(0.3 + 2.4)}


def _draw_body(s, y, body, ctx):
    for i, comp in enumerate(body):
        (name, val), = comp.items()
        if i:
            y += Inches(0.34) if name in ('takeaway', 'callout', 'band') else GAP
        ctx.reserve = sum(_EMPHASIS_H.get(next(iter(c)), 0) for c in body[i + 1:])
        y = COMPONENTS[name](s, y, val, ctx)
    return y


def _content(prs, n, f, ctx):
    s, y = kit.content(prs, eyebrow_text=_eyebrow(ctx.spec, f), heading=f['title'], sub=f.get('lede'),
                       source=f.get('source'))
    tree = s.shapes._spTree
    mark = len(tree)
    ctx.limit = _limit(f)
    end = _draw_body(s, y, f['body'], ctx)
    limit = _limit(f)
    if end > limit:
        over = (end - limit) / 914400
        ctx.findings.append(Finding('error', 'OVERFLOW', n,
                                    f'content runs {over:.2f} in past the {"source line" if f.get("source") else "footer"}. '
                                    'Cut text, use fewer items, or split the slide.'))
    else:
        slack = limit - end
        if slack > BALANCE_FROM:
            for el in list(tree)[mark:]:
                tree.remove(el)
            _draw_body(s, y + min(int(slack * 0.5), BALANCE_MAX), f["body"], ctx)
    return s


def _move_body(s, y, idxs, f):
    bottom = _limit(f)
    for idx in idxs:
        ph = kit.placeholder(s, idx)
        if ph is not None:
            kit._place(ph, y - Inches(0.12), bottom - (y - Inches(0.12)))


def _bullets(prs, n, f, ctx):
    s, y = kit.content(prs, eyebrow_text=_eyebrow(ctx.spec, f), heading=f['title'], sub=f.get('lede'),
                       source=f.get('source'), layout_name='Title and Content', extra={kit.PH_BODY: f['points']})
    _move_body(s, y, [kit.PH_BODY], f)
    return s


def _columns(prs, n, f, ctx):
    s, y = kit.content(prs, eyebrow_text=_eyebrow(ctx.spec, f), heading=f['title'], sub=f.get('lede'),
                       source=f.get('source'), layout_name='Two Content',
                       extra={kit.PH_BODY: f['left'], kit.PH_BODY2: f['right']})
    _move_body(s, y, [kit.PH_BODY, kit.PH_BODY2], f)
    return s


def _comparison(prs, n, f, ctx):
    s, _ = kit.content(prs, eyebrow_text=_eyebrow(ctx.spec, f), heading=f['title'], sub=f.get('lede'),
                       source=f.get('source'), layout_name='Comparison', measured=False,
                       extra={21: f['left']['heading'], 22: f['right']['heading'],
                              kit.PH_BODY: f['left']['points'], kit.PH_BODY2: f['right']['points']})
    return s


def _numbers(prs, n, f, ctx):
    extra = {}
    for i, st in enumerate(f['stats']):
        extra[15 + i], extra[18 + i] = st['value'], st['label']
    s, _ = kit.content(prs, eyebrow_text=_eyebrow(ctx.spec, f), heading=f['title'], sub=f.get('lede'),
                       source=f.get('source'), layout_name='Big Numbers', measured=False, extra=extra)
    return s


def _picture(prs, n, f, ctx):
    s = prs.slides.add_slide(kit.layout(prs, 'Picture and Text'))
    ph = kit.placeholder(s, kit.PH_PICTURE)
    ph.insert_picture(str(ctx.resolve(f['image'])))
    vals = {'title': f['title'], kit.PH_BODY: f['points'], kit.PH_PICTURE: None}
    if _eyebrow(ctx.spec, f):
        vals[kit.PH_EYEBROW] = _eyebrow(ctx.spec, f).upper()
    if f.get('source'):
        src = f['source']
        vals[kit.PH_SOURCE] = src if src.lower().startswith('source') else f'Source: {src}'
    kit.fill(s, vals)
    kit.add_slide_number(s)
    return s


def _cover(prs, n, f, ctx):
    pf = ctx.spec.get('prepared_for')
    pf = (pf['name'], ctx.resolve(pf['logo']) if pf.get('logo') else None) if pf else None
    return kit.cover(prs, eyebrow_text=f.get('eyebrow') or ctx.spec['client'], headline=f['title'],
                     subhead=f.get('subtitle', ''), meta=f.get('meta'), prepared_for=pf)


def _section(prs, n, f, ctx):
    """A divider that doubles as the agenda tracker: every section in the deck is listed
       along the bottom with this one lit, so the room always knows where it is."""
    every = list(slides(ctx.spec))
    close = next((m for m, k, _ in every if k == 'closing'), None)
    if close is not None and n > close:       # an appendix divider: not part of the agenda
        return kit.section(prs, number=f.get('number') or 'A', heading=f['title'], sub=f.get('lede'))
    secs = [(m, g) for m, k, g in every if k == 'section' and (close is None or m < close)]
    idx = next(i for i, (m, _) in enumerate(secs) if m == n)
    number = f.get('number') or f'{idx + 1:02d}'
    tracker = [g['title'] for _, g in secs] if len(secs) > 1 else None
    return kit.section(prs, number=number, heading=f['title'], sub=f.get('lede'), tracker=tracker, current=idx)


def _quote_slide(prs, n, f, ctx):
    return kit.quote(prs, text_=f['text'], attribution=f['attribution'])


def _full(prs, n, f, ctx):
    """The picture fills the slide edge to edge; alt text carries what it says for screen readers."""
    s = prs.slides.add_slide(kit.layout(prs, 'Blank'))
    pic = s.shapes.add_picture(str(ctx.resolve(f['image'])), 0, 0, width=kit.W, height=kit.H)
    pic._element.nvPicPr.cNvPr.set('descr', f['alt'])
    return s


def _closing(prs, n, f, ctx):
    return kit.closing(prs, headline=f['title'], lines=f['asks'], cta=f.get('cta'))


def _split(prs, n, f, ctx):
    side = lambda v: (ctx.resolve(v['image']), v['label'], v['stat'], v['headline'], v.get('line'))
    return B.split_hero(prs, side(f['left']), side(f['right']), tag=f.get('tag'), title_=f.get('title'),
                        workdir=ctx.workdir)


def _team(prs, n, f, ctx):
    people = [{**p, 'photo': str(ctx.resolve(p['photo'])), 'logos': [str(ctx.resolve(l)) for l in p.get('logos', [])]}
              for p in f['people']]
    return B.team_slide(prs, eyebrow_text=_eyebrow(ctx.spec, f) or '', title_=f['title'], points=f.get('points', []),
                        credentials=f.get('credentials'), credentials_label=f.get('credentials_label'),
                        people=people, workdir=ctx.workdir, proof=f.get('proof'),
                        stat=(f['stat']['value'], f['stat']['label']) if f.get('stat') else None,
                        logos=[str(ctx.resolve(l)) for l in f.get('logos', [])] or None,
                        logos_label=f.get('logos_label'))


SLIDES = {'split': _split, 'team': _team, 'cover': _cover, 'section': _section, 'content': _content, 'bullets': _bullets,
          'columns': _columns, 'comparison': _comparison, 'numbers': _numbers,
          'quote': _quote_slide, 'picture': _picture, 'closing': _closing, 'full': _full}


def render(spec, out_pptx, workdir=None):
    """Render a validated spec. Returns findings raised during layout (overflow)."""
    out_pptx = Path(out_pptx)
    ctx = Context(spec, workdir or out_pptx.parent)
    prs = kit.new_deck()
    for n, kind, f in slides(spec):
        s = SLIDES[kind](prs, n, f, ctx)
        kit.notes(s, f.get('notes'))
    _stamp(prs, spec)
    _copyright_year(prs, str(spec.get('date', ''))[:4])
    package.finalize(prs)          # whole-number coordinates, notes master listed: strict readers need both
    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_pptx))
    for part, problem in package.audit(out_pptx):
        ctx.findings.append(Finding('error', 'PPTX', 0, f'{part}: {problem}. Keynote and Google Slides will not open it.'))
    return ctx.findings


def _copyright_year(prs, year):
    """The master's copyright line says the template's year; a deck dated another year
       says its own. Only the 'Copyright' shapes on the master and layouts are touched."""
    import re as _re
    if not (year.isdigit() and len(year) == 4):
        return
    master = prs.slide_masters[0]
    for holder in [master, *master.slide_layouts]:
        for sh in holder.shapes:
            if sh.name == 'Copyright' and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.text = _re.sub(r'© \d{4}', f'© {year}', r.text)


def _stamp(prs, spec):
    """Provenance inside the file, so it travels with the deck through Drive."""
    cp = prs.core_properties
    cp.title = spec['deck'] if spec['client'].startswith('_') else f'{spec["client"]} {spec["deck"]}'
    cp.subject = spec['client']
    cp.author = kit.BRAND_NAME
    cp.category = f'{kit.BRAND_SHORT} documents'
    cp.version = __version__
    cp.identifier = f'spec:{spec["_sha"]}'
    cp.keywords = f'servantium-documents {__version__}; spec {spec["_sha"]}; date {spec["date"]}'
    cp.comments = ('Generated from a spec by servantium-documents. Edit freely. '
                   'To rebuild, use the spec in the _source folder next to this file.')
