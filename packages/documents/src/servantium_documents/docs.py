"""Word documents from a spec, on the brand document template.

Decks have their own engine (render.py). A document is simpler: a cover, then a flow of
headings, paragraphs, lists, tables, callouts and figures, set in the built-in styles of
the active brand's document template (kit.TEMPLATE_DOCX), so the result opens in Word and
imports into Google Docs with the same styles (see templates/build_docx.py for why
built-ins only). Contracts (numbered clauses, schedules, signature blocks) and business
forms (invoice, purchase order, change order, acceptance certificate) are documents too.

    python -m servantium_documents.docs <spec.yaml>    -> <spec dir>/out/<name>.docx, .pdf, pages
    build(spec_path) -> {'docx': Path, 'pdf': Path, 'contact': Path, 'pages': int, 'banned': [...]}

Spec:

    document: Discovery Plan            # the file name, after the client
    client: Acme Labs
    prepared_for: {name: Acme Labs, logo: logos/acme.png}   # optional; logo relative to the spec.
                                         # Puts a "Prepared for" lockup beside the brand logo
    draft: true                          # optional header label: true -> "Draft", or the label text
    watermark: Draft                     # optional diagonal watermark (Word and LibreOffice; Google drops it)
    footer: Acme Master Services Agreement  ·  Confidential   # optional; default kit.FOOTER_TEXT
    ref_main: the Agreement              # optional; how a schedule's {ref:} names the main body
    page_numbers: x_of_y                 # optional: "Page 3 of 9"; default is the bare number
    cover:
      eyebrow: Program · Document type
      title: The document title
      subtitle: One line on what it decides and who it is for.
      meta: [[Prepared for, Acme], [Prepared by, <brand>], [Date, October 1, 2026], [Status, Draft]]
      # or, for a working document with no cover page:
      # layout: plain                    # title block on page 1, header and footer on every page
      # byline: <brand> · October 1, 2026
      # or, for a form that carries its own header:
      # layout: none
    body:
      - h1: 1  Heading one               # h1, h2, h3; or {text: ..., page: true} for a new page
                                         # (h1 never breaks the page unless asked; h1_new_page: true
                                         # at the top level makes every h1 after the first break)
      - p: Text with **bold** and [links](https://example.com). {ref:fees} -> "Section 4"
      - bullets: [First point, Second point]
      - numbered: [First step, Second step]   # numbered by hand, so every list starts at 1
      - table:
          head: [Ask, Who, Why]
          rows: [[..., ..., ...]]
          widths: [2, 1, 3.5]            # inches; scaled to fit the page
          size: 9.5                      # body text size, default 10
          center: [1]                    # column indexes set centred
          right: [2]                     # column indexes set right (numeric columns are found anyway)
          money: [2, 3]                  # numbers in these columns print as currency
          currency: USD                  # USD (default), EUR, GBP, CAD, AUD
          rag: [1]                       # status columns: Green/Amber/Red (also G/A/R, On track,
                                         # At risk, Off track, Complete, Accepted, Rejected) become
                                         # chips in kit.SUCCESS / AMBER / CORAL (status roles, not
                                         # the brand primary); exact Green/Amber/Red are found anyway
          foot: [[Total, '', 42000]]     # rows set bold under a rule (totals)
          key: false                     # first column not set bold
          keep: true                     # hold the table on one page; default: only tables
                                         # under 12 rows; longer ones split, header repeated
          # head is optional; rows: [] with a foot draws the foot alone; nothing draws nothing
          strong: [P1]                   # cell values set bold wherever they appear
      - callout: {label: "Decision needed:", text: ...}
      - figure: {svg: figures/timeline.svg, width: 6.5, caption: Figure 1. ...}
      - caption: Table 1. Source ...
      - landscape: true                  # a new landscape section; `portrait: true` returns
      - pagebreak: true                  # never doubled, never before a block that breaks itself

  Contracts:

      - clauses:                         # real Word multilevel numbering: 1 / 1.1 / 1.1.1 / (a)
          - title: Definitions           # a top-level item with a title is an article heading
            id: defs                     # optional, for {ref:defs}
            text: Optional lead paragraph, unnumbered.
            clauses:
              - '"Affiliate" means ...'  # a "Quoted Term" is set bold where it is first defined
              - title: Fees              # a run-in title: "4.1  Fees. text"
                id: fees
                text: Customer pays the fees in {ref:schedule-a}.
                clauses: [..., ...]      # 4.1.1, then (a)
      - schedule: {label: Schedule A, title: Services, id: schedule-a}   # new page, numbering restarts
      - exhibit: {label: Exhibit 1, title: Data Processing Terms}        # the same, named "Exhibit"
      - signature:
          intro: The parties have signed this Agreement on the dates below.
          parties:
            - {company: Acme Labs, Inc., name: '', title: ''}
            - {company: <brand legal name>,   # default: the brand's legal name
               signed: {by: Jordan Lee, date: 'November 1, 2026'}}   # an executed copy
          fields: [By, Name, Title, Date]     # default; printed "By:", "Name:" ...

  Business forms:

      - form:
          type: invoice                  # invoice | purchase_order | change_order | acceptance_certificate
          number: INV-2026-0142
          fields: [[Invoice date, October 20, 2026], [Due date, November 19, 2026], [Terms, Net 30]]
          parties: [{label: Bill to, name: Acme Labs, lines: [100 Main Street, Springfield]}]
          sections: [{title: Description of change, text: ...}]   # before the line items
          items:
            head: [Description, Qty, Rate, Amount]   # a default per type
            rows: [[Discovery workshop, 2, 4500]]    # a missing last Amount is qty x rate
            currency: USD
          totals: {tax: {label: Sales tax, rate: 0.0625}, discount: 0, paid: 0,
                   original: 120000,     # change orders: original value, this change, revised value
                   original_label: ..., change_label: ..., revised_label: ...}   # optional wording
                                         # (also subtotal_label, discount_label, label, paid_label,
                                         # balance_label); a subtotal equal to the total is left out
          # prepared_for shows in the form header (under the number), with or without a cover
          statement: The Client confirms ...           # acceptance certificates
          terms: [Payment due within 30 days ...]
          approvals: [{company: Acme Labs}, {company: <brand>}]   # signature lines

Inline text understands **bold**, [label](url) and {ref:id} only.
"""
import hashlib
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Inches, Pt

from . import output
from . import kit as K
from .brand import hexof
from .spec import SpecError, _StrictLoader
from .templates.build_docx import BODY, borders, cell_margins, run, set_fonts, widths

TEMPLATE = K.TEMPLATE_DOCX
BLOCKS = ('h1', 'h2', 'h3', 'p', 'bullets', 'numbered', 'table', 'callout', 'figure', 'caption',
          'landscape', 'portrait', 'pagebreak', 'clauses', 'schedule', 'exhibit', 'signature', 'form')
FORMS = {
    'invoice':                {'title': 'Invoice', 'head': ['Description', 'Qty', 'Rate', 'Amount'],
                               'widths': [3.6, 0.7, 1.1, 1.1], 'money': [2, 3]},
    'purchase_order':         {'title': 'Purchase Order', 'head': ['Description', 'Qty', 'Unit price', 'Amount'],
                               'widths': [3.6, 0.7, 1.1, 1.1], 'money': [2, 3]},
    'change_order':           {'title': 'Change Order', 'head': ['Change', 'Hours', 'Rate', 'Amount'],
                               'widths': [3.6, 0.7, 1.1, 1.1], 'money': [2, 3]},
    'acceptance_certificate': {'title': 'Acceptance Certificate',
                               'head': ['Deliverable', 'Acceptance criteria', 'Result'],
                               'widths': [2.2, 3.1, 1.2], 'money': [], 'rag': [2]},
}
CURRENCY = {'USD': '$', 'EUR': '€', 'GBP': '£', 'CAD': 'CA$', 'AUD': 'A$'}
_INLINE = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)\s]+\))')
_REF = re.compile(r'\{ref:([A-Za-z0-9_.-]+)\}')
_TERM = re.compile(r'"([A-Z][A-Za-z0-9&\'’ -]{0,60})"')
_NUM = re.compile(r'^\(?[-−]?\s?(?:[A-Z]{0,3}[$€£])?\s?\d[\d,]*(?:\.\d+)?\s?%?\)?$')

# Status chips. A RAG cell is shaded in the brand's STATUS colours, never its primary (a
# brand's primary may be indigo or copper): kit.SUCCESS for green states, kit.AMBER (the
# brand's warning role) and kit.CORAL (its danger role). Text is white or ink by contrast.
# Google keeps cell shading, which is why the chip is the cell and not a shape.
_RAG = {'green': 'SUCCESS', 'g': 'SUCCESS', 'on track': 'SUCCESS', 'accepted': 'SUCCESS', 'done': 'SUCCESS',
        'complete': 'SUCCESS', 'completed': 'SUCCESS', 'closed': 'SUCCESS', 'pass': 'SUCCESS', 'passed': 'SUCCESS',
        'met': 'SUCCESS',
        'amber': 'AMBER', 'a': 'AMBER', 'at risk': 'AMBER', 'accepted with conditions': 'AMBER',
        'conditional': 'AMBER',
        'red': 'CORAL', 'r': 'CORAL', 'off track': 'CORAL', 'rejected': 'CORAL', 'blocked': 'CORAL',
        'fail': 'CORAL', 'failed': 'CORAL', 'not met': 'CORAL'}
_RAG_AUTO = {'green', 'amber', 'red'}
KEEP_ROWS = 12                       # tables with fewer body rows are kept on one page


# ── spec ─────────────────────────────────────────────────────────────────────
def load(path):
    path = Path(path).expanduser().resolve()
    raw = path.read_bytes()
    spec = yaml.load(raw, Loader=_StrictLoader)
    for key in ('document', 'client', 'body'):
        if key not in spec:
            raise SpecError(f'{path.name}: missing "{key}"')
    spec.setdefault('cover', {'layout': 'none'})
    c = spec['cover']
    if c.get('layout') not in (None, 'plain', 'none') :
        raise SpecError(f'{path.name}: cover.layout must be plain or none, got {c.get("layout")!r}')
    if c.get('layout') != 'none' and not c.get('title'):
        raise SpecError(f'{path.name}: cover needs a title (or layout: none)')
    pf = spec.get('prepared_for')
    if pf is not None:
        if not isinstance(pf, dict) or not pf.get('name'):
            raise SpecError(f'{path.name}: prepared_for needs a name: {{name: ..., logo: ...}}')
        if pf.get('logo') and not (path.parent / pf['logo']).exists():
            raise SpecError(f'{path.name}: prepared_for.logo not found: {pf["logo"]}')
    if spec.get('page_numbers', 'x') not in ('x', 'x_of_y'):
        raise SpecError(f'{path.name}: page_numbers must be x or x_of_y')
    for i, b in enumerate(spec['body'], 1):
        if not isinstance(b, dict) or len(b) != 1 or next(iter(b)) not in BLOCKS:
            raise SpecError(f'{path.name}: body item {i} must be one of {", ".join(BLOCKS)}: {b!r:.80}')
        (kind, v), = b.items()
        if kind == 'table':
            every = list(v.get('rows') or []) + list(v.get('foot') or [])
            n = len(v.get('head') or (every[0] if every else []))
            for r in every:
                if len(r) != n:
                    raise SpecError(f'{path.name}: body item {i}: row has {len(r)} cells, head has {n}: {r!r:.80}')
        elif kind == 'form':
            if v.get('type') not in FORMS:
                raise SpecError(f'{path.name}: body item {i}: form type must be one of {", ".join(FORMS)}')
        elif kind in ('schedule', 'exhibit'):
            if not isinstance(v, dict) or not v.get('label'):
                raise SpecError(f'{path.name}: body item {i}: {kind} needs a label, e.g. {{label: Schedule A, title: Services}}')
        elif kind == 'signature':
            if not isinstance(v, dict) or not v.get('parties'):
                raise SpecError(f'{path.name}: body item {i}: signature needs parties')
        elif kind == 'clauses':
            if not isinstance(v, list) or not v:
                raise SpecError(f'{path.name}: body item {i}: clauses must be a list')
    spec['_path'], spec['_sha'] = str(path), hashlib.sha256(raw).hexdigest()[:12]
    return spec


def words(spec):
    """Every piece of reader-facing text, for the banned-phrase check."""
    def walk(v):
        if isinstance(v, str):
            yield v
        elif isinstance(v, list):
            for x in v:
                yield from walk(x)
        elif isinstance(v, dict):
            for k, x in v.items():
                if k not in ('svg', 'widths', 'center', 'strong', 'width', 'logo', 'id', 'type', 'right',
                             'money', 'rag', 'currency'):
                    yield from walk(x)
    yield from walk(spec['cover'])
    yield from walk(spec['body'])


def banned_hits(spec):
    rules = []
    for line in (Path(__file__).parent / 'banned.txt').read_text().splitlines():
        if line.strip() and not line.startswith('#'):
            phrase, _, fix = line.partition('|')
            rules.append((phrase.strip(), fix.strip()))
    text = '\n'.join(words(spec)).lower()
    return [(p, f) for p, f in rules if re.search(rf'\b{re.escape(p.lower())}\b', text)]


# ── inline text ──────────────────────────────────────────────────────────────
def link(p, label, url, *, size=None, bold=False):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts')
    for slot in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        f.set(qn(slot), BODY)
    rpr.append(f)
    if bold:
        rpr.append(OxmlElement('w:b'))
    c = OxmlElement('w:color'); c.set(qn('w:val'), hexof(K.GREEN_600)); rpr.append(c)
    if size:
        s = OxmlElement('w:sz'); s.set(qn('w:val'), str(int(size * 2))); rpr.append(s)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = label; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r)
    p._p.append(h)


def smart(s):
    """Typographer's quotes, so a spec can be typed with straight ones."""
    s = re.sub(r'(^|[\s(\[/–—-])"', '\\1“', str(s)).replace('"', '”')
    s = re.sub(r"(^|[\s(\[])'", '\\1‘', s)
    return s.replace("'", '’')


def inline(p, text, *, size=None, color=None, bold=False):
    for part in _INLINE.split(str(text)):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run(p, smart(part[2:-2]), size=size, color=K.FOREST if color is None else color, bold=True)
        elif part.startswith('[') and re.fullmatch(r'\[([^\]]+)\]\(([^)\s]+)\)', part):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)\s]+)\)', part)
            link(p, smart(m.group(1)), m.group(2), size=size, bold=bold)
        else:                                   # plain text, including a literal "[" 
            run(p, smart(part), size=size, color=color, bold=bold or None)


# ── low-level helpers ────────────────────────────────────────────────────────
def shade(cell, hx):
    """Cell fill. Replaces any fill already set, so a status chip can sit on a zebra row."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hx)
    tcPr.append(shd)


def cell_border(cell, **edges):
    """edges: top/bottom/left/right = (hex, eighths-of-a-point) or None for no line."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcb = tcPr.find(qn('w:tcBorders'))
    if tcb is None:
        tcb = OxmlElement('w:tcBorders'); tcPr.append(tcb)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge not in edges:
            continue
        for old in tcb.findall(qn(f'w:{edge}')):
            tcb.remove(old)
        el = OxmlElement(f'w:{edge}')
        if edges[edge] is None:
            el.set(qn('w:val'), 'nil')
        else:
            hx, sz = edges[edge]
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), hx)
        tcb.append(el)


def no_borders(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'), 'nil'); b.append(el)
    tblPr.append(b)


def valign(cell, where='center'):
    cell.vertical_alignment = {'center': WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                               'bottom': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
                               'top': WD_CELL_VERTICAL_ALIGNMENT.TOP}[where]


def row_height(row, inches, exact=False):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(int(inches * 1440)))
    h.set(qn('w:hRule'), 'exact' if exact else 'atLeast')
    trPr.append(h)


def cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def tight(p, after=0, line=None):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if line:
        p.paragraph_format.line_spacing = line
    return p


def picture(p, path, *, max_w, max_h):
    """A logo fitted inside a box, whatever its aspect."""
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    if w / h > max_w / max_h:
        return p.add_run().add_picture(str(path), width=Inches(max_w))
    return p.add_run().add_picture(str(path), height=Inches(max_h))


def luminance(rgb):
    def ch(c):
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = rgb
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def ink_on(fill):
    """White or ink, whichever reads better on the fill (brand-agnostic contrast)."""
    lf = luminance(fill)
    on_white = (1.05) / (lf + 0.05)
    on_ink = (lf + 0.05) / (luminance(K.INK) + 0.05)
    return K.WHITE if on_white >= on_ink else K.INK


def money(x, cur='USD'):
    sym = CURRENCY.get(str(cur).upper(), str(cur) + ' ')
    if isinstance(x, str):
        return x
    s = f'{sym}{abs(x):,.2f}'
    return f'-{s}' if x < 0 else s


def qty(x):
    if isinstance(x, float) and not x.is_integer():
        return f'{x:,.2f}'.rstrip('0').rstrip('.')
    return f'{int(x):,}' if isinstance(x, (int, float)) else str(x)


def field(p, instr, *, size=9, color=None):
    """A real field (PAGE, NUMPAGES). Word, LibreOffice and Google render it live."""
    for kind in ('begin', 'instr', 'separate', 'text', 'end'):
        r = p.add_run()
        set_fonts(r, BODY); r.font.size = Pt(size); r.font.color.rgb = _C(color or K.INK_MUTED)
        if kind == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {instr} '; r._r.append(it)
        elif kind == 'text':
            t = OxmlElement('w:t'); t.text = '1'; r._r.append(t)
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)


def _C(rgb):
    from docx.shared import RGBColor
    return RGBColor(rgb[0], rgb[1], rgb[2])


# ── schema order ─────────────────────────────────────────────────────────────
# Word rejects ("found unreadable content") property elements out of the schema's
# sequence; LibreOffice and Google forgive it, so the bug only shows on a customer's PC.
# python-docx appends hand-made elements at the end, so every property container is
# put back in schema order (and a repeated single element keeps its last value) before save.
_ORDER = {
    'rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof '
           'snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd '
           'fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath',
    'pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd '
           'tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi '
           'adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc '
           'textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
    'tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc '
             'tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription',
    'tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark',
    'tcBorders': 'top start left bottom end right insideH insideV tl2br tr2bl',
    'tblBorders': 'top start left bottom end right insideH insideV',
    'pBdr': 'top left bottom right between bar',
}
_ORDER = {k: {t: i for i, t in enumerate(v.split())} for k, v in _ORDER.items()}
_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def schema_order(root):
    for local, rank in _ORDER.items():
        for el in root.iter(f'{_W}{local}'):
            kids = list(el)
            if not kids:
                continue
            seen = {}
            for k in kids:                      # one of each; the last one set wins
                if k.tag.startswith(_W) and k.tag[len(_W):] in rank:
                    if k.tag in seen:
                        el.remove(seen[k.tag])
                    seen[k.tag] = k
            kids = list(el)
            known = sorted((k for k in kids if k.tag.startswith(_W) and k.tag[len(_W):] in rank),
                           key=lambda k: rank[k.tag[len(_W):]])
            other = [k for k in kids if k not in known]
            for k in kids:
                el.remove(k)
            for k in known + other:
                el.append(k)


# ── clause numbering ─────────────────────────────────────────────────────────
# One multilevel list for every numbered clause. Level 0 is the article ("1."), then 1.1,
# 1.1.1 and (a). Each schedule or exhibit gets its own w:num on the same definition with
# a start override, so its numbering restarts at 1 and stays a live Word list.
_LEVELS = [  # (format, text, left in, hanging in, bold)
    ('decimal', '%1.', 0.45, 0.45, True),
    ('decimal', '%1.%2', 0.45, 0.45, True),
    ('decimal', '%1.%2.%3', 1.05, 0.6, False),
    ('lowerLetter', '(%4)', 1.45, 0.4, False),
]


def _lvl_xml(i, fmt, text, left, hang, bold):
    b = '<w:b/><w:bCs/>' if bold else ''
    lgl = '<w:isLgl/>' if i and fmt == 'decimal' else ''   # isLgl forces arabic: never on (a)
    f = BODY
    return (f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>{lgl}'
            f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:tabs><w:tab w:val="num" w:pos="{int(left * 1440)}"/></w:tabs>'
            f'<w:ind w:left="{int(left * 1440)}" w:hanging="{int(hang * 1440)}"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{f}" w:hAnsi="{f}" w:eastAsia="{f}" w:cs="{f}"/>{b}'
            f'<w:color w:val="{hexof(K.FOREST)}"/></w:rPr></w:lvl>')


class Numbering:
    def __init__(self, doc):
        self.el = doc.part.numbering_part.element
        ids = [int(a.get(qn('w:abstractNumId'))) for a in self.el.findall(qn('w:abstractNum'))]
        self.abstract = max(ids + [0]) + 1
        lv = ''.join(_lvl_xml(i, *l) for i, l in enumerate(_LEVELS))
        a = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{self.abstract}">'
                      f'<w:multiLevelType w:val="multilevel"/><w:name w:val="Clauses"/>{lv}</w:abstractNum>')
        nums = self.el.findall(qn('w:num'))
        (nums[0].addprevious(a) if nums else self.el.append(a))   # every abstractNum before any num

    def new(self):
        """A fresh list instance that starts at 1."""
        nid = max([int(n.get(qn('w:numId'))) for n in self.el.findall(qn('w:num'))] + [0]) + 1
        ov = ''.join(f'<w:lvlOverride w:ilvl="{i}"><w:startOverride w:val="1"/></w:lvlOverride>'
                     for i in range(len(_LEVELS)))
        self.el.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}">'
                                 f'<w:abstractNumId w:val="{self.abstract}"/>{ov}</w:num>'))
        return nid


def num_pr(p, num_id, ilvl):
    pPr = p._p.get_or_add_pPr()
    n = OxmlElement('w:numPr')
    a = OxmlElement('w:ilvl'); a.set(qn('w:val'), str(ilvl)); n.append(a)
    b = OxmlElement('w:numId'); b.set(qn('w:val'), str(num_id)); n.append(b)
    pPr.append(n)


def _items(v):
    return v if isinstance(v, list) else []


def number_clauses(body):
    """The numbers Word will show, worked out in advance so {ref:id} can print them.
       Returns ({path: number}, {id: (context, ref text)})."""
    nums, refs = {}, {}
    ctx = None
    counters = [0, 0, 0, 0]

    def walk(items, level, path):
        for j, it in enumerate(items):
            counters[level] += 1
            for d in range(level + 1, len(counters)):
                counters[d] = 0
            if level < 3:
                n = '.'.join(str(c) for c in counters[:level + 1])
            else:
                n = '.'.join(str(c) for c in counters[:3]) + f'({chr(96 + counters[3])})'
            nums[path + (j,)] = n
            if isinstance(it, dict):
                if it.get('id'):
                    refs[it['id']] = (ctx, n)
                walk(_items(it.get('clauses')), min(level + 1, 3), path + (j,))

    for i, b in enumerate(body):
        (kind, v), = b.items()
        if kind in ('schedule', 'exhibit'):
            ctx = v['label']
            counters[:] = [0, 0, 0, 0]
            if v.get('id'):
                refs[v['id']] = (None, None, v['label'])
        elif kind == 'clauses':
            walk(v, 0, (i,))
    return nums, refs


# ── the document ─────────────────────────────────────────────────────────────
class Doc:
    def __init__(self, spec):
        self.spec = spec
        self.root = Path(spec['_path']).parent
        if not Path(TEMPLATE).exists():
            raise SpecError(f'No document template for brand "{K.ACTIVE.id}" at {TEMPLATE}. '
                            'Build it: sv templates')
        self.d = Document(str(TEMPLATE))
        body = self.d.element.body
        for child in list(body):
            if child.tag != qn('w:sectPr'):
                body.remove(child)
        self.width = 6.5
        self.figures = 0
        self.nums, self.refs = number_clauses(spec['body'])
        self.numbering = None
        self.num_id = None
        self.context = None
        self.defined = set()
        self.landscape = False

    # text ─────────────────────────────────────────────────────────────────
    def ref(self, s):
        def sub(m):
            key = m.group(1)
            if key not in self.refs:
                raise SpecError(f'{Path(self.spec["_path"]).name}: {{ref:{key}}} names no clause, schedule or exhibit id')
            r = self.refs[key]
            if len(r) == 3:                       # a schedule or exhibit
                return r[2]
            ctx, n = r
            if ctx == self.context:
                return f'Section {n}'
            return f'Section {n} of {ctx or self.spec.get("ref_main", "the Agreement")}'
        return _REF.sub(sub, str(s))

    def inl(self, p, text, **kw):
        inline(p, self.ref(text), **kw)

    def terms(self, s):
        """A "Quoted Term" is set bold the first time the document defines it."""
        def sub(m):
            t = m.group(1)
            if t in self.defined:
                return m.group(0)
            self.defined.add(t)
            return f'**"{t}"**'
        return _TERM.sub(sub, s)

    def para(self, style=None):
        return self.d.add_paragraph(style=style)

    def spacer(self, pts=4):
        p = self.para()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(pts)
        return p

    # chrome ───────────────────────────────────────────────────────────────
    def footer(self, sec, width_in):
        """Company line on the left; the page number (or "Page X of Y") in the bottom-right
           corner, on a right tab at this section's margin."""
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        for el in list(fp._p):
            if el.tag != qn('w:pPr'):
                fp._p.remove(el)
        for extra in sec.footer.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        pPr = fp._p.get_or_add_pPr()
        for t in pPr.findall(qn('w:tabs')):
            pPr.remove(t)
        tabs = fp.paragraph_format.tab_stops
        for pos in (3.25, 6.5):            # the built-in Footer style's centre and right tabs
            if pos != width_in:
                tabs.add_tab_stop(Inches(pos), WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
        label = smart(self.spec.get('footer') or K.FOOTER_TEXT)
        of = self.spec.get('page_numbers') == 'x_of_y'
        room = width_in - K.text_w_in('Page 999 of 999' if of else '999', 9) - 0.3
        fsize = 9
        while fsize > 7.5 and K.text_w_in(label, fsize) > room:
            fsize -= 0.5
        if K.text_w_in(label, fsize) > room:
            while label and K.text_w_in(label + '…', fsize) > room:
                label = label[:-1]
            label = label.rstrip(' ·') + '…'
        run(fp, label + '\t', size=fsize, color=K.INK_MUTED)
        if of:
            run(fp, 'Page ', size=9, color=K.INK_MUTED)
            field(fp, 'PAGE')
            run(fp, ' of ', size=9, color=K.INK_MUTED)
            field(fp, 'NUMPAGES')
        else:
            field(fp, 'PAGE')

    def header(self, hdr, width_in, logo=True):
        """Draft label on the left (when asked), the brand logo on the right."""
        hdr.is_linked_to_previous = False
        hp = hdr.paragraphs[0]
        for el in list(hp._p):
            if el.tag != qn('w:pPr'):
                hp._p.remove(el)
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pPr = hp._p.get_or_add_pPr()
        for t in pPr.findall(qn('w:tabs')):
            pPr.remove(t)
        tabs = hp.paragraph_format.tab_stops
        for pos in (3.25, 6.5):            # the built-in Header style's centre and right tabs
            if pos != width_in:
                tabs.add_tab_stop(Inches(pos), WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
        draft = self.spec.get('draft')
        if draft:
            label = 'Draft' if draft is True else str(draft)
            run(hp, smart(label).upper(), size=8.5, bold=True, color=K.CORAL, spacing=1.2)
        if logo and K.LOGO_ON_LIGHT:
            run(hp, '\t')
            hp.add_run().add_picture(str(K.LOGO_ON_LIGHT), width=Inches(1.15))
        if self.spec.get('watermark'):
            hp._p.append(watermark_run(self.spec['watermark']))

    def chrome(self, sec, width_in, cover):
        # A form's own header carries the logo; the page header then holds only the draft label.
        logo = not any('form' in b for b in self.spec['body'])
        self.header(sec.header, width_in, logo=logo)
        if cover:
            # The cover carries its own logo: its header is the draft label and watermark only.
            self.header(sec.first_page_header, width_in, logo=False)
            sec.first_page_footer.is_linked_to_previous = False
        self.footer(sec, width_in)

    # cover ────────────────────────────────────────────────────────────────
    def lockup(self, pf, *, logo_w=1.9, big=True):
        """The brand logo, a hairline, then "Prepared for" and the client's logo or name."""
        t = self.d.add_table(rows=1, cols=2)
        no_borders(t)
        cell_margins(t, top=0, bottom=0, left=0, right=0)
        a, b = t.cell(0, 0), t.cell(0, 1)
        for c in (a, b):
            valign(c, 'center')
        pa = tight(a.paragraphs[0])
        if K.LOGO_ON_LIGHT:
            pa.add_run().add_picture(str(K.LOGO_ON_LIGHT), width=Inches(logo_w))
        cell_border(b, left=(hexof(K.CLOUD), 8))
        tcPr = b._tc.get_or_add_tcPr()
        mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:left w:w="{int(0.28 * 1440)}" w:type="dxa"/></w:tcMar>')
        tcPr.append(mar)
        p1 = tight(b.paragraphs[0], after=3)
        run(p1, 'PREPARED FOR', size=7.5, bold=True, color=K.INK_MUTED, spacing=1.0)
        p2 = tight(b.add_paragraph())
        if pf.get('logo'):
            picture(p2, self.root / pf['logo'], max_w=2.2, max_h=0.5 if big else 0.36)
        else:
            run(p2, smart(pf['name']), face=K.DISPLAY, size=16 if big else 13, color=K.FOREST)
        widths(t, [logo_w + 0.35, self.width - logo_w - 0.35])
        return t

    def cover(self, c):
        if c.get('layout') == 'plain':
            return self.title_block(c)
        if c.get('layout') == 'none':
            self.d.sections[0].different_first_page_header_footer = False
            return
        pf = self.spec.get('prepared_for')
        if pf:
            self.lockup(pf)
            sp = self.spacer(2)
            sp.paragraph_format.space_after = Pt(84)
        else:
            p = self.para()
            p.add_run().add_picture(str(K.LOGO_ON_LIGHT), width=Inches(1.9))
            p.paragraph_format.space_after = Pt(84)
        if c.get('eyebrow'):
            p = self.para()
            run(p, smart(c['eyebrow']).upper(), size=9.5, bold=True, color=K.GREEN_600, spacing=1.2)
            p.paragraph_format.space_after = Pt(6)
        self.d.add_paragraph(smart(c['title']), style='Title')
        if c.get('subtitle'):
            self.d.add_paragraph(smart(c['subtitle']), style='Subtitle')
        meta = [list(m) for m in (c.get('meta') or [])]
        if pf:
            keys = {str(k).lower() for k, _ in meta}
            if 'prepared for' not in keys:
                meta.insert(0, ['Prepared for', pf['name']])
            if 'prepared by' not in keys:
                meta.insert(1, ['Prepared by', K.BRAND_NAME])
        if meta:
            t = self.d.add_table(rows=len(meta), cols=2)
            borders(t, hexof(K.CLOUD))
            cell_margins(t, top=70, bottom=70, left=0, right=120)
            for i, (k, v) in enumerate(meta):
                a, b = t.cell(i, 0), t.cell(i, 1)
                for c_ in (a, b):
                    c_.paragraphs[0].paragraph_format.space_after = Pt(0)
                run(a.paragraphs[0], str(k).upper(), size=8.5, bold=True, color=K.INK_MUTED, spacing=0.6)
                self.inl(b.paragraphs[0], v, size=10.5, color=K.INK_800)
            widths(t, [1.6, 4.9])
        self.para().add_run().add_break(WD_BREAK.PAGE)

    def title_block(self, c):
        """A working document opens on its title, not a cover page."""
        self.d.sections[0].different_first_page_header_footer = False
        pf = self.spec.get('prepared_for')
        if pf:
            p = tight(self.para(), after=10)
            run(p, 'PREPARED FOR  ', size=7.5, bold=True, color=K.INK_MUTED, spacing=1.0)
            if pf.get('logo'):
                picture(p, self.root / pf['logo'], max_w=1.6, max_h=0.3)
            else:
                run(p, smart(pf['name']), size=10, bold=True, color=K.FOREST)
        t = self.d.add_paragraph(smart(c['title']), style='Title')
        t.paragraph_format.space_before = Pt(6)
        if c.get('subtitle'):
            sub = self.d.add_paragraph(smart(c['subtitle']), style='Subtitle')
            sub.paragraph_format.space_after = Pt(4)
        if c.get('byline'):
            p = self.para()
            p.paragraph_format.space_after = Pt(18)
            run(p, smart(c['byline']), size=9.5, color=K.INK_MUTED)

    def section(self, landscape):
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        sec.different_first_page_header_footer = False
        w, h = (Inches(11), Inches(8.5)) if landscape else (Inches(8.5), Inches(11))
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = w, h
        for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(sec, side, Inches(1))
        self.width = 9.0 if landscape else 6.5
        self.chrome(sec, self.width, cover=False)

    # lists ────────────────────────────────────────────────────────────────
    def numbered(self, items):
        for n, item in enumerate(items, 1):
            p = self.para()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent, pf.space_after = Inches(0.32), Inches(-0.32), Pt(6)
            pf.tab_stops.add_tab_stop(Inches(0.32))
            run(p, f'{n}.\t', bold=True, color=K.FOREST)
            self.inl(p, item)

    def bullets(self, items):
        for item in items:
            self.inl(self.d.add_paragraph(style='List Bullet'), item)

    # tables ───────────────────────────────────────────────────────────────
    def table(self, t, *, width=None, align=WD_TABLE_ALIGNMENT.LEFT, keep_next=False):
        head, rows, foot = list(t.get('head') or []), list(t.get('rows') or []), list(t.get('foot') or [])
        if not rows and not foot:
            return None                         # nothing to draw
        if not rows:
            head = []                           # a foot-only table: no empty header band
        ncol = len(head or (rows or foot)[0])
        width = width or self.width
        size = t.get('size', 10)
        cur = t.get('currency', 'USD')
        center, strong = set(t.get('center', [])), set(map(str, t.get('strong', [])))
        moneyc = set(t.get('money', []))
        ragc = set(t.get('rag', []))
        cells = lambda j: [r[j] for r in rows if r[j] not in (None, '')]
        right = set(t.get('right', [])) | moneyc
        for j in range(1, ncol):                # numeric columns sit right, so digits line up
            vals = cells(j)
            if vals and all(isinstance(v, (int, float)) or _NUM.match(str(v).strip()) for v in vals):
                right.add(j)
        right -= center
        ws = t.get('widths') or [width / ncol] * ncol
        ws = [w * width / sum(ws) for w in ws]
        hr = 1 if head else 0                   # header rows
        tb = self.d.add_table(rows=hr + len(rows) + len(foot), cols=ncol)
        tb.alignment = align
        borders(tb, hexof(K.CLOUD))
        cell_margins(tb, top=45, bottom=45, left=100, right=100)

        def fmt(j, v):
            if v is None:
                return ''
            if j in moneyc and isinstance(v, (int, float)):
                return money(v, cur)
            if isinstance(v, float):
                return qty(v)
            return str(v)

        def place(p, j):
            if j in center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif j in right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        for j, h in enumerate(head):
            c = tb.cell(0, j)
            shade(c, hexof(K.FOREST))
            valign(c, 'bottom')
            p = tight(c.paragraphs[0])
            place(p, j)
            run(p, str(h).upper(), size=8, bold=True, color=K.WHITE, spacing=0.5)
        if head:
            trPr = tb.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)   # repeats on every page
            cant_split(tb.rows[0])
        for i, row in enumerate(rows, start=hr):
            cant_split(tb.rows[i])
            for j, v in enumerate(row):
                c = tb.cell(i, j)
                if (i - hr) % 2 == 1:
                    shade(c, hexof(K.MIST))
                p = tight(c.paragraphs[0], line=1.05)
                place(p, j)
                s = fmt(j, v)
                chip = _RAG.get(s.strip().lower()) if s and (j in ragc or s.strip().lower() in _RAG_AUTO) else None
                if chip:
                    fill = getattr(K, chip)
                    shade(c, hexof(fill))
                    valign(c, 'center')
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run(p, s.upper(), size=size - 1.5, bold=True, color=ink_on(fill), spacing=0.6)
                elif j == 0 and t.get('key', True):
                    self.inl(p, s, size=size, color=K.FOREST, bold=True)
                elif s in strong:
                    run(p, s, size=size, bold=True, color=K.FOREST)
                else:
                    self.inl(p, s, size=size, color=K.INK_800)
        for k, row in enumerate(foot):
            i = hr + len(rows) + k
            cant_split(tb.rows[i])
            for j, v in enumerate(row):
                c = tb.cell(i, j)
                shade(c, hexof(K.MINT_TINT))
                if k == 0:
                    cell_border(c, top=(hexof(K.FOREST), 8))
                p = tight(c.paragraphs[0])
                place(p, j)
                run(p, smart(fmt(j, v)), size=size, bold=True, color=K.FOREST)
        # A short table (under KEEP_ROWS body rows) stays on one page. A long one splits across
        # pages with its header repeated, keeping only the header with the first row and the
        # last rows with the foot, so neither a lone header nor lone totals start a page.
        # `keep: true|false` in the spec overrides; keep_next also holds the last row to
        # whatever follows (a form's totals).
        n = len(tb.rows)
        keep = t.get('keep', len(rows) < KEEP_ROWS)
        held = set(range(n - 1)) if keep else (set(range(hr)) | set(range(max(hr, n - len(foot) - 2), n - 1)))
        if keep_next:
            held.add(n - 1)
        for ri in held:
            for c in tb.rows[ri].cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
        if ragc or any(_RAG.get(fmt(j, v).strip().lower()) for row in rows for j, v in enumerate(row)
                       if fmt(j, v).strip().lower() in _RAG_AUTO):
            for r in tb.rows[hr:]:
                for c in r.cells:
                    valign(c, 'center')
        widths(tb, ws)
        sp = self.spacer(6)
        sp.paragraph_format.keep_with_next = keep_next
        return tb

    def callout(self, c):
        tb = self.d.add_table(rows=1, cols=1)
        tb.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = tb.cell(0, 0)
        shade(cell, hexof(K.MINT_TINT))
        cell_border(cell, left=(hexof(K.GREEN), 24))
        cell_margins(tb, top=140, bottom=140, left=220, right=220)
        widths(tb, [self.width])
        cant_split(tb.rows[0])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run(p, smart(c['label']) + ' ', bold=True, color=K.FOREST)
        self.inl(p, c['text'], color=K.INK_800)
        self.spacer(8)

    def figure(self, f):
        svg = (self.root / f['svg']).resolve()
        png = self.out / (svg.stem + '.png')
        svg_to_png(svg, png)
        p = self.para()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(png), width=Inches(min(f.get('width', self.width), self.width)))
        if f.get('caption'):
            self.inl(self.d.add_paragraph(style='Caption'), f['caption'])

    # contracts ────────────────────────────────────────────────────────────
    def clauses(self, items, path):
        if self.numbering is None:
            self.numbering = Numbering(self.d)
        if self.num_id is None:
            self.num_id = self.numbering.new()
        self._clauses(items, 0, path)

    def _clauses(self, items, level, path):
        for j, it in enumerate(items):
            here = path + (j,)
            it = {'text': it} if isinstance(it, str) else it
            title, text = it.get('title'), it.get('text')
            if level == 0 and title:
                h = self.d.add_paragraph(smart(self.ref(title)), style='Heading 2')
                num_pr(h, self.num_id, 0)
                h.paragraph_format.keep_with_next = True
                if text:
                    p = self.para()
                    p.paragraph_format.left_indent = Inches(_LEVELS[0][2])
                    self.inl(p, self.terms(text))
            else:
                p = self.para()
                num_pr(p, self.num_id, level)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.22
                if title:
                    run(p, smart(self.ref(title)).rstrip('.') + '. ', bold=True, color=K.FOREST)
                if text:
                    self.inl(p, self.terms(text))
                if it.get('clauses'):
                    p.paragraph_format.keep_with_next = True
            self._clauses(_items(it.get('clauses')), min(level + 1, 3), here)

    def schedule(self, v, word):
        """Schedules and exhibits start a new page and restart clause numbering at 1."""
        self.context = v['label']
        if self.numbering is not None or any('clauses' in b for b in self.spec['body']):
            if self.numbering is None:
                self.numbering = Numbering(self.d)
            self.num_id = self.numbering.new()
        h = self.d.add_paragraph(smart(v['label'] + (f': {v["title"]}' if v.get('title') else '')),
                                 style='Heading 1')
        h.paragraph_format.page_break_before = True
        if v.get('text'):
            self.inl(self.para(), v['text'])

    def signature(self, v, *, heading=None, compact=False):
        """Two parties side by side: company, then a ruled line for each field. Each pair of
           parties is kept on one page (with its heading and intro). `compact` (a form's
           approvals) uses shorter rows, so the block fits under the totals on the form's page."""
        parties = [p if isinstance(p, dict) else {'company': p} for p in v['parties']]
        fields = v.get('fields') or ['By', 'Name', 'Title', 'Date']
        if heading:
            h = self.d.add_paragraph(heading, style='Heading 3')
            if compact:
                h.paragraph_format.space_before = Pt(4)
        if v.get('intro'):
            p = self.para()
            p.paragraph_format.keep_with_next = True
            self.inl(p, v['intro'])
        lab, gap = 0.75, 0.4
        line = (self.width - gap - 2 * lab) / 2
        for k in range(0, len(parties), 2):
            pair = parties[k:k + 2]
            tb = self.d.add_table(rows=1 + len(fields), cols=5)
            no_borders(tb)
            cell_margins(tb, top=20, bottom=20, left=0, right=80)
            for col, party in zip((0, 3), pair):
                head = tb.cell(0, col).merge(tb.cell(0, col + 1))
                p = tight(head.paragraphs[0], after=2, line=1.0)
                p.paragraph_format.keep_with_next = True
                run(p, smart(party.get('company') or K.ACTIVE.legal_name), size=10.5, bold=True, color=K.FOREST)
                for r, f in enumerate(fields, start=1):
                    a, b = tb.cell(r, col), tb.cell(r, col + 1)
                    for c in (a, b):
                        valign(c, 'bottom')
                    pa = tight(a.paragraphs[0], line=1.0)
                    pa.paragraph_format.keep_with_next = True
                    run(pa, str(f).rstrip(':') + ':', size=8.5, bold=True, color=K.INK_MUTED)
                    pb = tight(b.paragraphs[0], line=1.0)
                    pb.paragraph_format.keep_with_next = r < len(fields)
                    key = str(f).lower().rstrip(':')
                    signed = party.get('signed') or {}
                    if key in ('by', 'signature'):
                        # An executed copy prints who signed on the line; a blank copy leaves it to sign.
                        if signed.get('by'):
                            run(pb, smart(signed['by']), face=K.DISPLAY, size=13, color=K.FOREST)
                        val = None
                    else:
                        val = party.get(key) or (signed.get('by') if key == 'name' else None) \
                            or (signed.get('date') if key == 'date' else None)
                    if val:
                        run(pb, smart(val), size=10, color=K.INK_800)
                    cell_border(b, bottom=(hexof(K.INK_MUTED if r == 1 else K.CLOUD), 6 if r == 1 else 4))
            for r in range(len(fields) + 1):
                for c in tb.rows[r].cells:          # the gap column too: Normal's spacing sets row height
                    for p in c.paragraphs:
                        tight(p, line=1.0)
                        p.paragraph_format.keep_with_next = r < len(fields)
                cant_split(tb.rows[r])
                row_height(tb.rows[r], (0.42 if r == 1 else 0.25) if compact else (0.55 if r == 1 else 0.3))
            widths(tb, [lab, line, gap, lab, line])
            self.spacer(10)

    # forms ────────────────────────────────────────────────────────────────
    def form(self, f):
        kind = FORMS[f['type']]
        cur = (f.get('items') or {}).get('currency', 'USD')
        # Header: who is issuing it on the left, what it is on the right.
        t = self.d.add_table(rows=1, cols=2)
        no_borders(t)
        cell_margins(t, top=0, bottom=0, left=0, right=0)
        a, b = t.cell(0, 0), t.cell(0, 1)
        valign(a, 'top'); valign(b, 'top')
        pa = tight(a.paragraphs[0], after=8)
        if K.LOGO_ON_LIGHT:
            pa.add_run().add_picture(str(K.LOGO_ON_LIGHT), width=Inches(1.6))
        p = tight(a.add_paragraph(), after=1)
        run(p, smart(K.ACTIVE.legal_name), size=9, bold=True, color=K.INK_800)
        for line in K.ACTIVE.address or []:
            run(tight(a.add_paragraph(), after=0), smart(line), size=9, color=K.INK_MUTED)
        if K.ACTIVE.website:
            run(tight(a.add_paragraph(), after=0), K.ACTIVE.website, size=9, color=K.INK_MUTED)
        pt = b.paragraphs[0]
        pt.style = self.d.styles['Title']
        tight(pt, after=2, line=1.0)
        pt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        title = f.get('title') or kind['title']
        size = 26
        while size > 16 and K.text_w_in(title, size, font=K.DISPLAY) > self.width - 3.7:
            size -= 1                           # "Acceptance Certificate" stays on one line
        run(pt, title, face=K.DISPLAY, size=size)
        if f.get('number'):
            pn = tight(b.add_paragraph())
            pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run(pn, 'NO.  ', size=8, bold=True, color=K.INK_MUTED, spacing=0.8)
            run(pn, str(f['number']), size=11, bold=True, color=K.FOREST)
        if f.get('status'):
            ps = tight(b.add_paragraph(), after=0)
            ps.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run(ps, smart(f['status']).upper(), size=8, bold=True, color=K.CORAL, spacing=1.0)
        pf = self.spec.get('prepared_for')
        if pf:
            # The co-brand on a form: whom it is for, small, under the form's number.
            # One line (label, then the logo or name), so the header grows by a line, not two.
            pl = tight(b.add_paragraph())
            pl.paragraph_format.space_before = Pt(8)
            pl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run(pl, 'PREPARED FOR  ', size=7, bold=True, color=K.INK_MUTED, spacing=1.0)
            if pf.get('logo'):
                picture(pl, self.root / pf['logo'], max_w=1.4, max_h=0.28)
            else:
                run(pl, smart(pf['name']), size=10.5, bold=True, color=K.FOREST)
        widths(t, [3.6, self.width - 3.6])
        rule = tight(self.para(), after=14)
        pPr = rule._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="6" '
                             f'w:color="{hexof(K.GREEN)}"/></w:pBdr>'))
        rule.paragraph_format.line_spacing = Pt(6)

        if f.get('fields'):
            self.field_grid(f['fields'])
        if f.get('parties'):
            self.parties(f['parties'])
        for s in f.get('sections') or []:
            self.d.add_paragraph(smart(s['title']), style='Heading 3')
            for para in (s['text'] if isinstance(s['text'], list) else [s['text']]):
                self.inl(self.para(), para)
        items = f.get('items')
        subtotal = 0
        if items:
            head = items.get('head') or kind['head']
            moneyc = items.get('money', kind['money'])
            rows = []
            for r in items['rows']:
                r = list(r)
                if len(r) == len(head) - 1 and moneyc and head[-1].lower() in ('amount', 'total', 'fee'):
                    q, rate = r[-2], r[-1]
                    r.append((q or 0) * (rate or 0))
                if moneyc and isinstance(r[-1], (int, float)):
                    subtotal += r[-1]
                rows.append(r)
            if f.get('items_title'):
                self.d.add_paragraph(smart(f['items_title']), style='Heading 3')
            self.table({'head': head, 'rows': rows, 'widths': items.get('widths') or kind['widths'],
                        'money': moneyc, 'currency': cur, 'rag': items.get('rag', kind.get('rag', [])),
                        'size': items.get('size', 9.5), 'key': False}, keep_next=bool(moneyc))
            if moneyc:
                self.totals(f, subtotal, cur)
        if f.get('statement'):
            p = self.para()
            p.paragraph_format.space_before = Pt(6)
            self.inl(p, f['statement'])
        for label in ('notes', 'terms'):
            if f.get(label):
                self.d.add_paragraph(label.title(), style='Heading 3')
                for para in (f[label] if isinstance(f[label], list) else [f[label]]):
                    p = self.para()
                    p.paragraph_format.space_after = Pt(4)
                    self.inl(p, para, size=9.5, color=K.INK_800)
        if f.get('approvals'):
            self.spacer(6)
            self.signature({'parties': f['approvals'], 'intro': f.get('approvals_intro'),
                            'fields': f.get('approval_fields') or ['By', 'Name', 'Title', 'Date']},
                           heading=f.get('approvals_title', 'Approval'), compact=True)

    def field_grid(self, fields, cols=3):
        rows = (len(fields) + cols - 1) // cols
        t = self.d.add_table(rows=rows, cols=cols)
        no_borders(t)
        cell_margins(t, top=70, bottom=70, left=120, right=120)
        for i, (k, v) in enumerate(fields):
            c = t.cell(i // cols, i % cols)
            pk = tight(c.paragraphs[0], after=1)
            run(pk, str(k).upper(), size=7.5, bold=True, color=K.INK_MUTED, spacing=0.6)
            pv = tight(c.add_paragraph())
            self.inl(pv, '' if v is None else str(v), size=10.5, color=K.INK)
        for r in t.rows:
            cant_split(r)
            for c in r.cells:
                shade(c, hexof(K.MIST))
                cell_border(c, left=(hexof(K.WHITE), 18))
        widths(t, [self.width / cols] * cols)
        self.spacer(10)

    def parties(self, parties):
        t = self.d.add_table(rows=1, cols=len(parties))
        no_borders(t)
        cell_margins(t, top=0, bottom=0, left=0, right=200)
        for c, party in zip(t.rows[0].cells, parties):
            p = tight(c.paragraphs[0], after=3)
            run(p, smart(party.get('label', '')).upper(), size=7.5, bold=True, color=K.GREEN_600, spacing=1.0)
            p = tight(c.add_paragraph(), after=1)
            run(p, smart(party.get('name') or K.ACTIVE.legal_name), size=10.5, bold=True, color=K.FOREST)
            lines = party.get('lines')
            if lines is None and not party.get('name'):
                lines = K.ACTIVE.address
            for line in lines or []:
                run(tight(c.add_paragraph()), smart(line), size=9.5, color=K.INK_800)
        widths(t, [self.width / len(parties)] * len(parties))
        cant_split(t.rows[0])
        self.spacer(12)

    def totals(self, f, subtotal, cur):
        tot = f.get('totals') or {}
        lines = []
        if tot.get('original') is not None:
            lines += [(tot.get('original_label', 'Original contract value'), tot['original']),
                      (tot.get('change_label', 'This change order'), subtotal)]
            total = tot['original'] + subtotal
            total_label = tot.get('revised_label', 'Revised contract value')
        else:
            lines.append((tot.get('subtotal_label', 'Subtotal'), subtotal))
            total = subtotal
            if tot.get('discount'):
                lines.append((tot.get('discount_label', 'Discount'), -abs(tot['discount']))); total -= abs(tot['discount'])
            if tot.get('tax'):
                tx = tot['tax']
                amount = round(total * tx['rate'], 2)
                lines.append((f'{tx.get("label", "Tax")} ({tx["rate"] * 100:g}%)', amount)); total += amount
            if tot.get('shipping'):
                lines.append((tot.get('shipping_label', 'Shipping'), tot['shipping'])); total += tot['shipping']
            if len(lines) == 1:
                lines = []                      # a subtotal that equals the total says nothing
            total_label = tot.get('label', 'Total')
            if tot.get('paid'):
                lines += [(total_label, total), (tot.get('paid_label', 'Paid'), -abs(tot['paid']))]
                total, total_label = total - abs(tot['paid']), tot.get('balance_label', 'Balance due')
        # The label column is as wide as its longest label (the total's is set in tracked
        # capitals), so "Revised contract value" never wraps; the amount column is fixed.
        amt_w = 1.35
        need = max([K.text_w_in(l, 9.5) for l, _ in lines] +
                   [K.text_w_in(total_label.upper(), 8.5, bold=True) * 1.12]) + 0.25
        lab_w = min(max(1.85, need), self.width - amt_w - 0.5)
        t = self.d.add_table(rows=len(lines) + 1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        no_borders(t)
        cell_margins(t, top=50, bottom=50, left=100, right=100)
        for i, (label, amount) in enumerate(lines + [(total_label, total)]):
            a, b = t.cell(i, 0), t.cell(i, 1)
            last = i == len(lines)
            pa, pb = tight(a.paragraphs[0]), tight(b.paragraphs[0])
            pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            valign(a, 'center'); valign(b, 'center')
            pa.paragraph_format.keep_with_next = pb.paragraph_format.keep_with_next = not last
            if last:
                for c in (a, b):
                    shade(c, hexof(K.FOREST))
                run(pa, smart(label).upper(), size=8.5, bold=True, color=K.WHITE, spacing=0.6)
                run(pb, money(amount, cur), size=11.5, bold=True, color=K.WHITE)
            else:
                for c in (a, b):
                    cell_border(c, bottom=(hexof(K.CLOUD), 4))
                run(pa, smart(label), size=9.5, color=K.INK_MUTED)
                run(pb, money(amount, cur), size=10, color=K.INK_800)
            cant_split(t.rows[i])
        widths(t, [lab_w, amt_w])
        self.spacer(10)

    def page_break(self):
        """A break that never leaves a blank page. A break character in its own paragraph adds a
           page when the content before it already filled one (a form that ends at the foot of
           its page). Instead a hairline paragraph carries "page break before": at the top of a
           page already, it adds nothing. The empty spacer after a table is reused for it."""
        body = self.d.element.body
        kids = [k for k in body if k.tag != qn('w:sectPr')]
        last = kids[-1] if kids else None
        if (last is not None and last.tag == qn('w:p') and not last.findall(qn('w:r'))
                and last.find(f'{qn("w:pPr")}/{qn("w:sectPr")}') is None):
            p = last
        else:
            p = self.para()._p
        pPr = p.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
        for tag in ('w:spacing', 'w:keepNext', 'w:pageBreakBefore', 'w:rPr'):
            for old in pPr.findall(qn(tag)):
                pPr.remove(old)
        pPr.append(OxmlElement('w:pageBreakBefore'))
        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'))
        pPr.append(parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>'))

    def last_spacer(self):
        """A document that ends on a table must end on a paragraph after it. Shrink that empty
           paragraph to a hairline, so a table that fills the page does not push a blank page."""
        body = self.d.element.body
        kids = [k for k in body if k.tag != qn('w:sectPr')]
        if not kids or kids[-1].tag != qn('w:p') or kids[-1].findall(qn('w:r')):
            return
        p = kids[-1]
        pPr = p.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
        for old in pPr.findall(qn('w:spacing')):
            pPr.remove(old)
        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'))
        pPr.append(parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>'))

    # build ────────────────────────────────────────────────────────────────
    def build(self, out):
        self.out = out
        c = self.spec['cover']
        cover = c.get('layout') not in ('plain', 'none')
        self.cover(c)
        self.chrome(self.d.sections[0], 6.5, cover=cover)
        body = self.spec['body']
        for i, b in enumerate(body):
            (kind, v), = b.items()
            if kind in ('h1', 'h2', 'h3'):
                txt, page = (v.get('text', ''), v.get('page')) if isinstance(v, dict) else (v, None)
                h = self.d.add_paragraph(smart(self.ref(txt)), style=f'Heading {kind[1]}')
                if kind == 'h1':
                    # A new page only when asked (h1: {text, page: true} or h1_new_page: true),
                    # whatever the brand template's Heading 1 says. The first heading never breaks.
                    want = page if page is not None else bool(self.spec.get('h1_new_page'))
                    h.paragraph_format.page_break_before = bool(want) and i > 0
            elif kind == 'p':
                p = self.para()
                self.inl(p, v)
                # A lead-in paragraph stays on the page with the figure or table it introduces.
                if i + 1 < len(body) and ({'figure', 'table', 'signature'} & set(body[i + 1])):
                    p.paragraph_format.keep_with_next = True
            elif kind == 'caption':
                self.inl(self.d.add_paragraph(style='Caption'), v)
            elif kind == 'bullets':
                self.bullets(v)
            elif kind == 'numbered':
                self.numbered(v)
            elif kind == 'table':
                self.table(v)
            elif kind == 'callout':
                self.callout(v)
            elif kind == 'figure':
                self.figure(v)
            elif kind == 'clauses':
                self.clauses(v, (i,))
            elif kind in ('schedule', 'exhibit'):
                self.schedule(v, kind)
            elif kind == 'signature':
                self.signature(v)
            elif kind == 'form':
                self.form(v)
            elif kind == 'landscape':
                self.section(True)
            elif kind == 'portrait':
                self.section(False)
            elif kind == 'pagebreak':
                nxt = next(iter(body[i + 1])) if i + 1 < len(body) else None
                prev = next(iter(body[i - 1])) if i else None
                # Blocks that start their own page make a second break a blank page.
                if prev != 'pagebreak' and nxt not in (None, 'schedule', 'exhibit', 'landscape', 'portrait'):
                    self.page_break()
        self.last_spacer()
        cp = self.d.core_properties
        cp.title = c.get('title') or self.spec['document']
        cp.author = K.BRAND_NAME
        cp.last_modified_by = K.BRAND_NAME
        cp.subject = self.spec['document']
        cp.comments = f'spec {Path(self.spec["_path"]).name} sha {self.spec["_sha"]}'
        for part in [self.d.part] + [r.target_part for r in self.d.part.rels.values()
                                     if not r.is_external and r.reltype in (RT.HEADER, RT.FOOTER)]:
            schema_order(part.element)
        return self.d


def watermark_run(text):
    """Word's own watermark: a VML text path behind the page, as Word writes it. LibreOffice
       renders it; Google Docs drops it, so a draft that goes to Google should use `draft:`."""
    t = smart(str(text)).upper().replace('&', '&amp;').replace('"', '&quot;')
    color = hexof(K.CLOUD)
    font = K.UI.replace('"', '')
    return parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
        f'xmlns:w10="urn:schemas-microsoft-com:office:word"><w:rPr><w:noProof/></w:rPr><w:pict>'
        '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
        'path="m@7,l@8,m@5,21600l@6,21600e"><v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/>'
        '<v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/>'
        '<v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/>'
        '<v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>'
        '<v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" '
        'o:connectangles="270,180,90,0"/><v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" '
        'xrange="6629,14971"/></v:handles><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>'
        '<v:shape id="PowerPlusWaterMarkObject1" o:spid="_x0000_s2049" type="#_x0000_t136" '
        'style="position:absolute;margin-left:0;margin-top:0;width:440pt;height:110pt;rotation:315;'
        'z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        'mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" '
        f'fillcolor="#{color}" stroked="f"><v:fill opacity=".4"/>'
        f'<v:textpath style="font-family:&quot;{font}&quot;;font-size:1pt" string="{t}"/>'
        '<w10:wrap anchorx="margin" anchory="margin"/></v:shape></w:pict></w:r>')


def svg_to_png(svg, png):
    """Headless Chrome, the same printer `sv diagram` uses, so figures get brand fonts."""
    from . import diagram                  # only figures need Chrome
    js = Path(tempfile.gettempdir()) / 'sv-diagram-render.cjs'
    js.write_text(diagram._RENDER_JS)
    pdf = Path(tempfile.gettempdir()) / (svg.stem + '.pdf')
    env = {**os.environ, 'NODE_PATH': diagram._node_path()}
    subprocess.run(['node', str(js), str(svg), str(pdf), str(png)], check=True, env=env, capture_output=True)
    return png


def contact_sheet(pdf, out_png, *, cols=4, dpi=40):
    """Like output.contact_sheet, but every page sits in a square cell, so landscape
       sections do not run into the next tile."""
    import pymupdf
    from PIL import Image, ImageDraw
    doc = pymupdf.open(str(pdf))
    side = int(11 * dpi)
    tiles = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi)
        im = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
        cell = Image.new('RGB', (side, side), (154, 163, 173))
        cell.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        d = ImageDraw.Draw(cell)
        d.rectangle((0, 0, 34, 22), fill=tuple(K.FOREST))
        d.text((8, 5), str(i + 1), fill=(255, 255, 255))
        tiles.append(cell)
    cols = min(cols, len(tiles))
    rows, g = (len(tiles) + cols - 1) // cols, 10
    sheet = Image.new('RGB', (cols * side + (cols + 1) * g, rows * side + (rows + 1) * g), (154, 163, 173))
    for j, im in enumerate(tiles):
        sheet.paste(im, (g + (j % cols) * (side + g), g + (j // cols) * (side + g)))
    sheet.save(out_png)
    return out_png, len(doc)


def build(spec_path):
    spec = load(spec_path)
    out = Path(spec['_path']).parent / 'out'
    out.mkdir(exist_ok=True)
    stem = f'{spec["client"]} {spec["document"]}'
    docx = out / f'{stem}.docx'
    Doc(spec).build(out).save(str(docx))
    pdf = output.to_pdf(docx)
    sheet, pages = contact_sheet(pdf, out / f'{stem} contact.png')
    banned = banned_hits(spec)
    findings = [f'warning  "{p}": {f}' for p, f in banned]
    return {'docx': docx, 'pdf': pdf, 'contact': sheet, 'pages': pages, 'banned': banned, 'findings': findings}


if __name__ == '__main__':
    try:
        r = build(sys.argv[1])
    except SpecError as e:
        sys.exit(f'sv: {e}')
    for phrase, fix in r['banned']:
        print(f'  warning  "{phrase}": {fix}')
    print(f'  {r["pages"]} pages')
    for k in ('docx', 'pdf', 'contact'):
        print(f'  {k:<8}{r[k]}')
